"""Anforderungsbezogene Tests für Konten, Sitzungen, VP-Daten und Importe."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import DEFAULT, patch
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from defusedxml.ElementTree import fromstring as geschuetztes_xml
from defusedxml.common import EntitiesForbidden

from django.conf import settings
from django.contrib.auth import SESSION_KEY, get_user
from django.contrib.auth.hashers import get_hasher
from django.contrib.auth.password_validation import validate_password
from django.contrib.sessions.backends.db import SessionStore
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile, TemporaryUploadedFile
from django.db import DatabaseError, IntegrityError, transaction
from django.http import HttpRequest
from django.middleware.csrf import get_token
from django.test import Client, SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.external_link.external import (
    ExternalBook,
    ExternalLink,
    ExternalSheetNames,
)
from openpyxl.writer.excel import ExcelWriter
from openpyxl.xml.constants import CONTYPES_NS, PKG_REL_NS, VBA, XLSM, XLSX, XLTM

from vp_stunden.models import Benutzerkonto, Untersuchung, Teilnahme
from vp_stunden.sicherheit import (
    aendere_passwort,
    lege_konto_an,
    melde_ab,
    melde_an,
    pruefe_zugriff,
    setze_passwort_zurueck,
)
from vp_stunden.vp_daten import (
    berechne_stunden,
    korrigiere_untersuchung,
    loesche_teilnahmedaten,
    pruefe_einordnung,
    pruefe_lehrperson,
    pruefe_matrikelnummer,
    pruefe_studienname,
    pruefe_stundenwert,
    suche_teilnahmen,
    suche_untersuchungen,
)
from vp_stunden.reports import erzeuge_report
from vp_stunden.views import melde_an_view
from vp_stunden.importe import (
    importiere_uebergangsdatei,
    importiere_untersuchung,
)


class PasswortregelnTests(SimpleTestCase):
    """Prüft Mindestlänge und zulässige Passwortzeichen nach QA-01."""

    def test_qa01_14_zeichen_werden_abgewiesen(self):
        """Ein Passwort unterhalb der Mindestlänge wird abgewiesen."""
        with self.assertRaises(ValidationError):
            validate_password('a' * 14)

    def test_qa01_15_zeichen_werden_akzeptiert(self):
        """Ein Passwort genau an der Mindestlänge wird akzeptiert."""
        self.assertIsNone(validate_password('a' * 15))

    def test_qa01_64_zeichen_werden_akzeptiert(self):
        """Ein Passwort mit 64 Zeichen bleibt zulässig."""
        self.assertIsNone(validate_password('a' * 64))

    def test_qa01_65_zeichen_werden_akzeptiert(self):
        """Es gilt keine zusätzliche Obergrenze von 64 Zeichen."""
        self.assertIsNone(validate_password('a' * 65))

    def test_qa01_reines_zahlenpasswort_wird_akzeptiert(self):
        """Ausreichend lange Zahlenpasswörter brauchen keine weiteren Zeichenklassen."""
        self.assertIsNone(validate_password('123456789012345'))


class PasswortHashingTests(SimpleTestCase):
    """Prüft Passwortvergleich, Argon2id-Parameter und individuelle Salts."""

    def test_qa01_passwortvergleich(self):
        """Passwortvergleich, Hashparameter und neuer Salt erfüllen das gewählte Profil."""
        testpasswort = 'a' * 64
        konto = Benutzerkonto(username='testuser', rolle=Benutzerkonto.LEHRPERSON)
        konto.set_password(testpasswort)
        self.assertNotEqual(konto.password, testpasswort)
        self.assertTrue(konto.check_password(testpasswort))

        testpasswort_falsch = 'a' * 63 + 'b'
        self.assertFalse(konto.check_password(testpasswort_falsch))

        hasher = get_hasher()
        # decode liest Algorithmus und Parameter aus dem Hash, ohne das Passwort wiederherzustellen.
        hasher_decoded = hasher.decode(konto.password)
        self.assertEqual(hasher_decoded['algorithm'], 'argon2')
        self.assertEqual(hasher_decoded['time_cost'], 2)
        self.assertEqual(hasher_decoded['memory_cost'], 102400)
        self.assertEqual(hasher_decoded['parallelism'], 8)
        self.assertEqual(hasher_decoded['variety'], 'argon2id')

        konto.set_password(testpasswort)
        # Gleiches Passwort erneut hashen: Ein neuer Salt muss den gespeicherten Wert verändern.
        hasher_decoded_neu = hasher.decode(konto.password)

        self.assertNotEqual(hasher_decoded['salt'], hasher_decoded_neu['salt'])


class PasswortspeicherungTests(TestCase):
    """Prüft die unveränderte Speicherung des Passwort-Hashes."""

    def test_qa01_hash_wird_gespeichert(self):
        """Der erneut geladene Hash entspricht dem gespeicherten Wert und akzeptiert das Passwort."""
        testpasswort = 'a' * 64
        konto = Benutzerkonto(username='testuser', rolle=Benutzerkonto.LEHRPERSON)
        konto.set_password(testpasswort)
        passwort_hash = konto.password
        konto.save()

        # Erneut laden, damit die Assertions die Datenbankpersistenz nachweisen.
        konto_aus_db = Benutzerkonto.objects.get(username='testuser')
        self.assertEqual(konto_aus_db.password, passwort_hash)
        self.assertFalse(konto_aus_db.password == testpasswort)
        self.assertTrue(konto_aus_db.check_password(testpasswort))


class KontoErstellungTests(TestCase):
    """Prüft Kontenanlage und kontrollierte Abweisung ungültiger Eingaben."""

    def setUp(self):
        """Bereitet gültige Eingaben vor, ohne das zu prüfende Konto anzulegen."""
        self.benutzername = 'testnutzer'
        self.passwort = 'a' * 64
        self.rolle = Benutzerkonto.LEHRPERSON

    def test_qa01_kurzes_passwort_verhindert_kontoerstellung(self):
        """Ein zu kurzes Passwort verhindert jede Kontospeicherung."""
        self.passwort = 'a' * 14
        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNone(konto)
        self.assertTrue(fehlerliste)
        self.assertFalse(Benutzerkonto.objects.filter(username=self.benutzername).exists())

    def test_fa02_erfolgreiche_kontoerstellung(self):
        """Das Konto speichert Benutzername, Rolle und prüfbaren Passwort-Hash."""
        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)
        self.assertTrue(Benutzerkonto.objects.filter(username=self.benutzername).exists())

        # Erneut laden, damit die Assertions die Datenbankpersistenz nachweisen.
        benutzerkonto_aus_db = Benutzerkonto.objects.get(username=self.benutzername)
        self.assertTrue(benutzerkonto_aus_db.check_password(self.passwort))
        self.assertEqual(benutzerkonto_aus_db.rolle, self.rolle)
        self.assertNotEqual(benutzerkonto_aus_db.password, self.passwort)
        self.assertEqual(benutzerkonto_aus_db.username, self.benutzername)

    def test_fa02_ungueltige_rolle_verhindert_kontoerstellung(self):
        """Eine unzulässige Rolle wird ohne Kontospeicherung abgewiesen."""
        self.rolle = 'ungültig'
        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNone(konto)
        self.assertTrue(fehlerliste)
        self.assertFalse(Benutzerkonto.objects.filter(username=self.benutzername).exists())

    def test_fa02_doppelte_benutzername_verhindert_kontoerstellung(self):
        """Zwei aufeinanderfolgende Anlageversuche speichern denselben Namen nur einmal."""
        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        konto2, fehlerliste2 = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNone(konto2)
        self.assertTrue(fehlerliste2)
        self.assertTrue(Benutzerkonto.objects.filter(username=self.benutzername).count() == 1)


class PasswortZuruecksetzenTests(TestCase):
    """Prüft Passwortzurücksetzen und den Bestandserhalt bei Fehlern."""

    def setUp(self):
        """Bereitet Kontoeingaben vor und lässt den Fall ohne Konto weiterhin zu."""
        self.benutzername = 'testnutzer'
        self.passwort = 'a' * 64
        self.rolle = Benutzerkonto.LEHRPERSON

    def test_fa02_passwort_zuruecksetzen(self):
        """Nach dem Zurücksetzen gilt nur das neue Passwort."""
        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        neues_passwort = 'b' * 64

        fehlerliste_zuruecksetzen = setze_passwort_zurueck(konto.pk, neues_passwort)
        self.assertFalse(fehlerliste_zuruecksetzen)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertTrue(konto.check_password(neues_passwort))
        self.assertFalse(konto.check_password(self.passwort))

    def test_qa01_zu_kurzes_passwort_laesst_altes_passwort_unveraendert(self):
        """Ein zu kurzes neues Passwort lässt den gespeicherten Hash unverändert."""
        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        zu_kurzes_passwort = 'c' * 14
        # Den kodierten Hash vor der abgewiesenen Änderung als Vergleichswert sichern.
        passwort_hash_vorher = konto.password

        fehlerliste_zuruecksetzen = setze_passwort_zurueck(konto.pk, zu_kurzes_passwort)
        self.assertTrue(fehlerliste_zuruecksetzen)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.password, passwort_hash_vorher)
        self.assertTrue(konto.check_password(self.passwort))
        self.assertFalse(konto.check_password(zu_kurzes_passwort))

    def test_fa02_nicht_vorhandenes_konto_wird_abgewiesen(self):
        """Eine nachweislich fehlende Konto-ID wird kontrolliert abgewiesen."""
        konto_id = 1
        self.assertFalse(Benutzerkonto.objects.filter(pk=konto_id).exists())
        neues_passwort = 'b' * 64

        fehlerliste_zuruecksetzen = setze_passwort_zurueck(konto_id, neues_passwort)
        self.assertTrue(fehlerliste_zuruecksetzen)
        self.assertFalse(Benutzerkonto.objects.filter(pk=konto_id).exists())


class AnmeldeTests(TestCase):
    """Prüft Fehlversuchsfenster, Sperren und deren Rücksetzung nach QA-02."""

    def setUp(self):
        """Bereitet gemeinsame Eingaben und eine noch sitzungslose Anfrage vor."""
        self.benutzername = 'testnutzer'
        self.rolle = Benutzerkonto.LEHRPERSON
        self.anfrage = HttpRequest()

    def test_qa02_aktive_sperre_verhindert_anmeldung(self):
        """Eine aktive Sperre weist auch das richtige Passwort ohne Verlängerung ab."""
        passwort = 'a' * 64
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        konto.gesperrt_bis = timezone.now() + timezone.timedelta(minutes=5)
        sperrende_vorher = konto.gesperrt_bis
        konto.save()

        # Die Sitzung für diesen direkten Aufruf wie bei einer Browseranfrage bereitstellen.
        self.anfrage.session = self.client.session

        fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort)
        self.assertTrue(fehlerliste_anmeldung)
        self.assertIn("Anmeldung fehlgeschlagen", fehlerliste_anmeldung[0])

        self.assertFalse(get_user(self.anfrage).is_authenticated)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.gesperrt_bis, sperrende_vorher)

    def test_qa02_sperrende_setzt_sperrzustand_zurueck(self):
        """Genau am Sperrende werden Zähler, Fensterbeginn und Sperre zurückgesetzt."""
        passwort = 'a' * 64
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 5
        konto.fehlversuchszeitraum_beginn = testzeitpunkt - timezone.timedelta(minutes=10)
        konto.gesperrt_bis = testzeitpunkt
        konto.save()

        # Die Sitzung für diesen direkten Aufruf wie bei einer Browseranfrage bereitstellen.
        self.anfrage.session = self.client.session

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertIsNone(konto.gesperrt_bis)
        self.assertEqual(konto.fehlversuche, 0)
        self.assertIsNone(konto.fehlversuchszeitraum_beginn)

    def test_qa02_fensterende_setzt_fehlversuche_zurueck(self):
        """Eine erfolgreiche Anmeldung am Fensterende setzt die alte Zählung zurück."""
        passwort = 'a' * 64
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 3
        konto.fehlversuchszeitraum_beginn = testzeitpunkt - timezone.timedelta(minutes=15)
        konto.save()

        # Die Sitzung für diesen direkten Aufruf wie bei einer Browseranfrage bereitstellen.
        self.anfrage.session = self.client.session

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 0)
        self.assertIsNone(konto.fehlversuchszeitraum_beginn)
        self.assertIsNone(konto.gesperrt_bis)

    def test_qa02_fuenfter_fehlversuch_sperrt_fuenf_minuten(self):
        """Der fünfte Fehlversuch im laufenden Fenster sperrt für fünf Minuten."""
        passwort = 'a' * 64
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        fensterbeginn_vorher = testzeitpunkt - timezone.timedelta(minutes=1)

        konto.fehlversuche = 4
        konto.fehlversuchszeitraum_beginn = fensterbeginn_vorher
        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, 'falsches_passwort')
        self.assertTrue(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche,5)
        self.assertEqual(konto.fehlversuchszeitraum_beginn,fensterbeginn_vorher)
        self.assertEqual(testzeitpunkt + timezone.timedelta(minutes=5),konto.gesperrt_bis)

    def test_qa02_vierter_fehlversuch_sperrt_nicht(self):
        """Der vierte Fehlversuch wird gezählt und löst noch keine Sperre aus."""
        passwort = 'a' * 64
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 3
        konto.fehlversuchszeitraum_beginn = testzeitpunkt - timezone.timedelta(minutes=1)
        konto.save()

        # Die Sitzung für diesen direkten Aufruf wie bei einer Browseranfrage bereitstellen.
        self.anfrage.session = self.client.session

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste = melde_an(self.anfrage, self.benutzername, passwort='falschesPW')
        self.assertTrue(fehlerliste)

        self.assertFalse(get_user(self.anfrage).is_authenticated)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche,4)
        self.assertEqual(konto.fehlversuchszeitraum_beginn, testzeitpunkt - timezone.timedelta(minutes=1))
        self.assertIsNone(konto.gesperrt_bis)

    def test_qa02_sechster_Versuch_verlaengert_sperre_nicht(self):
        """Weitere Versuche während einer Sperre ändern weder Zählung noch Sperrende."""
        passwort = 'a' * 64
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 5
        konto.fehlversuchszeitraum_beginn = testzeitpunkt - timezone.timedelta(minutes=2)
        sperrende_vorher = testzeitpunkt + timezone.timedelta(minutes=4)
        konto.gesperrt_bis = sperrende_vorher

        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste = melde_an(self.anfrage, self.benutzername, passwort='falschesPW')
        self.assertTrue(fehlerliste)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche,5)
        self.assertEqual(konto.gesperrt_bis, sperrende_vorher)
        self.assertEqual(konto.fehlversuchszeitraum_beginn, testzeitpunkt - timezone.timedelta(minutes=2))

    def test_qa02_fehlversuch_am_fensterende_beginnt_neues_fenster(self):
        """Ein Fehlversuch genau am Fensterende beginnt eine neue Zählung."""
        passwort = 'a' * 15
        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 4
        konto.fehlversuchszeitraum_beginn = (testzeitpunkt - timezone.timedelta(minutes=15))
        konto.gesperrt_bis = None
        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt,):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort='falschesPW')
        self.assertTrue(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 1)
        self.assertEqual(konto.fehlversuchszeitraum_beginn, testzeitpunkt)
        self.assertIsNone(konto.gesperrt_bis)

    def test_qa02_fehlversuch_kurz_vor_fensterende_sperrt(self):
        """Eine Sekunde vor Fensterende zählt der fünfte Fehlversuch noch zur Sperrschwelle."""
        passwort = 'a' * 15

        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        fensterbeginn_vorher = testzeitpunkt - timezone.timedelta(minutes=14, seconds=59)

        konto.fehlversuche = 4
        konto.fehlversuchszeitraum_beginn = fensterbeginn_vorher
        konto.gesperrt_bis = None
        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now',return_value=testzeitpunkt,):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort= 'falschesPW')

        self.assertTrue(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 5)
        self.assertEqual(konto.fehlversuchszeitraum_beginn,fensterbeginn_vorher,)
        self.assertEqual(konto.gesperrt_bis,testzeitpunkt + timezone.timedelta(minutes=5))

    def test_qa02_fehlversuch_nach_fensterende_beginnt_neues_fenster(self):
        """Eine Sekunde nach Fensterende zählt der Fehlversuch als erster eines neuen Fensters."""
        passwort = 'a' * 15

        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 4
        konto.fehlversuchszeitraum_beginn = (testzeitpunkt - timezone.timedelta(minutes=15, seconds=1))
        konto.gesperrt_bis = None
        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now',return_value=testzeitpunkt,):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort='falschesPW')

        self.assertTrue(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 1)
        self.assertEqual(konto.fehlversuchszeitraum_beginn, testzeitpunkt)
        self.assertIsNone(konto.gesperrt_bis)

    def test_qa02_kurz_vor_sperrende_bleibt_konto_gesperrt(self):
        """Eine Sekunde vor Sperrende bleibt auch das richtige Passwort gesperrt."""
        passwort = 'a' * 15

        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        fensterbeginn_vorher = testzeitpunkt - timezone.timedelta(minutes=10)
        sperrende_vorher = testzeitpunkt + timezone.timedelta(seconds=1)

        konto.fehlversuche = 5
        konto.fehlversuchszeitraum_beginn = fensterbeginn_vorher
        konto.gesperrt_bis = sperrende_vorher
        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now',return_value=testzeitpunkt,):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort)

        self.assertTrue(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 5)
        self.assertEqual(konto.fehlversuchszeitraum_beginn,fensterbeginn_vorher,)
        self.assertEqual(konto.gesperrt_bis, sperrende_vorher)

    def test_qa02_fehlversuch_nach_sperrende_beginnt_neues_fenster(self):
        """Nach Sperrende beginnt ein falsches Passwort eine neue Zählung."""
        passwort = 'a' * 15

        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 5
        konto.fehlversuchszeitraum_beginn = (testzeitpunkt - timezone.timedelta(minutes=10))
        konto.gesperrt_bis = testzeitpunkt - timezone.timedelta(seconds=1)
        konto.save()

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now',return_value=testzeitpunkt,):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort= 'falschesPW')

        self.assertTrue(fehlerliste_anmeldung)

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 1)
        self.assertEqual(konto.fehlversuchszeitraum_beginn, testzeitpunkt)
        self.assertIsNone(konto.gesperrt_bis)

    def test_qa02_richtiges_passwort_setzt_fehlversuche_zurueck(self):
        """Ein richtiges Passwort setzt Fehlversuche zurück und erzeugt gültige Sitzungswerte."""
        passwort = 'a' * 15

        konto, fehlerliste = lege_konto_an(self.benutzername, passwort, self.rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        testzeitpunkt = timezone.now()
        konto.fehlversuche = 4
        konto.fehlversuchszeitraum_beginn = (testzeitpunkt - timezone.timedelta(minutes=1))
        konto.gesperrt_bis = None
        konto.save()

        # Die Sitzung für diesen direkten Aufruf wie bei einer Browseranfrage bereitstellen.
        self.anfrage.session = self.client.session

        # Den Anmeldeversuch genau am vorbereiteten Prüfzeitpunkt ausführen.
        with patch('vp_stunden.sicherheit.timezone.now',return_value=testzeitpunkt,):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, passwort)

        self.assertFalse(fehlerliste_anmeldung)

        angemeldetes_konto = get_user(self.anfrage)
        self.assertTrue(angemeldetes_konto.is_authenticated)
        self.assertEqual(angemeldetes_konto.pk, konto.pk)

        self.anfrage.session.save()
        # Eine neue SessionStore-Instanz liest die tatsächlich gespeicherten Sitzungswerte.
        gespeicherte_sitzung = SessionStore(session_key=self.anfrage.session.session_key)

        self.assertEqual(gespeicherte_sitzung['sitzungsbeginn'], testzeitpunkt.timestamp())
        self.assertEqual(gespeicherte_sitzung['letzte_aktivitaet'], testzeitpunkt.timestamp())

        # Den gespeicherten Kontozustand statt nur das geänderte Python-Objekt prüfen.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 0)
        self.assertIsNone(konto.fehlversuchszeitraum_beginn)
        self.assertIsNone(konto.gesperrt_bis)


class AbmeldeTests(TestCase):
    """Prüft die Ungültigkeit einer Sitzung nach der Abmeldung."""

    def test_qa03_abmeldung_macht_sitzung_ungueltig(self):
        """Die Abmeldung entfernt Authentifizierung und gespeicherte Sitzungskennung."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.LEHRPERSON
        konto, fehlerliste = lege_konto_an(benutzername, passwort, rolle)

        anfrage = HttpRequest()
        anfrage.session = self.client.session
        fehlerliste = melde_an(anfrage, benutzername, passwort=passwort)
        self.assertFalse(fehlerliste)
        self.assertTrue(get_user(anfrage).is_authenticated)

        anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = anfrage.session.session_key
        self.assertTrue(anfrage.session.exists(alte_sitzungskennung))

        melde_ab(anfrage)
        self.assertFalse(get_user(anfrage).is_authenticated)
        self.assertFalse(anfrage.session.exists(alte_sitzungskennung))


class ZugriffsTests(TestCase):
    """Prüft Rollenfreigabe, Sitzungsangaben und beide Sitzungsfristen."""

    def setUp(self):
        """Bereitet Eingaben und eine unangemeldete Anfrage mit Sitzung vor."""
        self.benutzername = 'testnutzer'
        self.passwort = 'a' * 64
        # Direkte Funktionsaufrufe umgehen die Middleware und brauchen diese Sitzung.
        self.anfrage = HttpRequest()
        self.anfrage.session = self.client.session

    def test_qa04_ohne_anmeldung_wird_zugriff_verweigert(self):
        """Eine unangemeldete Anfrage erhält keinen geschützten Zugriff."""
        fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.LEHRPERSON])
        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)

    def test_qa04_erlaubte_rolle_erhaelt_zugriff(self):
        """Die erlaubte Rolle erhält Zugriff und aktualisiert nur die letzte Aktivität."""
        rolle = Benutzerkonto.LEHRPERSON

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=5)

        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.LEHRPERSON])

        self.assertFalse(fehlerliste)
        self.assertTrue(get_user(self.anfrage).is_authenticated)

        self.anfrage.session.save()
        # Eine neue SessionStore-Instanz liest die tatsächlich gespeicherten Sitzungswerte.
        gespeicherte_sitzung = SessionStore(session_key=self.anfrage.session.session_key)
        self.assertEqual(gespeicherte_sitzung['sitzungsbeginn'], testzeitpunkt.timestamp())
        self.assertEqual(gespeicherte_sitzung['letzte_aktivitaet'], zugriffszeitpunkt.timestamp())

    def test_qa04_falsche_rolle_verweigert_zugriff(self):
        """Die falsche Rolle erhält keinen Zugriff und erneuert die Aktivitätszeit nicht."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=5)

        konto, fehlerliste_anmeldung = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.LEHRPERSON])

        self.assertTrue(fehlerliste)
        self.assertTrue(get_user(self.anfrage).is_authenticated)

        self.anfrage.session.save()
        # Eine neue SessionStore-Instanz liest die tatsächlich gespeicherten Sitzungswerte.
        gespeicherte_sitzung = SessionStore(
            session_key=self.anfrage.session.session_key
        )

        self.assertEqual(gespeicherte_sitzung['sitzungsbeginn'], testzeitpunkt.timestamp())
        self.assertEqual(gespeicherte_sitzung['letzte_aktivitaet'], testzeitpunkt.timestamp())

    def test_qa03_fehlender_sitzungsbeginn_beendet_sitzung(self):
        """Ein fehlender Sitzungsbeginn führt zur vollständigen Abmeldung."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=5)

        konto, fehlerliste_anmeldung = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Eine unvollständige Sitzung simulieren, die sicher beendet werden muss.
        del self.anfrage.session['sitzungsbeginn']
        self.anfrage.session.save()

        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)
        self.assertFalse(self.anfrage.session.exists(alte_sitzungskennung))

    def test_qa03_fehlende_letzte_aktivitaet_beendet_sitzung(self):
        """Eine fehlende letzte Aktivität führt zur vollständigen Abmeldung."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=5)

        konto, fehlerliste_anmeldung = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Eine unvollständige Sitzung simulieren, die sicher beendet werden muss.
        del self.anfrage.session['letzte_aktivitaet']
        self.anfrage.session.save()

        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)
        self.assertFalse(self.anfrage.session.exists(alte_sitzungskennung))

    def test_qa03_genau_30_minuten_inaktivitaet_beendet_sitzung(self):
        """Genau 30 Minuten Inaktivität beenden die Sitzung."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=30)

        konto, fehlerliste_anmeldung = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        self.anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)
        self.assertFalse(self.anfrage.session.exists(alte_sitzungskennung))

    def test_qa03_genau_29_minuten_59_sekunden_inaktivitaet_erlaubt_zugriff(self):
        """Eine Sekunde vor der Inaktivitätsgrenze bleibt die Sitzung nutzbar."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=29, seconds=59)

        konto, fehlerliste_anmeldung = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        self.anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertFalse(fehlerliste)
        self.assertTrue(get_user(self.anfrage).is_authenticated)
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        self.anfrage.session.save()
        # Eine neue SessionStore-Instanz liest die tatsächlich gespeicherten Sitzungswerte.
        gespeicherte_sitzung = SessionStore(
            session_key=self.anfrage.session.session_key
        )

        self.assertEqual(gespeicherte_sitzung['sitzungsbeginn'], testzeitpunkt.timestamp())
        self.assertEqual(gespeicherte_sitzung['letzte_aktivitaet'], zugriffszeitpunkt.timestamp())

    def test_qa03_genau_30_minuten_1_sekunde_inaktivitaet_beendet_sitzung(self):
        """Eine Sekunde nach der Inaktivitätsgrenze ist die Sitzung ungültig."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(minutes=30, seconds=1)

        konto, fehlerliste_anmeldung = lege_konto_an(self.benutzername, self.passwort, rolle)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        self.anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)
        self.assertFalse(self.anfrage.session.exists(alte_sitzungskennung))

    def test_qa03_genau_3_stunden_59_minuten_59_sekunden_sitzungsdauer_erlaubt_zugriff(self):
        """Eine Sekunde vor der absoluten Frist bleibt eine kürzlich aktive Sitzung nutzbar."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(hours=3, minutes=59, seconds=59)

        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Frische Aktivität schließt Inaktivität als Ursache für das Sitzungsende aus.
        aktivitaetszeitpunkt = zugriffszeitpunkt - timezone.timedelta(minutes=5)
        self.anfrage.session['letzte_aktivitaet'] = aktivitaetszeitpunkt.timestamp()

        self.anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertFalse(fehlerliste)
        self.assertTrue(get_user(self.anfrage).is_authenticated)
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        self.anfrage.session.save()
        # Eine neue SessionStore-Instanz liest die tatsächlich gespeicherten Sitzungswerte.
        gespeicherte_sitzung = SessionStore(session_key=self.anfrage.session.session_key)
        self.assertEqual(gespeicherte_sitzung['sitzungsbeginn'],testzeitpunkt.timestamp())
        self.assertEqual(gespeicherte_sitzung['letzte_aktivitaet'],zugriffszeitpunkt.timestamp())

    def test_qa03_genau_4_stunden_sitzungsdauer_beendet_sitzung(self):
        """Genau vier Stunden nach Anmeldung endet auch eine kürzlich aktive Sitzung."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(hours=4)

        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Frische Aktivität schließt Inaktivität als Ursache für das Sitzungsende aus.
        aktivitaetszeitpunkt = zugriffszeitpunkt - timezone.timedelta(minutes=5)
        self.anfrage.session['letzte_aktivitaet'] = aktivitaetszeitpunkt.timestamp()

        self.anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)
        self.assertFalse(self.anfrage.session.exists(alte_sitzungskennung))

    def test_qa03_genau_4_stunden_1_sekunde_sitzungsdauer_beendet_sitzung(self):
        """Eine Sekunde nach der absoluten Frist bleibt die Sitzung ungültig."""
        rolle = Benutzerkonto.ADMINISTRATOR

        testzeitpunkt = timezone.now()
        zugriffszeitpunkt = testzeitpunkt + timezone.timedelta(hours=4, seconds=1)

        konto, fehlerliste = lege_konto_an(self.benutzername, self.passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Einen festen Sitzungsbeginn für die spätere Fristprüfung setzen.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=testzeitpunkt):
            fehlerliste_anmeldung = melde_an(self.anfrage, self.benutzername, self.passwort)
        self.assertFalse(fehlerliste_anmeldung)

        # Frische Aktivität schließt Inaktivität als Ursache für das Sitzungsende aus.
        aktivitaetszeitpunkt = zugriffszeitpunkt - timezone.timedelta(minutes=5)
        self.anfrage.session['letzte_aktivitaet'] = aktivitaetszeitpunkt.timestamp()

        self.anfrage.session.save()
        # Die bisherige Kennung für die spätere Existenzprüfung festhalten.
        alte_sitzungskennung = self.anfrage.session.session_key
        self.assertTrue(self.anfrage.session.exists(alte_sitzungskennung))

        # Den späteren Zugriff gezielt simulieren, ohne im Test zu warten.
        with patch('vp_stunden.sicherheit.timezone.now', return_value=zugriffszeitpunkt):
            fehlerliste = pruefe_zugriff(self.anfrage, [Benutzerkonto.ADMINISTRATOR])

        self.assertTrue(fehlerliste)
        self.assertFalse(get_user(self.anfrage).is_authenticated)
        self.assertFalse(self.anfrage.session.exists(alte_sitzungskennung))


class AnmeldeViewTests(TestCase):
    """Prüft HTTP-Zugriff, Anmeldung und Sitzungszuordnung der Anmelde-View."""

    def test_fa01_get_anfrage_wird_abgewiesen(self):
        """Eine GET-Anfrage wird mit POST als einzig erlaubter Methode abgewiesen."""
        antwort = self.client.get('/anmelden/')

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort['Allow'], 'POST')
        self.assertFalse(antwort.json()['erfolg'])

    def test_fa01_post_ohne_anmeldedaten_wird_abgewiesen(self):
        """Fehlende Anmeldedaten werden vor dem Aufruf der Anmeldefunktion abgewiesen."""
        # Der Ersatz erfasst Aufrufe: Fehlende Felder müssen bereits die View stoppen.
        with patch('vp_stunden.views.melde_an') as anmeldefunktion:
            antwort = self.client.post('/anmelden/', data={})

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()['erfolg'])
        anmeldefunktion.assert_not_called()

    def test_fa01_erfolgreiche_anmeldung(self):
        """Gültige Anmeldedaten ordnen die Sitzung dem angemeldeten Konto zu."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.LEHRPERSON
        konto, fehlerliste = lege_konto_an(benutzername, passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        antwort = self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': passwort
            }
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])

        # Django speichert die Konto-ID unter SESSION_KEY als Text in der Sitzung.
        self.assertEqual(self.client.session[SESSION_KEY], str(konto.pk))

    def test_fa01_falsches_passwort_verhindert_anmeldung(self):
        """Ein falsches Passwort verhindert die Anmeldung und erhöht den Fehlversuchszähler."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.LEHRPERSON
        konto, fehlerliste = lege_konto_an(benutzername, passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        antwort = self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': 'falsches_passwort'
            }
        )

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()['erfolg'])
        self.assertNotIn(SESSION_KEY, self.client.session)

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertEqual(konto.fehlversuche, 1)


class AbmeldeViewTests(TestCase):
    """Prüft die Sitzungsbeendigung und die Beschränkung der Abmeldung auf POST."""

    def setUp(self):
        """Bereitet für jeden Test eine tatsächlich angemeldete Sitzung vor."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.LEHRPERSON
        self.konto, fehlerliste = lege_konto_an(benutzername, passwort, rolle)
        self.assertIsNotNone(self.konto)
        self.assertFalse(fehlerliste)

        self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': passwort
            }
        )

    def test_qa03_abmeldung_beendet_gespeicherte_sitzung(self):
        """POST beendet die Anmeldung und löscht die bisher gespeicherte Sitzung."""
        # Die bisherige Kennung erlaubt anschließend die Prüfung der gespeicherten Sitzung.
        sitzung_vorher = self.client.session
        alte_sitzungskennung = sitzung_vorher.session_key

        self.assertTrue(sitzung_vorher.exists(alte_sitzungskennung))
        self.assertEqual(sitzung_vorher[SESSION_KEY], str(self.konto.pk))

        antwort = self.client.post(
            '/abmelden/',
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])
        self.assertNotIn(SESSION_KEY, self.client.session)
        self.assertFalse(sitzung_vorher.exists(alte_sitzungskennung))

    def test_qa03_get_anfrage_beendet_sitzung_nicht(self):
        """GET wird abgewiesen und erhält die angemeldete Sitzung."""
        # Die bisherige Kennung erlaubt anschließend die Prüfung der gespeicherten Sitzung.
        sitzung_vorher = self.client.session
        alte_sitzungskennung = sitzung_vorher.session_key

        self.assertTrue(sitzung_vorher.exists(alte_sitzungskennung))
        self.assertEqual(sitzung_vorher[SESSION_KEY], str(self.konto.pk))

        antwort = self.client.get(
            '/abmelden/',
        )

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort['Allow'], 'POST')
        self.assertFalse(antwort.json()['erfolg'])
        self.assertTrue(sitzung_vorher.exists(alte_sitzungskennung))
        self.assertEqual(self.client.session[SESSION_KEY], str(self.konto.pk))


class PasswortAenderungTests(TestCase):
    """Prüft eigene Passwortänderungen und den Erhalt der bestehenden Sitzung."""

    def setUp(self):
        """Legt für jeden Test dasselbe Konto ohne vorherige Anmeldung an."""
        self.benutzername = 'testnutzer'
        self.altes_passwort = 'a' * 64
        rolle = Benutzerkonto.LEHRPERSON
        self.konto, fehlerliste = lege_konto_an(
            self.benutzername, self.altes_passwort, rolle
        )
        self.assertIsNotNone(self.konto)
        self.assertFalse(fehlerliste)

    def test_fa03_erfolgreiche_eigene_passwortaenderung(self):
        """Ein neues Passwort wird gespeichert und die bestehende Anmeldung bleibt erhalten."""
        neues_passwort = 'b' * 64

        # Beim direkten Funktionsaufruf muss der Test die Sitzung selbst bereitstellen.
        anfrage = HttpRequest()
        anfrage.session = self.client.session

        fehlerliste_anmeldung = melde_an(anfrage, self.benutzername, self.altes_passwort)
        self.assertFalse(fehlerliste_anmeldung)

        anfrage.session.save()
        sitzungsbeginn_vorher = anfrage.session['sitzungsbeginn']

        # Die spätere Testzeit macht einen unzulässig erneuerten Sitzungsbeginn erkennbar.
        aenderungszeitpunkt = timezone.now() + timezone.timedelta(minutes=5)
        with patch('vp_stunden.sicherheit.timezone.now', return_value=aenderungszeitpunkt):
            fehlerliste = aendere_passwort(anfrage, self.altes_passwort, neues_passwort)

        self.assertFalse(fehlerliste)
        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        self.konto.refresh_from_db()
        self.assertTrue(self.konto.check_password(neues_passwort))
        self.assertFalse(self.konto.check_password(self.altes_passwort))

        anfrage.session.save()
        # Neu laden, damit die Anmeldung anhand der tatsächlich gespeicherten Sitzung geprüft wird.
        anfrage.session = SessionStore(
            session_key=anfrage.session.session_key
        )

        angemeldetes_konto = get_user(anfrage)
        self.assertTrue(angemeldetes_konto.is_authenticated)
        self.assertEqual(angemeldetes_konto.pk, self.konto.pk)
        self.assertEqual(
            anfrage.session['sitzungsbeginn'], sitzungsbeginn_vorher
        )

    def test_fa03_falsches_bisheriges_passwort_verhindert_aenderung(self):
        """Ein falsches bisheriges Passwort wird bei erhaltener Anmeldung abgewiesen."""
        neues_passwort = 'b' * 64

        # Beim direkten Funktionsaufruf muss der Test die Sitzung selbst bereitstellen.
        anfrage = HttpRequest()
        anfrage.session = self.client.session

        fehlerliste_anmeldung = melde_an(anfrage, self.benutzername, self.altes_passwort)
        self.assertFalse(fehlerliste_anmeldung)

        anfrage.session.save()
        sitzungsbeginn_vorher = anfrage.session['sitzungsbeginn']

        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = self.konto.password

        # Die spätere Testzeit macht einen unzulässig erneuerten Sitzungsbeginn erkennbar.
        aenderungszeitpunkt = timezone.now() + timezone.timedelta(minutes=5)
        with patch('vp_stunden.sicherheit.timezone.now', return_value=aenderungszeitpunkt):
            fehlerliste = aendere_passwort(anfrage, 'falsches_passwort', neues_passwort)

        self.assertTrue(fehlerliste)

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        self.konto.refresh_from_db()
        self.assertEqual(self.konto.password, passwort_hash_vorher)
        self.assertTrue(self.konto.check_password(self.altes_passwort))
        self.assertFalse(self.konto.check_password(neues_passwort))

        anfrage.session.save()
        # Neu laden, damit die Anmeldung anhand der tatsächlich gespeicherten Sitzung geprüft wird.
        anfrage.session = SessionStore(
            session_key=anfrage.session.session_key
        )

        angemeldetes_konto = get_user(anfrage)
        self.assertTrue(angemeldetes_konto.is_authenticated)
        self.assertEqual(angemeldetes_konto.pk, self.konto.pk)
        self.assertEqual(
            anfrage.session['sitzungsbeginn'], sitzungsbeginn_vorher
        )

    def test_qa01_zu_kurzes_neues_passwort_verhindert_aenderung(self):
        """Ein zu kurzes neues Passwort wird bei erhaltener Anmeldung abgewiesen."""
        neues_passwort = 'b' * 14

        # Beim direkten Funktionsaufruf muss der Test die Sitzung selbst bereitstellen.
        anfrage = HttpRequest()
        anfrage.session = self.client.session

        fehlerliste_anmeldung = melde_an(anfrage, self.benutzername, self.altes_passwort)
        self.assertFalse(fehlerliste_anmeldung)

        anfrage.session.save()
        sitzungsbeginn_vorher = anfrage.session['sitzungsbeginn']

        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = self.konto.password

        # Die spätere Testzeit macht einen unzulässig erneuerten Sitzungsbeginn erkennbar.
        aenderungszeitpunkt = timezone.now() + timezone.timedelta(minutes=5)
        with patch('vp_stunden.sicherheit.timezone.now', return_value=aenderungszeitpunkt):
            fehlerliste = aendere_passwort(anfrage, self.altes_passwort, neues_passwort)

        self.assertTrue(fehlerliste)

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        self.konto.refresh_from_db()
        self.assertEqual(self.konto.password, passwort_hash_vorher)
        self.assertTrue(self.konto.check_password(self.altes_passwort))
        self.assertFalse(self.konto.check_password(neues_passwort))

        anfrage.session.save()
        # Neu laden, damit die Anmeldung anhand der tatsächlich gespeicherten Sitzung geprüft wird.
        anfrage.session = SessionStore(
            session_key=anfrage.session.session_key
        )

        angemeldetes_konto = get_user(anfrage)
        self.assertTrue(angemeldetes_konto.is_authenticated)
        self.assertEqual(angemeldetes_konto.pk, self.konto.pk)
        self.assertEqual(
            anfrage.session['sitzungsbeginn'], sitzungsbeginn_vorher
        )

    def test_fa03_ohne_anmeldung_wird_passwortaenderung_verweigert(self):
        """Ohne Anmeldung werden weder das Passwort geändert noch eine Anmeldung hergestellt."""
        neues_passwort = 'b' * 64

        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = self.konto.password

        # Beim direkten Funktionsaufruf muss der Test die Sitzung selbst bereitstellen.
        anfrage = HttpRequest()
        anfrage.session = self.client.session
        self.assertFalse(get_user(anfrage).is_authenticated)

        fehlerliste = aendere_passwort(
            anfrage, self.altes_passwort, neues_passwort
        )

        self.assertTrue(fehlerliste)

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        self.konto.refresh_from_db()
        self.assertEqual(self.konto.password, passwort_hash_vorher)
        self.assertTrue(self.konto.check_password(self.altes_passwort))
        self.assertFalse(self.konto.check_password(neues_passwort))
        self.assertFalse(get_user(anfrage).is_authenticated)


class PasswortAenderungViewTests(TestCase):
    """Prüft Formularprüfung und Zugriffsschutz der Passwortänderungs-View."""

    def test_fa03_post_anfrage_aendert_passwort(self):
        """Vollständige gültige Formulardaten ändern das Passwort des angemeldeten Kontos."""
        benutzername = 'testnutzer'
        altes_passwort = 'a' * 64
        neues_passwort = 'b' * 64
        rolle = Benutzerkonto.LEHRPERSON

        konto, fehlerliste = lege_konto_an(benutzername, altes_passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        antwort = self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': altes_passwort
            }
        )
        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])

        antwort = self.client.post(
            '/passwort_aendern/',
            data={
                'bisheriges_passwort': altes_passwort,
                'neues_passwort': neues_passwort,
                'wiederholung_neues_passwort': neues_passwort
            }
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertTrue(konto.check_password(neues_passwort))

    def test_fa03_abweichende_passwortwiederholung_verhindert_aenderung(self):
        """Eine abweichende Passwortwiederholung lässt das gespeicherte Passwort unverändert."""
        benutzername = 'testnutzer'
        altes_passwort = 'a' * 64
        neues_passwort = 'b' * 64
        rolle = Benutzerkonto.LEHRPERSON

        konto, fehlerliste = lege_konto_an(
            benutzername, altes_passwort, rolle
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        antwort = self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': altes_passwort
            }
        )
        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])

        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = konto.password

        antwort = self.client.post(
            '/passwort_aendern/',
            data={
                'bisheriges_passwort': altes_passwort,
                'neues_passwort': neues_passwort,
                'wiederholung_neues_passwort': neues_passwort + 'x'
            }
        )

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertEqual(konto.password, passwort_hash_vorher)
        self.assertTrue(konto.check_password(altes_passwort))
        self.assertFalse(konto.check_password(neues_passwort))

    def test_fa03_post_ohne_anmeldung_wird_abgewiesen(self):
        """Eine Passwortänderung ohne Anmeldung wird ohne Datenänderung abgewiesen."""
        benutzername = 'testnutzer'
        altes_passwort = 'a' * 64
        neues_passwort = 'b' * 64
        rolle = Benutzerkonto.LEHRPERSON

        konto, fehlerliste = lege_konto_an(
            benutzername, altes_passwort, rolle
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = konto.password

        antwort = self.client.post(
            '/passwort_aendern/',
            data={
                'bisheriges_passwort': altes_passwort,
                'neues_passwort': neues_passwort,
                'wiederholung_neues_passwort': neues_passwort
            }
        )

        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertEqual(konto.password, passwort_hash_vorher)
        self.assertTrue(konto.check_password(altes_passwort))
        self.assertFalse(konto.check_password(neues_passwort))

    def test_fa03_get_anfrage_wird_abgewiesen(self):
        """Die Passwortänderungs-View weist GET ab und erlaubt ausschließlich POST."""
        antwort = self.client.get('/passwort_aendern/')

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort['Allow'], 'POST')
        self.assertFalse(antwort.json()['erfolg'])


class KontoErstellungViewTests(TestCase):
    """Prüft Kontenanlage und Rollenprüfung über die Kontoerstellungs-View."""

    def test_fa02_administrator_kann_konto_anlegen(self):
        """Ein Administratorkonto kann ein Konto mit den angefragten Daten anlegen."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.LEHRPERSON

        admin_konto, fehlerliste = lege_konto_an('admin', 'adminpasswort12', Benutzerkonto.ADMINISTRATOR)
        self.assertIsNotNone(admin_konto)
        self.assertFalse(fehlerliste)

        antwort_anmeldung = self.client.post(
            '/anmelden/',
            data={
                'benutzername': 'admin',
                'passwort': 'adminpasswort12'
            }
        )

        self.assertEqual(antwort_anmeldung.status_code, 200)
        self.assertTrue(antwort_anmeldung.json()['erfolg'])
        self.assertFalse(Benutzerkonto.objects.filter(username=benutzername).exists())

        antwort = self.client.post(
            '/konto_anlegen/',
            data={
                'benutzername': benutzername,
                'passwort': passwort,
                'rolle': rolle
            }
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])

        neues_konto = Benutzerkonto.objects.get(username=benutzername)

        self.assertEqual(neues_konto.username, benutzername)
        self.assertEqual(neues_konto.rolle, rolle)
        self.assertTrue(neues_konto.check_password(passwort))

    def test_qa04_lehrperson_kann_kein_konto_anlegen(self):
        """Eine Lehrperson darf auch mit vollständigen Daten kein Konto anlegen."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.ADMINISTRATOR

        lehrperson_konto, fehlerliste = lege_konto_an('lehrperson', 'lehrpersonpasswort12', Benutzerkonto.LEHRPERSON)
        self.assertIsNotNone(lehrperson_konto)
        self.assertFalse(fehlerliste)

        antwort_anmeldung = self.client.post(
            '/anmelden/',
            data={
                'benutzername': 'lehrperson',
                'passwort': 'lehrpersonpasswort12'
            }
        )

        self.assertEqual(antwort_anmeldung.status_code, 200)
        self.assertTrue(antwort_anmeldung.json()['erfolg'])
        self.assertFalse(Benutzerkonto.objects.filter(username=benutzername).exists())

        kontozahl_vorher = Benutzerkonto.objects.count()

        antwort = self.client.post(
            '/konto_anlegen/',
            data={
                'benutzername': benutzername,
                'passwort': passwort,
                'rolle': rolle
            }
        )

        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()['erfolg'])
        self.assertFalse(
            Benutzerkonto.objects.filter(username=benutzername).exists()
        )
        self.assertEqual(
            Benutzerkonto.objects.count(), kontozahl_vorher
        )

    def test_qa04_kontoanlage_ohne_anmeldung_wird_abgewiesen(self):
        """Ohne Anmeldung bleibt die Kontenanlage ohne Datenänderung verweigert."""
        benutzername = 'testnutzer'
        passwort = 'a' * 64
        rolle = Benutzerkonto.ADMINISTRATOR

        self.assertFalse(
            Benutzerkonto.objects.filter(username=benutzername).exists()
        )

        kontozahl_vorher = Benutzerkonto.objects.count()

        antwort = self.client.post(
            '/konto_anlegen/',
            data={
                'benutzername': benutzername,
                'passwort': passwort,
                'rolle': rolle
            }
        )

        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()['erfolg'])
        self.assertFalse(
            Benutzerkonto.objects.filter(username=benutzername).exists()
        )

        self.assertEqual(
            Benutzerkonto.objects.count(), kontozahl_vorher
        )

    def test_fa02_get_anfrage_wird_abgewiesen(self):
        """Die View weist GET ab und nennt POST als einzig erlaubte Methode."""
        antwort = self.client.get('/konto_anlegen/')

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort['Allow'], 'POST')
        self.assertFalse(antwort.json()['erfolg'])


class PasswortZuruecksetzenViewTests(TestCase):
    """Prüft administrative Passwortrücksetzung und abgewiesene Zugriffe."""

    def test_fa02_administrator_kann_passwort_zuruecksetzen(self):
        """Ein Administratorkonto kann das Passwort eines separaten Zielkontos zurücksetzen."""
        admin_benutzername = 'testadmin'
        admin_passwort = 'a' * 64
        admin_konto, fehlerliste = lege_konto_an(
            admin_benutzername, admin_passwort, Benutzerkonto.ADMINISTRATOR
        )
        self.assertIsNotNone(admin_konto)
        self.assertFalse(fehlerliste)

        passwort = 'b' * 64
        neues_passwort = 'c' * 64
        konto, fehlerliste = lege_konto_an(
            'testnutzer', passwort, Benutzerkonto.LEHRPERSON
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        antwort_anmeldung = self.client.post(
            '/anmelden/',
            data={
                'benutzername': admin_benutzername,
                'passwort': admin_passwort
            }
        )
        self.assertEqual(antwort_anmeldung.status_code, 200)
        self.assertTrue(antwort_anmeldung.json()['erfolg'])

        antwort = self.client.post(
            '/passwort_zuruecksetzen/',
            data={
                'konto_id': str(konto.pk),
                'neues_passwort': neues_passwort
            }
        )
        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertTrue(konto.check_password(neues_passwort))
        self.assertFalse(konto.check_password(passwort))

    def test_qa04_lehrperson_kann_kein_passwort_zuruecksetzen(self):
        """Eine Lehrperson kann das Passwort eines fremden Administratorkontos nicht ändern."""
        benutzername = 'testlehrperson'
        passwort = 'a' * 64
        lehrperson_konto, fehlerliste = lege_konto_an(
            benutzername, passwort, Benutzerkonto.LEHRPERSON
        )
        self.assertIsNotNone(lehrperson_konto)
        self.assertFalse(fehlerliste)

        zielpasswort = 'b' * 64
        neues_passwort = 'c' * 64
        konto, fehlerliste = lege_konto_an(
            'testadmin', zielpasswort, Benutzerkonto.ADMINISTRATOR
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)
        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = konto.password

        antwort_anmeldung = self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': passwort
            }
        )
        self.assertEqual(antwort_anmeldung.status_code, 200)
        self.assertTrue(antwort_anmeldung.json()['erfolg'])

        antwort = self.client.post(
            '/passwort_zuruecksetzen/',
            data={
                'konto_id': str(konto.pk),
                'neues_passwort': neues_passwort
            }
        )
        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertEqual(konto.password, passwort_hash_vorher)
        self.assertTrue(konto.check_password(zielpasswort))
        self.assertFalse(konto.check_password(neues_passwort))

    def test_qa04_ruecksetzung_ohne_anmeldung_wird_abgewiesen(self):
        """Ohne Anmeldung bleibt die Passwortrücksetzung ohne Datenänderung verweigert."""
        passwort = 'a' * 64
        neues_passwort = 'b' * 64
        konto, fehlerliste = lege_konto_an(
            'testnutzer', passwort, Benutzerkonto.LEHRPERSON
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)
        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = konto.password

        antwort = self.client.post(
            '/passwort_zuruecksetzen/',
            data={
                'konto_id': str(konto.pk),
                'neues_passwort': neues_passwort
            }
        )
        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertEqual(konto.password, passwort_hash_vorher)
        self.assertTrue(konto.check_password(passwort))
        self.assertFalse(konto.check_password(neues_passwort))

    def test_fa02_get_anfrage_wird_abgewiesen(self):
        """Die View weist GET ab und nennt POST als einzig erlaubte Methode."""
        antwort = self.client.get('/passwort_zuruecksetzen/')

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort['Allow'], 'POST')
        self.assertFalse(antwort.json()['erfolg'])

    def test_fa02_ungueltige_konto_id_wird_abgewiesen(self):
        """Eine nicht numerische Konto-ID wird vor dem Aufruf der Rücksetzfunktion abgewiesen."""
        benutzername = 'testadmin'
        passwort = 'a' * 64
        admin_konto, fehlerliste = lege_konto_an(
            benutzername, passwort, Benutzerkonto.ADMINISTRATOR
        )
        self.assertIsNotNone(admin_konto)
        self.assertFalse(fehlerliste)

        antwort_anmeldung = self.client.post(
            '/anmelden/',
            data={
                'benutzername': benutzername,
                'passwort': passwort
            }
        )
        self.assertEqual(antwort_anmeldung.status_code, 200)
        self.assertTrue(antwort_anmeldung.json()['erfolg'])

        # Die View muss die ID prüfen, bevor sie die Rücksetzfunktion aufruft.
        with patch('vp_stunden.views.setze_passwort_zurueck') as ruecksetzfunktion:
            antwort = self.client.post(
                '/passwort_zuruecksetzen/',
                data={
                    'konto_id': 'keine_zahl',
                    'neues_passwort': 'b' * 64
                }
            )

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()['erfolg'])

        ruecksetzfunktion.assert_not_called()

    def test_qa01_zu_kurzes_passwort_verhindert_ruecksetzung(self):
        """Ein zu kurzes neues Passwort lässt das gespeicherte Zielpasswort unverändert."""
        admin_benutzername = 'testadmin'
        admin_passwort = 'a' * 64
        admin_konto, fehlerliste = lege_konto_an(
            admin_benutzername, admin_passwort, Benutzerkonto.ADMINISTRATOR
        )
        self.assertIsNotNone(admin_konto)
        self.assertFalse(fehlerliste)

        passwort = 'b' * 64
        neues_passwort = 'c' * 14
        konto, fehlerliste = lege_konto_an(
            'testnutzer', passwort, Benutzerkonto.LEHRPERSON
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)
        # Auch der gespeicherte Hash muss bei einer Abweisung unverändert bleiben.
        passwort_hash_vorher = konto.password

        antwort_anmeldung = self.client.post(
            '/anmelden/',
            data={
                'benutzername': admin_benutzername,
                'passwort': admin_passwort
            }
        )
        self.assertEqual(antwort_anmeldung.status_code, 200)
        self.assertTrue(antwort_anmeldung.json()['erfolg'])

        antwort = self.client.post(
            '/passwort_zuruecksetzen/',
            data={
                'konto_id': str(konto.pk),
                'neues_passwort': neues_passwort
            }
        )
        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()['erfolg'])

        # Für die folgenden Prüfungen den tatsächlich gespeicherten Zustand laden.
        konto.refresh_from_db()
        self.assertEqual(konto.password, passwort_hash_vorher)
        self.assertTrue(konto.check_password(passwort))
        self.assertFalse(konto.check_password(neues_passwort))


class VPDatenmodellTests(TestCase):
    """Prüft die Eindeutigkeitsregeln für Untersuchungen und Teilnahmen."""

    def test_fa08_vergleichswert_ist_eindeutig(self):
        """Die Datenbank verhindert zwei Untersuchungen mit demselben Studienvergleichswert."""
        anzeigename1 = 'Teststudie A'
        vergleichswert1 = 'teststudie a'

        studie1 = Untersuchung.objects.create(
            studienname=anzeigename1,
            studienname_vergleich=vergleichswert1,
            stundenwert=1.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson='testdozentA'
        )

        anzeigename2 = 'TESTSTUDIE A '
        vergleichswert2 = 'teststudie a'

        # Der innere Sicherungspunkt wird vor dem Abfangen des Fehlers zurückgerollt.
        # Dadurch bleiben die anschließenden Datenbankabfragen im Test möglich.
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Untersuchung.objects.create(
                    studienname=anzeigename2,
                    studienname_vergleich=vergleichswert2,
                    stundenwert=1.5,
                    einordnung=Untersuchung.PRAESENZ,
                    lehrperson='testdozentB',
                )

        self.assertEqual(Untersuchung.objects.count(), 1)

    def test_fa12_teilnahmezuordnung_ist_eindeutig(self):
        """Nur die Kombination aus Untersuchung und Matrikelnummer muss eindeutig sein."""
        studie1 = Untersuchung.objects.create(
            studienname='Studie 1',
            studienname_vergleich='studie 1',
            stundenwert=1.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson='dozent1',
        )

        studie2 = Untersuchung.objects.create(
            studienname='Studie 2',
            studienname_vergleich='studie 2',
            stundenwert=2.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson='dozent2',
        )

        teilnahme1 = Teilnahme.objects.create(
            untersuchung=studie1,
            matrikelnummer='1234567',
        )

        teilnahme2 = Teilnahme.objects.create(
            untersuchung=studie2,
            matrikelnummer='1234567',
        )

        teilnahme3 = Teilnahme.objects.create(
            untersuchung=studie1,
            matrikelnummer='7654321',
        )

        # Der innere Sicherungspunkt wird vor dem Abfangen des Fehlers zurückgerollt.
        # Dadurch bleiben die anschließenden Datenbankabfragen im Test möglich.
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Teilnahme.objects.create(
                    untersuchung=studie1,
                    matrikelnummer='1234567',
                )

        self.assertEqual(Teilnahme.objects.count(), 3)


class EinordnungTests(SimpleTestCase):
    """Prüft zulässige Einordnungswerte ohne Datenbankzugriff."""

    def test_fa07_gueltige_einordnung_bleibt_unveraendert(self):
        """Die zulässigen Werte P und O werden unverändert zurückgegeben."""
        self.assertEqual(pruefe_einordnung('P'), 'P')
        self.assertEqual(pruefe_einordnung('O'), 'O')

    def test_fa07_ungueltige_einordnung_wird_abgewiesen(self):
        """Andere Schreibweisen und Datentypen lösen einen Validierungsfehler aus."""
        for ungueltige_einordnung in ['p', 'o', ' P ', 'P ', '', 'X', None, 1, True]:
            with self.subTest(wert=ungueltige_einordnung):
                with self.assertRaises(ValidationError):
                    pruefe_einordnung(ungueltige_einordnung)


class MatrikelnummerTests(SimpleTestCase):
    """Prüft Zahlenformat und Wertebereich der Matrikelnummer ohne Datenbankzugriff."""

    def test_fa07_gueltige_matrikelnummer_wird_als_text_zurueckgegeben(self):
        """Gültige siebenstellige Ganzzahlen werden als Text zurückgegeben."""
        self.assertEqual(pruefe_matrikelnummer(1000000), '1000000')
        self.assertEqual(pruefe_matrikelnummer(1234567), '1234567')
        self.assertEqual(pruefe_matrikelnummer(9999999), '9999999')

    def test_fa07_ungueltige_matrikelnummer_wird_abgewiesen(self):
        """Werte außerhalb des Bereichs und unzulässige Datentypen werden abgewiesen."""
        for ungueltige_matrikelnummer in [999_999, 10_000_000, 0, -1234567, '1234567', 1234567.0, 1234567.5, None, True, False]:
            with self.subTest(wert=ungueltige_matrikelnummer):
                with self.assertRaises(ValidationError):
                    pruefe_matrikelnummer(ungueltige_matrikelnummer)


class LehrpersonTests(SimpleTestCase):
    """Prüft Bereinigung und Längengrenzen der Lehrpersonenangabe."""

    def test_fa07_gueltiger_lehrpersonenname_wird_bereinigt(self):
        """Zulässige Namen verlieren nur äußere Leerzeichen und behalten innere Leerzeichen."""
        self.assertEqual(
            pruefe_lehrperson('  Dr. Müller  '),
            'Dr. Müller'
        )

        self.assertEqual(
            pruefe_lehrperson('Herr  Müller'),
            'Herr  Müller'
        )

        self.assertEqual(pruefe_lehrperson('A'), 'A')
        self.assertEqual(
            # Die Längengrenze gilt nach Entfernung der äußeren Leerzeichen.
            pruefe_lehrperson('  ' + 'A' * 100 + '  '),
            'A' * 100
        )

    def test_fa07_ungueltiger_lehrpersonenname_wird_abgewiesen(self):
        """Leere und überlange Namen sowie unzulässige Datentypen werden abgewiesen."""
        for ungueltiger_name in ['', '   ', 'A' * 101, None, 123, True]:
            with self.subTest(wert=ungueltiger_name):
                with self.assertRaises(ValidationError):
                    pruefe_lehrperson(ungueltiger_name)


class StudiennameTests(SimpleTestCase):
    """Prüft Anzeigenamen und Vergleichswerte der Studiennamen."""

    def test_fa07_gueltiger_studienname_liefert_anzeigename_und_vergleichswert(self):
        """Zulässige Studiennamen liefern bereinigte Anzeigenamen und Vergleichswerte."""
        self.assertEqual(
            pruefe_studienname('  Teststudie A  '),
            ('Teststudie A', 'teststudie a')
        )

        self.assertEqual(
            pruefe_studienname('TESTSTUDIE A'),
            ('TESTSTUDIE A', 'teststudie a')
        )

        self.assertEqual(
            pruefe_studienname('Test  Studie'),
            ('Test  Studie', 'test  studie')
        )

        self.assertEqual(
            pruefe_studienname('A'),
            ('A', 'a')
        )

        self.assertEqual(
            # casefold darf den Vergleichswert verlängern: ß wird zu ss.
            pruefe_studienname('  ' + 'ß' * 100 + '  '),
            ('ß' * 100, 'ss' * 100)
        )

    def test_fa07_ungueltiger_studienname_wird_abgewiesen(self):
        """Leere und überlange Studiennamen sowie unzulässige Datentypen werden abgewiesen."""
        for ungueltiger_name in ['', '   ', 'A' * 101, None, 123, True]:
            with self.subTest(wert=ungueltiger_name):
                with self.assertRaises(ValidationError):
                    pruefe_studienname(ungueltiger_name)


class StundenwertTests(SimpleTestCase):
    """Prüft zulässige Stundenwerte und ihre Rückgabe als Fließkommazahl."""

    def test_fa07_gueltiger_stundenwert_wird_als_float_zurueckgegeben(self):
        """Zulässige positive Schritte von 0,5 VP-Stunden werden als float zurückgegeben."""
        testfaelle = [
            (0.5, 0.5),
            (1, 1.0),
            (1.5, 1.5),
            (2.0, 2.0),
            (1000, 1000.0),
        ]

        for eingabe, erwartet in testfaelle:
            # Jeder Teilfall erhält bei einem Fehler eine eigene Fundstelle.
            with self.subTest(wert=eingabe):
                ergebnis = pruefe_stundenwert(eingabe)

                self.assertEqual(ergebnis, erwartet)
                self.assertIsInstance(ergebnis, float)

    def test_fa07_ungueltiger_stundenwert_wird_abgewiesen(self):
        """Unzulässige Schritte, nicht positive Werte und falsche Datentypen werden abgewiesen."""
        for ungueltiger_wert in [
            0, -0.5, 0.25, 1.25, '1.5', '', None, True, False
        ]:

            with self.subTest(wert=ungueltiger_wert):
                with self.assertRaises(ValidationError):
                    pruefe_stundenwert(ungueltiger_wert)


class StundenberechnungTests(SimpleTestCase):
    """Prüft Stundensummen und Erfüllungsgrenzen ohne gespeicherte Modelldaten."""

    def test_fa16_leere_liste_liefert_nullsummen(self):
        """Eine leere Teilnahmeliste liefert Fließkomma-Nullsummen und einen unerfüllten Status."""
        ergebnis = berechne_stunden([])

        self.assertEqual(ergebnis, {
            "praesenzstunden": 0.0,
            "onlinestunden": 0.0,
            "gesamtstunden": 0.0,
            "erfuellt": False,
        })

        # Auch Nullsummen müssen den vereinbarten Zahlentyp behalten.
        for feldname in ["praesenzstunden", "onlinestunden", "gesamtstunden"]:
            self.assertIsInstance(ergebnis[feldname], float)

    def test_fa16_stunden_werden_nach_einordnung_addiert(self):
        """Mehrere Teilnahmen werden getrennt nach Einordnung und insgesamt addiert."""
        testfaelle = [
            (Untersuchung.PRAESENZ, 1.5),
            (Untersuchung.PRAESENZ, 2.5),
            (Untersuchung.ONLINE, 0.5),
            (Untersuchung.ONLINE, 1.0),
        ]

        teilnahmen = []

        for einordnung, stundenwert in testfaelle:
            # Ohne save() bleiben diese Modellobjekte ausschließlich im Arbeitsspeicher.
            untersuchung = Untersuchung(
                einordnung=einordnung,
                stundenwert=stundenwert,
            )
            teilnahmen.append(Teilnahme(untersuchung=untersuchung))

        ergebnis = berechne_stunden(teilnahmen)

        self.assertEqual(ergebnis, {
            "praesenzstunden": 4.0,
            "onlinestunden": 1.5,
            "gesamtstunden": 5.5,
            "erfuellt": False,
        })

    def test_fa16_erfuellungsgrenzen_werden_beachtet(self):
        """Gesamt- und Präsenzminimum werden unabhängig und einschließlich der Grenzen geprüft."""
        testfaelle = [
            (15.0, 15.0, 30.0, True),    # Beide Mindestwerte genau erreicht.
            (15.0, 14.5, 29.5, False),   # Zu wenige Gesamtstunden.
            (14.5, 15.5, 30.0, False),   # Zu wenige Präsenzstunden.
            (30.0, 0.0, 30.0, True),     # Onlinestunden sind nicht erforderlich.
        ]

        for praesenzstunden, onlinestunden, gesamtstunden, erfuellt in testfaelle:
            with self.subTest(
                praesenzstunden=praesenzstunden,
                onlinestunden=onlinestunden,
            ):

                # Die Berechnung benötigt Modellobjekte, aber keine gespeicherten Datensätze.
                praesenzuntersuchung = Untersuchung(
                    einordnung=Untersuchung.PRAESENZ,
                    stundenwert=praesenzstunden,
                )

                teilnahmen = [
                    Teilnahme(untersuchung=praesenzuntersuchung),
                ]

                # Null Onlinestunden bedeuten hier, dass keine Online-Teilnahme vorliegt.
                if onlinestunden > 0:
                    onlineuntersuchung = Untersuchung(
                        einordnung=Untersuchung.ONLINE,
                        stundenwert=onlinestunden,
                    )

                    teilnahmen.append(
                        Teilnahme(untersuchung=onlineuntersuchung)
                    )

                ergebnis = berechne_stunden(teilnahmen)

                self.assertEqual(ergebnis, {
                    "praesenzstunden": praesenzstunden,
                    "onlinestunden": onlinestunden,
                    "gesamtstunden": gesamtstunden,
                    "erfuellt": erfuellt,
                })

                # Identität prüft zusätzlich den bool-Typ statt nur die Gleichheit mit 0 oder 1.
                self.assertIs(ergebnis["erfuellt"], erfuellt)


class TeilnahmesucheTests(TestCase):
    """Prüft Auswahl, Sortierung und mitgeladene Untersuchungsangaben."""

    def setUp(self):
        """Legt sortierbare Teilnahmen und einen fremden Eintrag als Vergleich an."""
        # Entgegen der Namensreihenfolge anlegen, damit eine fehlende Sortierung auffällt.
        untersuchung_beta = Untersuchung.objects.create(
            studienname="Beta",
            studienname_vergleich="beta",
            stundenwert=1.5,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Testlehrperson Beta",
        )

        untersuchung_alpha = Untersuchung.objects.create(
            studienname="alpha",
            studienname_vergleich="alpha",
            stundenwert=2.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alpha",
        )

        self.teilnahme_beta = Teilnahme.objects.create(
            untersuchung=untersuchung_beta,
            matrikelnummer="1234567",
        )

        self.teilnahme_alpha = Teilnahme.objects.create(
            untersuchung=untersuchung_alpha,
            matrikelnummer="1234567",
        )

        # Eine fremde Matrikelnummer darf weder Suche noch Berechnung beeinflussen.
        Teilnahme.objects.create(
            untersuchung=untersuchung_beta,
            matrikelnummer="7654321",
        )

    def test_fa14_passende_teilnahmen_werden_sortiert_geladen(self):
        """Die Suche liefert passende Teilnahmen sortiert und ohne weitere Nachladeabfragen."""
        teilnahmen = suche_teilnahmen("1234567")

        self.assertIsInstance(teilnahmen, list)
        self.assertEqual(
            teilnahmen,
            [self.teilnahme_alpha, self.teilnahme_beta],
        )

        # Untersuchungsangaben erstmals hier lesen: Dabei darf keine weitere Abfrage entstehen.
        with self.assertNumQueries(0):
            self.assertEqual(
                teilnahmen[0].untersuchung.studienname,
                "alpha",
            )
            self.assertEqual(
                teilnahmen[0].untersuchung.lehrperson,
                "Testlehrperson Alpha",
            )
            self.assertEqual(
                teilnahmen[1].untersuchung.studienname,
                "Beta",
            )
            self.assertEqual(
                teilnahmen[1].untersuchung.lehrperson,
                "Testlehrperson Beta",
            )

            berechnung = berechne_stunden(teilnahmen)

        self.assertEqual(berechnung, {
            "praesenzstunden": 2.5,
            "onlinestunden": 1.5,
            "gesamtstunden": 4.0,
            "erfuellt": False,
        })

    def test_fa14_fehlende_matrikelnummer_liefert_leere_liste(self):
        """Eine unbekannte Matrikelnummer liefert trotz anderer Teilnahmen eine leere Liste."""
        ergebnis = suche_teilnahmen("9999999")

        self.assertEqual(ergebnis, [])


class UntersuchungskorrekturTests(TestCase):
    """Prüft Korrekturen, Namenskonflikte und unveränderte Bestände bei Fehlern."""

    def setUp(self):
        """Legt zwei Untersuchungen, zugehörige Teilnahmen und feste Vergleichsbestände an."""
        self.untersuchung = Untersuchung.objects.create(
            studienname="Ausgangsstudie",
            studienname_vergleich="ausgangsstudie",
            stundenwert=1.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Alte Lehrperson",
        )

        # Die zweite Untersuchung macht Konflikte mit bereits vergebenen Namen prüfbar.
        self.andere_untersuchung = Untersuchung.objects.create(
            studienname="Andere Studie",
            studienname_vergleich="andere studie",
            stundenwert=2.0,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Andere Lehrperson",
        )

        for matrikelnummer in ["1234567", "7654321"]:
            Teilnahme.objects.create(
                untersuchung=self.untersuchung,
                matrikelnummer=matrikelnummer,
            )

        self.uploadzeitpunkt = self.untersuchung.uploadzeitpunkt

        # Feldwerte sofort als Liste festhalten, einschließlich IDs und Uploadzeitpunkt.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )

        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

    def test_fa15_gueltige_korrektur_wird_gespeichert(self):
        """Eine gültige Korrektur ändert Untersuchungsangaben und erhält Identität und Teilnahmen."""
        ergebnis, fehlerliste = korrigiere_untersuchung(
            untersuchung_id=self.untersuchung.pk,
            studienname="  Neue Studie ",
            stundenwert=2.5,
            einordnung=Untersuchung.ONLINE,
            lehrperson="  Neue Lehrperson  ",
        )

        self.assertEqual(fehlerliste, [])

        self.assertIsInstance(ergebnis, Untersuchung)

        self.assertEqual(ergebnis.pk, self.untersuchung.pk)

        self.assertEqual(ergebnis.studienname, "Neue Studie")

        # Neu aus der Datenbank laden, damit tatsächlich gespeicherte Werte geprüft werden.
        self.untersuchung.refresh_from_db()

        self.assertEqual(self.untersuchung.studienname, "Neue Studie")
        self.assertEqual(
            self.untersuchung.studienname_vergleich,
            "neue studie",
        )

        self.assertEqual(self.untersuchung.stundenwert, 2.5)
        self.assertEqual(self.untersuchung.einordnung, Untersuchung.ONLINE)
        self.assertEqual(self.untersuchung.lehrperson, "Neue Lehrperson")

        self.assertEqual(
            self.untersuchung.uploadzeitpunkt,
            self.uploadzeitpunkt,
        )

        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

        self.assertEqual(Untersuchung.objects.count(), 2)

        # Die zweite, später angelegte Untersuchung muss vollständig unverändert bleiben.
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values())[1],
            self.untersuchungsbestand[1],
        )

    def test_fa15_eigener_studienname_bleibt_erlaubt(self):
        """Eine andere Schreibweise des eigenen Studiennamens löst keinen Namenskonflikt aus."""
        # Der eigene Vergleichswert darf beim Namensabgleich keinen Konflikt auslösen.
        ergebnis, fehlerliste = korrigiere_untersuchung(
            untersuchung_id=self.untersuchung.pk,
            studienname="  AUSGANGSSTUDIE  ",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Neue Lehrperson",
        )

        self.assertEqual(fehlerliste, [])
        self.assertIsInstance(ergebnis, Untersuchung)
        self.assertEqual(ergebnis.pk, self.untersuchung.pk)

        self.untersuchung.refresh_from_db()

        self.assertEqual(self.untersuchung.studienname, "AUSGANGSSTUDIE")
        self.assertEqual(
            self.untersuchung.studienname_vergleich,
            "ausgangsstudie",
        )
        self.assertEqual(self.untersuchung.stundenwert, 1.5)

    def test_fa15_ungueltige_angaben_werden_gemeinsam_abgewiesen(self):
        """Vier ungültige Angaben werden gemeinsam gemeldet, ohne den Bestand zu ändern."""
        ergebnis, fehlerliste = korrigiere_untersuchung(
            untersuchung_id=self.untersuchung.pk,
            studienname="",
            stundenwert=0,
            einordnung="X",
            lehrperson="   ",
        )

        self.assertIsNone(ergebnis)

        self.assertEqual(len(fehlerliste), 4)

        meldungen = " ".join(fehlerliste)

        # subTest kennzeichnet Fehler je Feld und führt die übrigen Einzelprüfungen fort.
        for feldname in ["Studienname", "Stundenwert", "Einordnung", "Lehrperson"]:
            with self.subTest(feldname=feldname):
                self.assertIn(f"{feldname}:", meldungen)

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa15_name_einer_anderen_untersuchung_wird_abgewiesen(self):
        """Der Name einer anderen Untersuchung wird auch in abweichender Schreibweise abgewiesen."""
        # Großschreibung und äußere Leerzeichen dürfen den Namenskonflikt nicht verdecken.
        ergebnis, fehlerliste = korrigiere_untersuchung(
            untersuchung_id=self.untersuchung.pk,
            studienname="  ANDERE STUDIE  ",
            stundenwert=3.0,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Neue Lehrperson",
        )

        self.assertIsNone(ergebnis)
        self.assertEqual(len(fehlerliste), 1)

        self.assertTrue(fehlerliste[0].startswith("Studienname:"))

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa15_fehlende_untersuchung_wird_abgewiesen(self):
        """Eine fehlende Untersuchungs-ID führt zu einer kontrollierten Abweisung ohne Änderungen."""
        # Eine freie ID ableiten und ausdrücklich prüfen, dass sie nicht belegt ist.
        fehlende_id = self.andere_untersuchung.pk + 1

        self.assertFalse(
            Untersuchung.objects.filter(pk=fehlende_id).exists()
        )

        ergebnis, fehlerliste = korrigiere_untersuchung(
            untersuchung_id=fehlende_id,
            studienname="Neue Studie",
            stundenwert=2.5,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Neue Lehrperson",
        )

        self.assertIsNone(ergebnis)
        self.assertEqual(len(fehlerliste), 1)
        self.assertTrue(fehlerliste[0].startswith("Untersuchung:"))

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa15_speicherfehler_wird_kontrolliert_behandelt(self):
        """Ein Speicherfehler wird kontrolliert gemeldet und erhält den ursprünglichen Bestand."""
        # side_effect löst nur im Block einen Speicherfehler aus, ohne Daten zu schreiben.
        with patch.object(
            Untersuchung,
            "save",
            side_effect=DatabaseError("Synthetischer Speicherfehler"),
        ) as speichern:
            ergebnis, fehlerliste = korrigiere_untersuchung(
                untersuchung_id=self.untersuchung.pk,
                studienname="Neue Studie",
                stundenwert=2.5,
                einordnung=Untersuchung.ONLINE,
                lehrperson="Neue Lehrperson",
            )

        # Der Test muss den vorgesehenen Speicheraufruf tatsächlich erreicht haben.
        speichern.assert_called_once()

        self.assertIsNone(ergebnis)
        self.assertEqual(fehlerliste, [
            "Untersuchung: Die Änderungen konnten nicht gespeichert werden."
        ])

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )


class TeilnahmedatenLoeschungTests(TestCase):
    """Prüft Löschvoraussetzungen, verbleibende Daten und Transaktionsrücknahme."""

    def setUp(self):
        """Legt einen erfüllten synthetischen Bestand und feste Vergleichswerte an."""
        self.matrikelnummer = "1234567"

        self.praesenzuntersuchung = Untersuchung.objects.create(
            studienname="Präsenzstudie",
            studienname_vergleich="präsenzstudie",
            stundenwert=15.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Präsenz",
        )

        self.onlineuntersuchung = Untersuchung.objects.create(
            studienname="Onlinestudie",
            studienname_vergleich="onlinestudie",
            stundenwert=15.0,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Testlehrperson Online",
        )

        # Beide Teilnahmen zusammen erfüllen Gesamt- und Präsenzstundengrenze.
        for untersuchung in [
            self.praesenzuntersuchung,
            self.onlineuntersuchung,
        ]:
            Teilnahme.objects.create(
                untersuchung=untersuchung,
                matrikelnummer=self.matrikelnummer,
            )

        # Alle Feldwerte festhalten, um auch unbemerkte Änderungen bei gleicher Anzahl zu erkennen.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

    def test_fa18_erfuellte_teilnahmen_werden_geloescht(self):
        """Bestätigte, erfüllte Teilnahmen und dadurch leere Untersuchungen werden gelöscht."""
        fehlerliste = loesche_teilnahmedaten(
            matrikelnummer=self.matrikelnummer,
            ablage_bestaetigt=True,
        )

        self.assertEqual(fehlerliste, [])

        self.assertEqual(Teilnahme.objects.count(), 0)
        self.assertEqual(Untersuchung.objects.count(), 0)

    def test_fa18_ohne_ablagebestaetigung_bleibt_bestand_erhalten(self):
        """Ohne Bestätigung der externen Reportablage bleiben alle Daten erhalten."""
        fehlerliste = loesche_teilnahmedaten(
            matrikelnummer=self.matrikelnummer,
            ablage_bestaetigt=False,
        )

        self.assertTrue(fehlerliste)

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa18_aktuell_zu_wenige_gesamtstunden_verhindern_loeschung(self):
        """Nachträglich gesunkene Gesamtstunden verhindern die Löschung."""
        # 29,5 Gesamtstunden unterschreiten nur die Gesamtstundengrenze.
        self.onlineuntersuchung.stundenwert = 14.5
        self.onlineuntersuchung.save(update_fields=["stundenwert"])

        # Die vorbereitete Änderung gehört zum neuen Vergleichsstand.
        untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )

        fehlerliste = loesche_teilnahmedaten(
            matrikelnummer=self.matrikelnummer,
            ablage_bestaetigt=True,
        )

        self.assertTrue(fehlerliste)

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            untersuchungsbestand,
        )

        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa18_aktuell_zu_wenige_praesenzstunden_verhindern_loeschung(self):
        """Zu wenige aktuelle Präsenzstunden verhindern trotz ausreichender Gesamtstunden die Löschung."""
        # Online einordnen: 30 Gesamtstunden bleiben erhalten, Präsenzstunden sinken auf 0.
        self.praesenzuntersuchung.einordnung = Untersuchung.ONLINE
        self.praesenzuntersuchung.save(update_fields=["einordnung"])

        untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )

        fehlerliste = loesche_teilnahmedaten(
            matrikelnummer=self.matrikelnummer,
            ablage_bestaetigt=True,
        )

        self.assertTrue(fehlerliste)

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            untersuchungsbestand,
        )

        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa18_fehlende_matrikelnummer_veraendert_keine_daten(self):
        """Eine Matrikelnummer ohne Teilnahmen führt zu einer Abweisung ohne Änderungen."""
        fehlende_matrikelnummer = "9999999"

        self.assertFalse(
            Teilnahme.objects.filter(
                matrikelnummer=fehlende_matrikelnummer,
            ).exists()
        )

        fehlerliste = loesche_teilnahmedaten(
            matrikelnummer=fehlende_matrikelnummer,
            ablage_bestaetigt=True,
        )

        self.assertTrue(fehlerliste)

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )

        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa18_untersuchung_mit_fremder_teilnahme_bleibt_erhalten(self):
        """Eine Untersuchung mit einer fremden Teilnahme bleibt bei der Löschung erhalten."""
        fremde_teilnahme = Teilnahme.objects.create(
            untersuchung=self.onlineuntersuchung,
            matrikelnummer="7654321",
        )

        # Nur die fremde Teilnahme und ihre weiterhin benötigte Untersuchung müssen verbleiben.
        verbleibender_untersuchungsbestand = list(
            Untersuchung.objects.filter(
                pk=self.onlineuntersuchung.pk,
            ).values()
        )
        verbleibender_teilnahmebestand = list(
            Teilnahme.objects.filter(
                pk=fremde_teilnahme.pk,
            ).values()
        )

        fehlerliste = loesche_teilnahmedaten(
            matrikelnummer=self.matrikelnummer,
            ablage_bestaetigt=True,
        )

        self.assertEqual(fehlerliste, [])

        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            verbleibender_untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            verbleibender_teilnahmebestand,
        )

    def test_fa18_datenbankfehler_setzt_teilnahmeloeschung_zurueck(self):
        """Ein nachgelagerter Datenbankfehler nimmt die bereits erfolgte Teilnahmelöschung zurück."""
        # Erst nach der Teilnahmelöschung einen Fehler beim Untersuchungszugriff auslösen.
        with patch.object(
            Untersuchung.objects,
            "filter",
            side_effect=DatabaseError("Synthetischer Löschfehler"),
        ) as untersuchungsabfrage:
            fehlerliste = loesche_teilnahmedaten(
                matrikelnummer=self.matrikelnummer,
                ablage_bestaetigt=True,
            )

        untersuchungsabfrage.assert_called_once()

        self.assertEqual(fehlerliste, [
            "Die Teilnahmedaten konnten nicht gelöscht werden."
        ])

        # Die Transaktion muss auch die bereits gelöschten Teilnahmen wiederherstellen.
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )


class TeilnahmedatenLoeschungViewTests(TestCase):
    """Prüft Löschanfragen, Rollenrechte und die Bestätigung der Reportablage."""

    def setUp(self):
        """Legt einen erfüllten Datenbestand und feste Vergleichswerte für Löschanfragen an."""
        self.matrikelnummer = "1234567"

        self.untersuchung = Untersuchung.objects.create(
            studienname="Teststudie Löschview",
            studienname_vergleich="teststudie löschview",
            stundenwert=30.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson",
        )

        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer=self.matrikelnummer,
        )

        # Den vollständigen Bestand vor der Anfrage als feste Feldwerte sichern.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

    def melde_testkonto_an(self, rolle):
        """Legt ein synthetisches Rollenkonto an und richtet seine Sitzung per Anmeldung ein."""
        benutzername = "testkonto"
        passwort = "a" * 64

        konto, fehlerliste = lege_konto_an(
            benutzername,
            passwort,
            rolle,
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Die reguläre Anmeldung richtet zugleich die erforderlichen Sitzungszeitwerte ein.
        antwort = self.client.post(
            "/anmelden/",
            data={
                "benutzername": benutzername,
                "passwort": passwort,
            },
        )
        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()["erfolg"])

    def pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche gespeicherten Feldwerte mit dem vorbereiteten Bestand."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa18_administrator_kann_teilnahmedaten_loeschen(self):
        """Ein Administrationskonto kann bei erfüllten Voraussetzungen die Daten löschen."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.post(
            "/teilnahmedaten_loeschen/",
            data={
                "matrikelnummer": self.matrikelnummer,
                "ablage_bestaetigt": "true",
            },
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()["erfolg"])

        self.assertEqual(Teilnahme.objects.count(), 0)
        self.assertEqual(Untersuchung.objects.count(), 0)

    def test_qa04_lehrperson_kann_keine_teilnahmedaten_loeschen(self):
        """Eine Lehrperson darf trotz erfüllter Löschvoraussetzungen keine Daten löschen."""
        self.melde_testkonto_an(Benutzerkonto.LEHRPERSON)

        # Gültige Löschdaten isolieren die fehlende Rollenberechtigung als Abweisungsgrund.
        antwort = self.client.post(
            "/teilnahmedaten_loeschen/",
            data={
                "matrikelnummer": self.matrikelnummer,
                "ablage_bestaetigt": "true",
            },
        )

        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()["erfolg"])
        self.pruefe_unveraenderten_bestand()

    def test_qa04_loeschung_ohne_anmeldung_wird_abgewiesen(self):
        """Eine nicht angemeldete Löschanfrage wird ohne Datenänderung abgewiesen."""
        antwort = self.client.post(
            "/teilnahmedaten_loeschen/",
            data={
                "matrikelnummer": self.matrikelnummer,
                "ablage_bestaetigt": "true",
            },
        )

        self.assertEqual(antwort.status_code, 403)
        self.assertFalse(antwort.json()["erfolg"])
        self.pruefe_unveraenderten_bestand()

    def test_fa18_get_anfrage_wird_abgewiesen(self):
        """Eine GET-Anfrage wird abgewiesen und weist POST als erlaubte Methode aus."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.get("/teilnahmedaten_loeschen/")

        self.assertEqual(antwort.status_code, 405)
        self.assertFalse(antwort.json()["erfolg"])

        self.assertEqual(antwort.headers["Allow"], "POST")
        self.pruefe_unveraenderten_bestand()

    def test_fa18_fehlende_matrikelnummer_wird_abgewiesen(self):
        """Eine fehlende Matrikelnummer wird vor dem Aufruf der Löschfunktion abgewiesen."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Der Mock erfasst, ob die View trotz fehlender Matrikelnummer die Löschung aufruft.
        with patch(
            "vp_stunden.views.loesche_teilnahmedaten"
        ) as loeschfunktion:
            antwort = self.client.post(
                "/teilnahmedaten_loeschen/",
                data={
                    "ablage_bestaetigt": "true",
                },
            )

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()["erfolg"])

        loeschfunktion.assert_not_called()
        self.pruefe_unveraenderten_bestand()

    def test_fa18_ohne_ablagebestaetigung_wird_loeschung_abgewiesen(self):
        """Eine fehlende oder ausdrücklich verneinte Ablagebestätigung verhindert die Löschung."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Ein fehlendes Feld und der Formulartext "false" müssen beide die Löschung verhindern.
        testfaelle = [
            {
                "matrikelnummer": self.matrikelnummer,
            },
            {
                "matrikelnummer": self.matrikelnummer,
                "ablage_bestaetigt": "false",
            },
        ]

        for daten in testfaelle:
            with self.subTest(daten=daten):
                antwort = self.client.post(
                    "/teilnahmedaten_loeschen/",
                    data=daten,
                )

                self.assertEqual(antwort.status_code, 400)
                self.assertFalse(antwort.json()["erfolg"])
                self.assertTrue(antwort.json()["fehlerliste"])
                self.pruefe_unveraenderten_bestand()

    def test_fa18_nicht_erfuellte_mindestvorgaben_verhindern_loeschung(self):
        """Aktuell unterschrittene Mindestvorgaben verhindern die Löschung über die View."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        self.untersuchung.stundenwert = 29.5
        self.untersuchung.save(update_fields=["stundenwert"])

        # Die vorbereitete Änderung gehört zum neuen Vergleichsstand.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )

        antwort = self.client.post(
            "/teilnahmedaten_loeschen/",
            data={
                "matrikelnummer": self.matrikelnummer,
                "ablage_bestaetigt": "true",
            },
        )

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()["erfolg"])
        self.assertTrue(antwort.json()["fehlerliste"])
        self.pruefe_unveraenderten_bestand()


class TeilnahmesucheViewTests(TestCase):
    """Prüft Suchantworten, Stundenberechnung und Zugriffsabweisungen."""

    def setUp(self):
        """Legt einen erfüllten Suchbestand, eine fremde Teilnahme und Vergleichswerte an."""
        self.matrikelnummer = "1234567"

        # Entgegen der Namensreihenfolge anlegen, damit eine fehlende Sortierung auffällt.
        self.untersuchung_beta = Untersuchung.objects.create(
            studienname="Beta-Studie",
            studienname_vergleich="beta-studie",
            stundenwert=15.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Beta",
        )

        self.untersuchung_alpha = Untersuchung.objects.create(
            studienname="Alpha-Studie",
            studienname_vergleich="alpha-studie",
            stundenwert=15.0,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Testlehrperson Alpha",
        )

        # Ohne Mikrosekunden lässt sich der gespeicherte Zeitpunkt mit der JSON-Ausgabe vergleichen.
        self.uploadzeitpunkt = timezone.now().replace(microsecond=0)
        Untersuchung.objects.update(
            uploadzeitpunkt=self.uploadzeitpunkt,
        )

        for untersuchung in [
            self.untersuchung_beta,
            self.untersuchung_alpha,
        ]:
            Teilnahme.objects.create(
                untersuchung=untersuchung,
                matrikelnummer=self.matrikelnummer,
            )

        # Die fremde Teilnahme darf weder Suchergebnis noch Stundensummen beeinflussen.
        Teilnahme.objects.create(
            untersuchung=self.untersuchung_beta,
            matrikelnummer="7654321",
        )

        # Alle Feldwerte vor der Anfrage festhalten, um unbeabsichtigte Änderungen zu erkennen.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

    def melde_testkonto_an(self, rolle):
        """Legt ein synthetisches Rollenkonto an und richtet seine Sitzung per Anmeldung ein."""
        benutzername = "testkonto"
        passwort = "a" * 64

        konto, fehlerliste = lege_konto_an(
            benutzername,
            passwort,
            rolle,
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Die reguläre Anmeldung richtet zugleich die erforderlichen Sitzungszeitwerte ein.
        antwort = self.client.post(
            "/anmelden/",
            data={
                "benutzername": benutzername,
                "passwort": passwort,
            },
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()["erfolg"])

    def pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche gespeicherten Feldwerte mit dem vorbereiteten Bestand."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa14_administrator_erhaelt_sortierte_teilnahmen(self):
        """Die Suchantwort enthält passende Teilnahmen in Namensreihenfolge und ihre Stundensummen."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.post(
            "/teilnahmen_suchen/",
            data={
                "matrikelnummer": self.matrikelnummer,
            },
        )

        self.assertEqual(antwort.status_code, 200)
        ergebnis = antwort.json()
        self.assertTrue(ergebnis["erfolg"])

        self.assertEqual(len(ergebnis["teilnahmen"]), 2)

        erwartete_untersuchungen = [
            self.untersuchung_alpha,
            self.untersuchung_beta,
        ]

        # zip verbindet jede Antwortzeile mit ihrer erwarteten Untersuchung und Position.
        for zeile, untersuchung in zip(
            ergebnis["teilnahmen"],
            erwartete_untersuchungen,
        ):
            with self.subTest(studienname=untersuchung.studienname):
                self.assertEqual(
                    zeile["matrikelnummer"],
                    self.matrikelnummer,
                )
                self.assertEqual(
                    zeile["studienname"],
                    untersuchung.studienname,
                )
                self.assertEqual(
                    zeile["stundenwert"],
                    untersuchung.stundenwert,
                )
                self.assertEqual(
                    zeile["einordnung"],
                    untersuchung.einordnung,
                )
                self.assertEqual(
                    zeile["lehrperson"],
                    untersuchung.lehrperson,
                )

                # Den JSON-Text für den Vergleich zurück in einen Zeitpunkt umwandeln.
                self.assertEqual(
                    datetime.fromisoformat(zeile["uploadzeitpunkt"]),
                    self.uploadzeitpunkt,
                )

        self.assertEqual(ergebnis["berechnung"], {
            "praesenzstunden": 15.0,
            "onlinestunden": 15.0,
            "gesamtstunden": 30.0,
            "erfuellt": True,
        })

        self.pruefe_unveraenderten_bestand()

    def test_fa16_nicht_erfuellter_status_wird_angezeigt(self):
        """Eine erfolgreiche Suche zeigt auch einen noch nicht erfüllten Stundenstatus an."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        self.untersuchung_alpha.stundenwert = 14.5
        self.untersuchung_alpha.save(update_fields=["stundenwert"])

        # Die vorbereitete Änderung gehört zum neuen Vergleichsstand.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )

        antwort = self.client.post(
            "/teilnahmen_suchen/",
            data={
                "matrikelnummer": self.matrikelnummer,
            },
        )

        # Ein nicht erfüllter Stundenstatus ist ein gültiges Suchergebnis.
        self.assertEqual(antwort.status_code, 200)
        ergebnis = antwort.json()
        self.assertTrue(ergebnis["erfolg"])
        self.assertEqual(len(ergebnis["teilnahmen"]), 2)

        self.assertEqual(ergebnis["berechnung"], {
            "praesenzstunden": 15.0,
            "onlinestunden": 14.5,
            "gesamtstunden": 29.5,
            "erfuellt": False,
        })

        self.pruefe_unveraenderten_bestand()

    def test_qa04_lehrperson_kann_keine_teilnahmen_suchen(self):
        """Eine Lehrperson erhält weder Teilnahmeangaben noch berechnete Stundensummen."""
        self.melde_testkonto_an(Benutzerkonto.LEHRPERSON)

        antwort = self.client.post(
            "/teilnahmen_suchen/",
            data={
                "matrikelnummer": self.matrikelnummer,
            },
        )

        self.assertEqual(antwort.status_code, 403)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])

        # Eine verweigerte Anfrage darf auch keine Teil- oder Berechnungsergebnisse liefern.
        self.assertNotIn("teilnahmen", ergebnis)
        self.assertNotIn("berechnung", ergebnis)
        self.pruefe_unveraenderten_bestand()

    def test_qa04_suche_ohne_anmeldung_wird_abgewiesen(self):
        """Eine nicht angemeldete Suche wird ohne Ergebnisdaten und Änderungen abgewiesen."""
        antwort = self.client.post(
            "/teilnahmen_suchen/",
            data={
                "matrikelnummer": self.matrikelnummer,
            },
        )

        self.assertEqual(antwort.status_code, 403)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("teilnahmen", ergebnis)
        self.assertNotIn("berechnung", ergebnis)
        self.pruefe_unveraenderten_bestand()

    def test_fa14_get_anfrage_wird_abgewiesen(self):
        """Eine GET-Suche wird abgewiesen und weist POST als erlaubte Methode aus."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.get(
            "/teilnahmen_suchen/",
            data={
                "matrikelnummer": self.matrikelnummer,
            },
        )

        self.assertEqual(antwort.status_code, 405)
        self.assertFalse(antwort.json()["erfolg"])
        self.assertEqual(antwort.headers["Allow"], "POST")
        self.pruefe_unveraenderten_bestand()

    def test_fa14_fehlende_matrikelnummer_wird_abgewiesen(self):
        """Eine fehlende Matrikelnummer wird vor dem Aufruf der Suchfunktion abgewiesen."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Der Mock erfasst, ob die unvollständige Anfrage trotzdem eine Suche auslöst.
        with patch(
            "vp_stunden.views.suche_teilnahmen"
        ) as suchfunktion:
            antwort = self.client.post(
                "/teilnahmen_suchen/",
                data={},
            )

        self.assertEqual(antwort.status_code, 400)
        self.assertFalse(antwort.json()["erfolg"])

        suchfunktion.assert_not_called()
        self.pruefe_unveraenderten_bestand()

    def test_fa14_fehlende_treffer_werden_gemeldet(self):
        """Eine Suche ohne Treffer meldet den fehlenden Bestand ohne Teilnahmeangaben oder Summen."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        fehlende_matrikelnummer = "9999999"

        self.assertFalse(
            Teilnahme.objects.filter(
                matrikelnummer=fehlende_matrikelnummer,
            ).exists()
        )

        antwort = self.client.post(
            "/teilnahmen_suchen/",
            data={
                "matrikelnummer": fehlende_matrikelnummer,
            },
        )

        self.assertEqual(antwort.status_code, 404)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])

        self.assertNotIn("teilnahmen", ergebnis)
        self.assertNotIn("berechnung", ergebnis)
        self.pruefe_unveraenderten_bestand()


class UntersuchungssucheViewTests(TestCase):
    """Prüft Untersuchungsübersichten, Teilnamensuche und Zugriffsabweisungen."""

    def setUp(self):
        """Legt drei synthetische Untersuchungen außerhalb der erwarteten Namensreihenfolge an."""
        # Entgegen der Namensreihenfolge anlegen, damit eine fehlende Sortierung auffällt.
        self.beta = Untersuchung.objects.create(
            studienname="Beta-Studie",
            studienname_vergleich="beta-studie",
            stundenwert=1.5,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Testlehrperson Beta",
        )
        self.alpha = Untersuchung.objects.create(
            studienname="Alpha-Studie",
            studienname_vergleich="alpha-studie",
            stundenwert=2.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alpha",
        )
        self.gamma = Untersuchung.objects.create(
            studienname="Gamma-Projekt",
            studienname_vergleich="gamma-projekt",
            stundenwert=1.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Gamma",
        )

    def melde_testkonto_an(self, rolle):
        """Legt ein synthetisches Rollenkonto an und richtet seine Sitzung per Anmeldung ein."""
        testpasswort = "a" * 64
        konto, fehlerliste = lege_konto_an(
            "testkonto",
            testpasswort,
            rolle,
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Die reguläre Anmeldung richtet zugleich die erforderlichen Sitzungsdaten ein.
        anmeldung = self.client.post(
            "/anmelden/",
            data={
                "benutzername": "testkonto",
                "passwort": testpasswort,
            },
        )
        self.assertEqual(anmeldung.status_code, 200)
        self.assertTrue(anmeldung.json()["erfolg"])

    def test_fa14_administrator_erhaelt_passende_sortierte_untersuchungen(self):
        """Die Teilnamensuche ignoriert Großschreibung und liefert passende Untersuchungen sortiert."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Gemischte Großschreibung und ein Teilname prüfen denselben Suchbegriff.
        antwort = self.client.post(
            "/untersuchungen_suchen/",
            data={"suchtext": "sTuDiE"},
        )
        self.assertEqual(antwort.status_code, 200)

        ergebnis = antwort.json()
        self.assertTrue(ergebnis["erfolg"])
        self.assertEqual(len(ergebnis["untersuchungen"]), 2)

        # Antwortzeilen mit ihren erwarteten Untersuchungen und Positionen verbinden.
        for zeile, untersuchung in zip(
            ergebnis["untersuchungen"],
            [self.alpha, self.beta],
        ):
            self.assertEqual(zeile["untersuchung_id"], untersuchung.pk)
            self.assertEqual(zeile["studienname"], untersuchung.studienname)
            self.assertEqual(zeile["stundenwert"], untersuchung.stundenwert)
            self.assertEqual(zeile["einordnung"], untersuchung.einordnung)
            self.assertEqual(zeile["lehrperson"], untersuchung.lehrperson)

            # Den JSON-Text für den Vergleich zurück in einen Zeitpunkt umwandeln.
            self.assertEqual(
                datetime.fromisoformat(zeile["uploadzeitpunkt"]),
                untersuchung.uploadzeitpunkt,
            )

    def test_qa04_suche_ohne_anmeldung_wird_abgewiesen(self):
        """Eine nicht angemeldete Anfrage erhält keine Untersuchungsangaben."""
        antwort = self.client.post(
            "/untersuchungen_suchen/",
            data={"suchtext": "Studie"},
        )

        self.assertEqual(antwort.status_code, 403)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("untersuchungen", ergebnis)

    def test_qa04_lehrperson_kann_keine_untersuchungen_suchen(self):
        """Eine Lehrperson erhält keinen Zugriff auf die Untersuchungssuche."""
        self.melde_testkonto_an(Benutzerkonto.LEHRPERSON)

        antwort = self.client.post(
            "/untersuchungen_suchen/",
            data={"suchtext": "Studie"},
        )

        self.assertEqual(antwort.status_code, 403)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("untersuchungen", ergebnis)

    def test_fa14_get_anfrage_wird_abgewiesen(self):
        """Eine GET-Suche wird abgewiesen und weist POST als erlaubte Methode aus."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.get(
            "/untersuchungen_suchen/",
            data={"suchtext": "Studie"},
        )

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort.headers["Allow"], "POST")
        self.assertFalse(antwort.json()["erfolg"])

    def test_fa14_fehlender_suchtext_wird_abgewiesen(self):
        """Ein fehlender Suchparameter wird vor dem Suchaufruf abgewiesen."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Der Mock macht einen unzulässigen Suchaufruf bei fehlendem Parameter sichtbar.
        with patch(
            "vp_stunden.views.suche_untersuchungen"
        ) as suchfunktion:
            antwort = self.client.post(
                "/untersuchungen_suchen/",
                data={},
            )

        self.assertEqual(antwort.status_code, 400)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("untersuchungen", ergebnis)
        suchfunktion.assert_not_called()

    def test_fa14_fehlende_treffer_werden_gemeldet(self):
        """Eine erfolglose Suche meldet fehlende Treffer ohne Untersuchungsdaten."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.post(
            "/untersuchungen_suchen/",
            data={"suchtext": "Unbekannte Untersuchung"},
        )

        self.assertEqual(antwort.status_code, 404)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("untersuchungen", ergebnis)

    def test_fa14_leerer_suchtext_liefert_alle_untersuchungen(self):
        """Ein ausdrücklich leerer Suchtext liefert die vollständige, alphabetische Übersicht."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Ein leerer Wert ist zulässig und unterscheidet sich vom fehlenden Parameter.
        antwort = self.client.post(
            "/untersuchungen_suchen/",
            data={"suchtext": ""},
        )
        self.assertEqual(antwort.status_code, 200)

        ergebnis = antwort.json()
        self.assertTrue(ergebnis["erfolg"])

        gefundene_ids = []
        for zeile in ergebnis["untersuchungen"]:
            gefundene_ids.append(zeile["untersuchung_id"])

        self.assertEqual(
            gefundene_ids,
            [self.alpha.pk, self.beta.pk, self.gamma.pk],
        )


class UntersuchungskorrekturViewTests(TestCase):
    """Prüft Korrekturanfragen, Eingabeprüfung und Rollenrechte."""

    def setUp(self):
        """Legt eine Untersuchung, vollständige Korrekturdaten und den Vergleichsbestand an."""
        self.untersuchung = Untersuchung.objects.create(
            studienname="Ausgangsstudie",
            studienname_vergleich="ausgangsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Alte Lehrperson",
        )

        self.untersuchung_id = self.untersuchung.pk
        self.uploadzeitpunkt = self.untersuchung.uploadzeitpunkt

        # Der vollständige Formularsatz ändert ausschließlich den Stundenwert.
        self.daten = {
            "untersuchung_id": str(self.untersuchung_id),
            "studienname": "Ausgangsstudie",
            "stundenwert": "2.5",
            "einordnung": Untersuchung.PRAESENZ,
            "lehrperson": "Alte Lehrperson",
        }

        # Alle gespeicherten Feldwerte als festen Vergleichsstand sichern.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )

    def melde_testkonto_an(self, rolle):
        """Legt ein synthetisches Rollenkonto an und richtet seine Sitzung per Anmeldung ein."""
        testpasswort = "a" * 64
        konto, fehlerliste = lege_konto_an(
            "testkonto",
            testpasswort,
            rolle,
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        anmeldung = self.client.post(
            "/anmelden/",
            data={
                "benutzername": "testkonto",
                "passwort": testpasswort,
            },
        )
        self.assertEqual(anmeldung.status_code, 200)
        self.assertTrue(anmeldung.json()["erfolg"])

    def pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche gespeicherten Feldwerte mit dem vorbereiteten Bestand."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )

    def test_fa15_einzelner_stundenwert_wird_korrigiert(self):
        """Eine gültige Anfrage ändert ausschließlich den Stundenwert in Antwort und Datenbank."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self.client.post(
            "/untersuchung_korrigieren/",
            data=self.daten,
        )
        self.assertEqual(antwort.status_code, 200)

        ergebnis = antwort.json()
        self.assertTrue(ergebnis["erfolg"])
        zeile = ergebnis["untersuchung"]

        self.assertEqual(zeile["untersuchung_id"], self.untersuchung_id)
        self.assertEqual(zeile["studienname"], "Ausgangsstudie")
        self.assertEqual(zeile["stundenwert"], 2.5)
        self.assertEqual(zeile["einordnung"], Untersuchung.PRAESENZ)
        self.assertEqual(zeile["lehrperson"], "Alte Lehrperson")
        self.assertEqual(
            datetime.fromisoformat(zeile["uploadzeitpunkt"]),
            self.uploadzeitpunkt,
        )

        # Neu aus der Datenbank laden, damit die Antwort allein nicht als Speichernachweis gilt.
        self.untersuchung.refresh_from_db()
        self.assertEqual(self.untersuchung.pk, self.untersuchung_id)
        self.assertEqual(self.untersuchung.studienname, "Ausgangsstudie")
        self.assertEqual(
            self.untersuchung.studienname_vergleich,
            "ausgangsstudie",
        )
        self.assertEqual(self.untersuchung.stundenwert, 2.5)
        self.assertEqual(
            self.untersuchung.einordnung,
            Untersuchung.PRAESENZ,
        )
        self.assertEqual(self.untersuchung.lehrperson, "Alte Lehrperson")
        self.assertEqual(
            self.untersuchung.uploadzeitpunkt,
            self.uploadzeitpunkt,
        )

    def test_qa04_korrektur_ohne_anmeldung_wird_abgewiesen(self):
        """Eine nicht angemeldete Korrektur wird ohne Datenänderung abgewiesen."""
        antwort = self.client.post(
            "/untersuchung_korrigieren/",
            data=self.daten,
        )

        self.assertEqual(antwort.status_code, 403)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("untersuchung", ergebnis)
        self.pruefe_unveraenderten_bestand()

    def test_qa04_lehrperson_kann_keine_untersuchung_korrigieren(self):
        """Eine Lehrperson darf keine gespeicherte Untersuchung korrigieren."""
        self.melde_testkonto_an(Benutzerkonto.LEHRPERSON)

        antwort = self.client.post(
            "/untersuchung_korrigieren/",
            data=self.daten,
        )

        self.assertEqual(antwort.status_code, 403)
        ergebnis = antwort.json()
        self.assertFalse(ergebnis["erfolg"])
        self.assertNotIn("untersuchung", ergebnis)
        self.pruefe_unveraenderten_bestand()

    def test_fa15_get_anfrage_wird_abgewiesen(self):
        """Eine GET-Korrektur wird abgewiesen und weist POST als erlaubte Methode aus."""
        antwort = self.client.get(
            "/untersuchung_korrigieren/",
            data=self.daten,
        )

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort.headers["Allow"], "POST")
        self.assertFalse(antwort.json()["erfolg"])
        self.pruefe_unveraenderten_bestand()

    def test_fa15_fehlende_parameter_werden_abgewiesen(self):
        """Jeder fehlende Pflichtparameter wird vor der Korrekturfunktion gemeldet."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        for schluessel, feldname in [
            ("untersuchung_id", "Untersuchung"),
            ("studienname", "Studienname"),
            ("stundenwert", "Stundenwert"),
            ("einordnung", "Einordnung"),
            ("lehrperson", "Lehrperson"),
        ]:
            with self.subTest(parameter=schluessel):
                daten = self.daten.copy()
                daten.pop(schluessel)

                # Der Mock prüft, dass fehlende Angaben vor der eigentlichen Korrektur abgefangen werden.
                with patch(
                    "vp_stunden.views.korrigiere_untersuchung"
                ) as korrekturfunktion:
                    antwort = self.client.post(
                        "/untersuchung_korrigieren/",
                        data=daten,
                    )

                self.assertEqual(antwort.status_code, 400)
                ergebnis = antwort.json()
                self.assertFalse(ergebnis["erfolg"])
                self.assertTrue(ergebnis["fehlerliste"])
                self.assertIn(
                    feldname,
                    " ".join(ergebnis["fehlerliste"]),
                )
                korrekturfunktion.assert_not_called()
                self.pruefe_unveraenderten_bestand()

    def test_fa15_ungueltige_zahlen_werden_abgewiesen(self):
        """Nicht lesbare Zahlenwerte werden vor der Korrekturfunktion feldbezogen gemeldet."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        for schluessel, feldname in [
            ("untersuchung_id", "Untersuchung"),
            ("stundenwert", "Stundenwert"),
        ]:
            with self.subTest(parameter=schluessel):
                # Eine Kopie hält die Teilfälle voneinander unabhängig.
                daten = self.daten.copy()
                daten[schluessel] = "keine Zahl"

                with patch(
                    "vp_stunden.views.korrigiere_untersuchung"
                ) as korrekturfunktion:
                    antwort = self.client.post(
                        "/untersuchung_korrigieren/",
                        data=daten,
                    )

                self.assertEqual(antwort.status_code, 400)
                ergebnis = antwort.json()
                self.assertFalse(ergebnis["erfolg"])
                self.assertTrue(ergebnis["fehlerliste"])
                self.assertIn(
                    feldname,
                    " ".join(ergebnis["fehlerliste"]),
                )
                korrekturfunktion.assert_not_called()
                self.pruefe_unveraenderten_bestand()

    def test_fa15_inhaltliche_fehler_werden_zurueckgegeben(self):
        """Inhaltlich ungültige Angaben liefern Feldfehler und lassen den Bestand unverändert."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Lesbare Zahlen und vorhandene Felder erreichen die nachgelagerten Inhaltsprüfungen.
        for schluessel, wert, feldname in [
            ("studienname", "", "Studienname"),
            ("stundenwert", "1.25", "Stundenwert"),
        ]:
            with self.subTest(parameter=schluessel):
                daten = self.daten.copy()
                daten[schluessel] = wert

                antwort = self.client.post(
                    "/untersuchung_korrigieren/",
                    data=daten,
                )

                self.assertEqual(antwort.status_code, 400)
                ergebnis = antwort.json()
                self.assertFalse(ergebnis["erfolg"])
                self.assertTrue(ergebnis["fehlerliste"])
                self.assertIn(
                    feldname,
                    " ".join(ergebnis["fehlerliste"]),
                )
                self.assertNotIn("untersuchung", ergebnis)
                self.pruefe_unveraenderten_bestand()


class ReporterzeugungViewTests(TestCase):
    """Prüft PDF-Antworten, Reportvoraussetzungen und Fehlerbehandlung."""

    def setUp(self):
        """Legt einen erfüllten Reportbestand mit abweichenden Namens- und Teilnahmefolgen an."""
        self.matrikelnummer = "1234567"
        self.daten = {"matrikelnummer": self.matrikelnummer}

        # Namensreihenfolge und spätere Teilnahmereihenfolge bewusst gegeneinander setzen.
        self.untersuchung_alpha = Untersuchung.objects.create(
            studienname="Alpha-Studie",
            studienname_vergleich="alpha-studie",
            stundenwert=15.0,
            einordnung=Untersuchung.ONLINE,
            lehrperson="Testlehrperson Alpha",
        )
        self.untersuchung_beta = Untersuchung.objects.create(
            studienname="Beta-Studie",
            studienname_vergleich="beta-studie",
            stundenwert=15.0,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Beta",
        )

        # Die Anlagereihenfolge der Teilnahmen bestimmt die erwartete Reihenfolge im Report.
        self.teilnahme_beta = Teilnahme.objects.create(
            untersuchung=self.untersuchung_beta,
            matrikelnummer=self.matrikelnummer,
        )
        self.teilnahme_alpha = Teilnahme.objects.create(
            untersuchung=self.untersuchung_alpha,
            matrikelnummer=self.matrikelnummer,
        )

        # Die fremde Teilnahme darf weder Reportinhalt noch Stundensummen beeinflussen.
        Teilnahme.objects.create(
            untersuchung=self.untersuchung_alpha,
            matrikelnummer="7654321",
        )

        # Alle Feldwerte festhalten, da die Reporterzeugung keine Daten ändern darf.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

    def melde_testkonto_an(self, rolle):
        """Legt ein synthetisches Rollenkonto an und richtet seine Sitzung per Anmeldung ein."""
        testpasswort = "a" * 64
        konto, fehlerliste = lege_konto_an(
            "testkonto",
            testpasswort,
            rolle,
        )
        self.assertIsNotNone(konto)
        self.assertFalse(fehlerliste)

        # Die reguläre Anmeldung richtet zugleich die erforderlichen Sitzungsdaten ein.
        anmeldung = self.client.post(
            "/anmelden/",
            data={
                "benutzername": "testkonto",
                "passwort": testpasswort,
            },
        )
        self.assertEqual(anmeldung.status_code, 200)
        self.assertTrue(anmeldung.json()["erfolg"])

    def pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche gespeicherten Feldwerte mit dem vorbereiteten Bestand."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def pruefe_abweisung(self, daten, status):
        """Prüft HTTP-Abweisung, ausgebliebene PDF-Erzeugung und unveränderte Daten."""
        # Der Mock prüft zusätzlich zum HTTP-Status, dass keine PDF-Erzeugung beginnt.
        with patch(
            "vp_stunden.views.erzeuge_report"
        ) as reportfunktion:
            antwort = self.client.post(
                "/report_erzeugen/",
                data=daten,
            )

        self.assertEqual(antwort.status_code, status)
        self.assertFalse(antwort.json()["erfolg"])
        reportfunktion.assert_not_called()
        self.pruefe_unveraenderten_bestand()

    def test_fa17_administrator_erhaelt_pdf_in_importreihenfolge(self):
        """Die Reportanfrage liefert PDF-Daten und übergibt passende Teilnahmen in Importreihenfolge."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # wraps führt die echte Reporterzeugung aus und zeichnet ihre Aufrufargumente auf.
        with patch(
            "vp_stunden.views.erzeuge_report",
            wraps=erzeuge_report,
        ) as reportfunktion:
            antwort = self.client.post(
                "/report_erzeugen/",
                data=self.daten,
            )

        self.assertEqual(antwort.status_code, 200)
        self.assertEqual(
            antwort.headers["Content-Type"],
            "application/pdf",
        )
        self.assertEqual(
            antwort.headers["Content-Disposition"],
            'attachment; filename="vp_stunden_1234567.pdf"',
        )

        # PDF-Markierungen prüfen das Ausgabeformat, nicht den vollständigen Dokumentinhalt.
        self.assertTrue(antwort.content.startswith(b"%PDF-"))
        self.assertTrue(antwort.content.rstrip().endswith(b"%%EOF"))

        reportfunktion.assert_called_once()
        matrikelnummer, teilnahmen, berechnung = (
            reportfunktion.call_args.args
        )

        self.assertEqual(matrikelnummer, self.matrikelnummer)

        # Die übergebenen IDs prüfen Auswahl und Importreihenfolge unabhängig von den Namen.
        teilnahme_ids = [teilnahme.pk for teilnahme in teilnahmen]
        self.assertEqual(
            teilnahme_ids,
            [self.teilnahme_beta.pk, self.teilnahme_alpha.pk],
        )

        self.assertEqual(
            berechnung,
            {
                "praesenzstunden": 15.0,
                "onlinestunden": 15.0,
                "gesamtstunden": 30.0,
                "erfuellt": True,
            },
        )
        self.pruefe_unveraenderten_bestand()

    def test_qa04_report_ohne_anmeldung_wird_abgewiesen(self):
        """Eine nicht angemeldete Reportanfrage wird ohne PDF-Erzeugung abgewiesen."""
        self.pruefe_abweisung(self.daten, 403)

    def test_qa04_lehrperson_kann_keinen_report_erzeugen(self):
        """Eine Lehrperson darf keinen Report erzeugen."""
        self.melde_testkonto_an(Benutzerkonto.LEHRPERSON)
        self.pruefe_abweisung(self.daten, 403)

    def test_fa17_get_anfrage_wird_abgewiesen(self):
        """Eine GET-Reportanfrage wird ohne PDF-Erzeugung abgewiesen."""
        with patch(
            "vp_stunden.views.erzeuge_report"
        ) as reportfunktion:
            antwort = self.client.get(
                "/report_erzeugen/",
                data=self.daten,
            )

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort.headers["Allow"], "POST")
        self.assertFalse(antwort.json()["erfolg"])
        reportfunktion.assert_not_called()
        self.pruefe_unveraenderten_bestand()

    def test_fa17_fehlende_matrikelnummer_wird_abgewiesen(self):
        """Eine fehlende oder leere Matrikelnummer verhindert die Reportverarbeitung."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        for daten in [{}, {"matrikelnummer": ""}]:
            with self.subTest(daten=daten):
                self.pruefe_abweisung(daten, 400)

    def test_fa17_fehlende_teilnahmen_werden_gemeldet(self):
        """Eine Matrikelnummer ohne Teilnahmen erhält keinen Report."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        self.pruefe_abweisung(
            {"matrikelnummer": "9999999"},
            404,
        )

    def test_fa17_nicht_erfuellter_status_verhindert_report(self):
        """Zu wenige Gesamt- oder Präsenzstunden verhindern jeweils die Reporterzeugung."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        # Jeder Teilfall unterschreitet genau eine der beiden Mindestgrenzen.
        for online, praesenz in [(14.5, 15.0), (15.5, 14.5)]:
            with self.subTest(online=online, praesenz=praesenz):
                self.untersuchung_alpha.stundenwert = online
                self.untersuchung_alpha.save(
                    update_fields=["stundenwert"]
                )
                self.untersuchung_beta.stundenwert = praesenz
                self.untersuchung_beta.save(
                    update_fields=["stundenwert"]
                )

                # Die vorbereiteten Änderungen gehören zum neuen Vergleichsstand.
                self.untersuchungsbestand = list(
                    Untersuchung.objects.order_by("pk").values()
                )

                self.pruefe_abweisung(self.daten, 400)

    def test_fa17_erzeugungsfehler_wird_kontrolliert_beantwortet(self):
        """Ein Erzeugungsfehler wird ohne interne Fehlerdetails und Datenänderungen beantwortet."""
        self.melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        interner_fehler = "INTERNER_TESTFEHLER"

        # Den Erzeugungsfehler gezielt auslösen und seine internen Details im Antworttext ausschließen.
        with patch(
            "vp_stunden.views.erzeuge_report",
            side_effect=RuntimeError(interner_fehler),
        ) as reportfunktion:
            antwort = self.client.post(
                "/report_erzeugen/",
                data=self.daten,
            )

        self.assertEqual(antwort.status_code, 500)
        self.assertFalse(antwort.json()["erfolg"])
        self.assertNotIn(
            interner_fehler,
            antwort.content.decode("utf-8"),
        )
        reportfunktion.assert_called_once()
        self.pruefe_unveraenderten_bestand()


class UebergangsimportZuordnungTests(TestCase):
    """Prüft Zuordnung, Datenübernahme und vollständige Abweisung."""

    def setUp(self):
        """Bereitet Altbestand und eine gültige Übergangsdatei für jeden Test neu vor."""
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="7654321",
        )

        # list(values()) hält alle Feldwerte sofort fest, auch IDs und Zeitpunkt.
        # Spätere Vergleiche erkennen damit Änderungen trotz gleicher Datensatzanzahl.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        # Die XLSX-Datei entsteht im Speicher. addCleanup schließt die Arbeitsmappe.
        self.arbeitsmappe = Workbook()
        self.addCleanup(self.arbeitsmappe.close)
        self.blatt = self.arbeitsmappe.active
        self.blatt.title = "Übergangsimport"
        self.blatt.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])

        # Abweichende Schreibweise und äußere Leerzeichen dürfen die Zuordnung
        # zur vorhandenen Untersuchung in Zeile 2 nicht verhindern.
        self.blatt.append([
            1234567,
            "  BESTANDSSTUDIE  ",
            1.5,
            Untersuchung.PRAESENZ,
            "  TESTLEHRPERSON ALT  ",
        ])

        # Die Leerzeile darf die Prüfung der folgenden Zeile nicht beenden.
        self.blatt.append([None, None, None, None, None])
        self.blatt.append([
            1234567,
            "  Neue Studie  ",
            2.5,
            Untersuchung.ONLINE,
            "  Testlehrperson Neu  ",
        ])

    def test_fa11_neue_und_vorhandene_untersuchung_werden_importiert(self):
        """Ordnet vorhandene Untersuchungen zu und speichert neue Untersuchungen mit Teilnahmen."""
        # BytesIO sammelt XLSX-Bytes im Speicher, SimpleUploadedFile bildet den Upload nach.
        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        beginn = timezone.now()
        fehlerliste = importiere_uebergangsdatei(datei)
        ende = timezone.now()

        self.assertEqual(fehlerliste, [])

        neue_untersuchung = Untersuchung.objects.get(
            studienname_vergleich="neue studie"
        )
        self.assertEqual(
            (
                neue_untersuchung.studienname,
                neue_untersuchung.stundenwert,
                neue_untersuchung.einordnung,
                neue_untersuchung.lehrperson,
            ),
            ("Neue Studie", 2.5, Untersuchung.ONLINE, "Testlehrperson Neu"),
        )
        self.assertGreaterEqual(neue_untersuchung.uploadzeitpunkt, beginn)
        self.assertLessEqual(neue_untersuchung.uploadzeitpunkt, ende)

        # Der Altbestand muss einschließlich Schreibweise und Zeitpunkt unverändert bleiben.
        self.assertEqual(
            list(
                Untersuchung.objects.exclude(
                    pk=neue_untersuchung.pk
                ).order_by("pk").values()
            ),
            self.untersuchungsbestand,
        )

        # Die Untersuchungs-IDs prüfen Zuordnung und Anlage in der Reihenfolge der Zeilen.
        self.assertEqual(
            list(
                Teilnahme.objects.filter(
                    matrikelnummer="1234567"
                ).order_by("pk").values_list("untersuchung_id", flat=True)
            ),
            [self.untersuchung.pk, neue_untersuchung.pk],
        )
        self.assertEqual(
            list(
                Teilnahme.objects.exclude(
                    matrikelnummer="1234567"
                ).order_by("pk").values()
            ),
            self.teilnahmebestand,
        )

    def test_fa11_unterschiedliche_matrikelnummern_werden_abgewiesen(self):
        """Weist verschiedene Matrikelnummern ohne Änderung am Bestand ab."""
        self.blatt["A4"] = 2345678

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        # Die Fundstelle muss vorkommen, der übrige Meldungstext bleibt frei.
        self.assertTrue(
            any(meldung.startswith("A4:") for meldung in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa12_doppelte_untersuchung_in_datei_wird_abgewiesen(self):
        """Weist doppelte Studiennamen trotz anderer Schreibweise ohne Teilimport ab."""
        # Auch mit anderer Schreibweise ist Zeile 5 eine Dublette von Zeile 4.
        self.blatt.append([
            1234567,
            "  NEUE STUDIE  ",
            2.5,
            Untersuchung.ONLINE,
            "Testlehrperson Neu",
        ])

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        self.assertTrue(
            any(meldung.startswith("B5:") for meldung in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa12_bereits_gespeicherte_teilnahme_wird_abgewiesen(self):
        """Weist eine bereits gespeicherte Teilnahme ab und erhält den gesamten Bestand."""
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="1234567",
        )

        # Diese zusätzliche Teilnahme gehört jetzt zum vorbereiteten Bestand.
        teilnahmebestand_vor_import = list(
            Teilnahme.objects.order_by("pk").values()
        )

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        self.assertTrue(
            any(meldung.startswith("Zeile 2:") for meldung in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            teilnahmebestand_vor_import,
        )

    def test_fa12_abweichende_bestandsangaben_werden_abgewiesen(self):
        """Weist jede Abweichung von gemeinsamen Bestandsangaben ohne Datenänderung ab."""
        # Jeden sonst gültigen Wert einzeln verändern, damit kein anderer Fehler
        # eine fehlende Prüfung des Bestandsvergleichs verdeckt.
        for adresse, abweichender_wert in [
            ("C2", 2.0),
            ("D2", Untersuchung.ONLINE),
            ("E2", "Andere Testlehrperson"),
        ]:
            with self.subTest(zelle=adresse):
                urspruenglicher_wert = self.blatt[adresse].value
                self.blatt[adresse] = abweichender_wert

                # Der vorherige Upload ist geschlossen, jeder Teilfall braucht einen neuen.
                with BytesIO() as puffer:
                    self.arbeitsmappe.save(puffer)
                    datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

                fehlerliste = importiere_uebergangsdatei(datei)

                # subTest setzt Änderungen nicht zurück: den Ausgangswert selbst wiederherstellen.
                self.blatt[adresse] = urspruenglicher_wert

                self.assertTrue(
                    any(
                        meldung.startswith(f"{adresse}:")
                        for meldung in fehlerliste
                    ),
                    fehlerliste,
                )
                self.assertEqual(
                    list(Untersuchung.objects.order_by("pk").values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.order_by("pk").values()),
                    self.teilnahmebestand,
                )

    def test_fa11_zellfehler_werden_gesammelt_und_datei_abgewiesen(self):
        """Meldet mehrere Zellfehler gemeinsam und speichert auch gültige Zeilen nicht."""
        # Zwei unabhängige Fehler in Zeile 4 dürfen sich nicht gegenseitig verdecken.
        # Auch die gültige Zeile 2 darf bei der Abweisung nicht gespeichert werden.
        self.blatt["C4"] = 0.25
        self.blatt["E4"] = None

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        for fundstelle in ["C4:", "E4:"]:
            self.assertTrue(
                any(meldung.startswith(fundstelle) for meldung in fehlerliste),
                fehlerliste,
            )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa11_speicherfehler_nimmt_alle_aenderungen_zurueck(self):
        """Nimmt bei einem Speicherfehler neue Untersuchungen und Teilnahmen vollständig zurück."""
        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        # wraps verwendet die echte Speichermethode, DEFAULT lässt den ersten Aufruf durch.
        # Der zweite scheitert nach Anlage der neuen Untersuchung. Die Transaktion muss
        # diese Untersuchung und die zuvor gespeicherte Teilnahme zurücknehmen.
        with patch(
            "vp_stunden.importe.Teilnahme.objects.create",
            wraps=Teilnahme.objects.create,
            side_effect=[DEFAULT, DatabaseError("INTERNER_TESTFEHLER")],
        ) as teilnahmespeicherung:
            fehlerliste = importiere_uebergangsdatei(datei)

        self.assertEqual(teilnahmespeicherung.call_count, 2)
        self.assertEqual(
            fehlerliste,
            ["Datei: Der Übergangsimport konnte nicht übernommen werden."],
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )


class RegulaererImportTests(TestCase):
    """Prüft Datenübernahme und vollständige Abweisung beim regulären Import."""

    def setUp(self):
        """Bereitet Altbestand und eine gültige reguläre Importdatei für jeden Test neu vor."""
        # Der synthetische Altbestand entsteht vor jedem Test in der Testdatenbank neu.
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="9876543",
        )

        # Alle Feldwerte sofort festhalten, um auch Änderungen bei gleicher Anzahl zu erkennen.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        # Die XLSX-Datei entsteht im Speicher. addCleanup schließt die Arbeitsmappe.
        self.arbeitsmappe = Workbook()
        self.addCleanup(self.arbeitsmappe.close)
        self.blatt = self.arbeitsmappe.active
        self.blatt.title = "Untersuchung"

        self.blatt["A1"] = "Studienname"
        self.blatt["A2"] = "Stundenwert (VP-Stunden)"
        self.blatt["A3"] = "Einordnung (P/O)"
        self.blatt["A4"] = "Verantwortliche Lehrperson"
        self.blatt["A6"] = "Matrikelnummer"

        self.blatt["B1"] = "  Neue Studie  "
        self.blatt["B2"] = 2.5
        self.blatt["B3"] = Untersuchung.ONLINE
        self.blatt["B4"] = "  Testlehrperson Neu  "

        # Numerische Excel-Werte müssen als Zeichenfolgen gespeichert werden.
        # A106 prüft zugleich, dass die große Lücke den Import nicht beendet.
        self.blatt["A7"] = 1234567
        self.blatt["A106"] = 7654321

    def test_fa10_untersuchung_und_teilnahmen_werden_gespeichert(self):
        """Speichert eine neue Untersuchung mit allen Teilnahmen und erhält den Altbestand."""
        # BytesIO sammelt XLSX-Bytes im Speicher, SimpleUploadedFile bildet den Upload nach.
        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())

        beginn = timezone.now()
        fehlerliste = importiere_untersuchung(datei)
        ende = timezone.now()

        self.assertEqual(fehlerliste, [])

        neue_untersuchung = Untersuchung.objects.get(
            studienname_vergleich="neue studie"
        )
        self.assertEqual(
            (
                neue_untersuchung.studienname,
                neue_untersuchung.stundenwert,
                neue_untersuchung.einordnung,
                neue_untersuchung.lehrperson,
            ),
            ("Neue Studie", 2.5, Untersuchung.ONLINE, "Testlehrperson Neu"),
        )
        self.assertGreaterEqual(neue_untersuchung.uploadzeitpunkt, beginn)
        self.assertLessEqual(neue_untersuchung.uploadzeitpunkt, ende)

        # Die neue Untersuchung herausnehmen und sämtliche alten Feldwerte vergleichen.
        self.assertEqual(
            list(
                Untersuchung.objects.exclude(
                    pk=neue_untersuchung.pk
                ).order_by("pk").values()
            ),
            self.untersuchungsbestand,
        )

        # Beide Nummern müssen derselben Untersuchung zugeordnet sein, auch die aus A106.
        self.assertEqual(
            list(
                Teilnahme.objects.filter(
                    untersuchung=neue_untersuchung
                ).order_by("pk").values_list("matrikelnummer", flat=True)
            ),
            ["1234567", "7654321"],
        )
        self.assertEqual(
            list(
                Teilnahme.objects.exclude(
                    untersuchung=neue_untersuchung
                ).order_by("pk").values()
            ),
            self.teilnahmebestand,
        )

    def test_fa08_doppelte_matrikelnummer_wird_abgewiesen(self):
        """Weist eine wiederholte Matrikelnummer ohne Datenübernahme ab."""
        self.blatt["A106"] = 1234567

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())

        fehlerliste = importiere_untersuchung(datei)

        # Mindestens eine Meldung muss die verursachende Zelle nennen.
        self.assertTrue(
            any(meldung.startswith("A106:") for meldung in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa08_bereits_gespeicherte_untersuchung_wird_abgewiesen(self):
        """Erkennt vorhandene Studiennamen trotz anderer Schreibweise und weist den Import ab."""
        # Andere Großschreibung und äußere Leerzeichen dürfen den Bestandskonflikt nicht verdecken.
        self.blatt["B1"] = "  BESTANDSSTUDIE  "

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())

        fehlerliste = importiere_untersuchung(datei)

        self.assertTrue(
            any(meldung.startswith("B1:") for meldung in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa09_unabhaengige_zellfehler_werden_gesammelt(self):
        """Sammelt Fehler aller Eingabefelder und erhält den vollständigen Bestand."""
        # Die Feldprüfungen sind separat getestet. Hier muss der Import alle Fehler
        # sammeln, auch wenn bereits der Studienname ungültig ist.
        self.blatt["B1"] = "   "
        self.blatt["B2"] = 0.25
        self.blatt["B3"] = "p"
        self.blatt["B4"] = None
        self.blatt["A106"] = "7654321"

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())

        fehlerliste = importiere_untersuchung(datei)

        for fundstelle in ["B1:", "B2:", "B3:", "B4:", "A106:"]:
            self.assertTrue(
                any(meldung.startswith(fundstelle) for meldung in fehlerliste),
                fehlerliste,
            )

        # Auch die weiterhin gültige Nummer in A7 darf nicht gespeichert werden.
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa06_ohne_matrikelnummer_wird_datei_abgewiesen(self):
        """Weist eine Datei ohne Teilnahmen ab und erhält den Bestand."""
        self.blatt["A7"] = None
        self.blatt["A106"] = None

        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())

        fehlerliste = importiere_untersuchung(datei)

        self.assertTrue(
            any(meldung.startswith("A7:A106:") for meldung in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa09_speicherfehler_nimmt_alle_aenderungen_zurueck(self):
        """Nimmt Untersuchung und erste Teilnahme nach einem späteren Speicherfehler zurück."""
        with BytesIO() as puffer:
            self.arbeitsmappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())

        # wraps verwendet die echte Speichermethode, DEFAULT lässt den ersten Aufruf durch.
        # Der zweite scheitert. Die Transaktion muss Untersuchung und erste Teilnahme zurücknehmen.
        with patch(
            "vp_stunden.importe.Teilnahme.objects.create",
            wraps=Teilnahme.objects.create,
            side_effect=[DEFAULT, DatabaseError("INTERNER_TESTFEHLER")],
        ) as teilnahmespeicherung:
            fehlerliste = importiere_untersuchung(datei)

        self.assertEqual(teilnahmespeicherung.call_count, 2)
        self.assertEqual(
            fehlerliste,
            ["Datei: Der Import konnte nicht übernommen werden."],
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )


class ImportVorlagenstrukturTests(TestCase):
    """Prüft Vorlagenaufbau und Teilnahmezeilengrenzen nach FA-06/FA-11."""

    def setUp(self):
        """Bereitet Altbestand und beide gültigen Vorlagen für jeden Test neu vor."""
        # Der Altbestand macht unbeabsichtigte Änderungen bei einer Abweisung sichtbar.
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="9876543",
        )
        # Vollständige Feldwerte vergleichen, damit unveränderte Anzahlen allein nicht genügen.
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        # Gültige Vorlagen mit je einem Eintrag bilden die Grundlage der gezielten Änderungen.
        # Arbeitsmappen und XLSX-Uploads bleiben im Speicher, addCleanup schließt die Mappen.
        self.regulaere_mappe = Workbook()
        self.addCleanup(self.regulaere_mappe.close)
        self.regulaeres_blatt = self.regulaere_mappe.active
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "Neue reguläre Studie",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
        }.items():
            self.regulaeres_blatt[adresse] = wert

        self.uebergangsmappe = Workbook()
        self.addCleanup(self.uebergangsmappe.close)
        self.uebergangsblatt = self.uebergangsmappe.active
        self.uebergangsblatt.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        self.uebergangsblatt.append([
            1234567,
            "Übergangsstudie 1",
            0.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Neu",
        ])

    def test_fa06_regulaere_beschriftungen_muessen_exakt_stimmen(self):
        """Weist jede fehlende oder abweichende reguläre Beschriftung ohne Datenänderung ab."""
        for adresse in ["A1", "A2", "A3", "A4", "A6"]:
            original = self.regulaeres_blatt[adresse].value

            # Auch ein zusätzliches Leerzeichen verletzt die exakte Beschriftung.
            for falsch in [None, original + " "]:
                with self.subTest(zelle=adresse, inhalt=falsch):
                    self.regulaeres_blatt[adresse] = falsch
                    with BytesIO() as puffer:
                        self.regulaere_mappe.save(puffer)
                        datei = SimpleUploadedFile("regulaer.xlsx", puffer.getvalue())

                    fehlerliste = importiere_untersuchung(datei)

                    # subTest setzt nichts zurück: den Ausgangswert vor den Assertions wiederherstellen.
                    self.regulaeres_blatt[adresse] = original
                    self.assertTrue(
                        any(m.startswith(f"{adresse}:") for m in fehlerliste),
                        fehlerliste,
                    )
                    self.assertEqual(
                        list(Untersuchung.objects.order_by("pk").values()),
                        self.untersuchungsbestand,
                    )
                    self.assertEqual(
                        list(Teilnahme.objects.order_by("pk").values()),
                        self.teilnahmebestand,
                    )

    def test_fa11_uebergangsbeschriftungen_muessen_exakt_stimmen(self):
        """Weist jede fehlende oder abweichende Übergangsbeschriftung ohne Datenänderung ab."""
        for adresse in ["A1", "B1", "C1", "D1", "E1"]:
            original = self.uebergangsblatt[adresse].value

            for falsch in [None, original + " "]:
                with self.subTest(zelle=adresse, inhalt=falsch):
                    self.uebergangsblatt[adresse] = falsch
                    with BytesIO() as puffer:
                        self.uebergangsmappe.save(puffer)
                        datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

                    fehlerliste = importiere_uebergangsdatei(datei)

                    # Den Ausgangswert vor den Assertions für den nächsten Teilfall wiederherstellen.
                    self.uebergangsblatt[adresse] = original
                    self.assertTrue(
                        any(m.startswith(f"{adresse}:") for m in fehlerliste),
                        fehlerliste,
                    )
                    self.assertEqual(
                        list(Untersuchung.objects.order_by("pk").values()),
                        self.untersuchungsbestand,
                    )
                    self.assertEqual(
                        list(Teilnahme.objects.order_by("pk").values()),
                        self.teilnahmebestand,
                    )

    def test_fa06_regulaer_unzulaessige_zellinhalte_werden_abgewiesen(self):
        """Weist zusätzliche Inhalte und scheinbar leere Matrikelnummernzellen vollständig ab."""
        # Auch Leerzeichen, 0 und False sind außerhalb der Eingabezellen unzulässiger Inhalt.
        # Innerhalb des Bereichs ist eine Zelle mit Leerzeichen ebenfalls nicht leer.
        for adresse, wert in [
            ("A5", "Zusatz"),
            ("B6", "Zusatz"),
            ("B7", "Zusatz"),
            ("C1", "Zusatz"),
            ("C1", "   "),
            ("C1", 0),
            ("C1", False),
            ("A8", "   "),
        ]:
            with self.subTest(zelle=adresse, inhalt=wert):
                self.regulaeres_blatt[adresse] = wert
                with BytesIO() as puffer:
                    self.regulaere_mappe.save(puffer)
                    datei = SimpleUploadedFile("regulaer.xlsx", puffer.getvalue())

                fehlerliste = importiere_untersuchung(datei)

                # Den leeren Ausgangszustand für den nächsten Teilfall wiederherstellen.
                self.regulaeres_blatt[adresse] = None
                self.assertTrue(
                    any(m.startswith(f"{adresse}:") for m in fehlerliste),
                    fehlerliste,
                )
                self.assertEqual(
                    list(Untersuchung.objects.order_by("pk").values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.order_by("pk").values()),
                    self.teilnahmebestand,
                )

    def test_fa11_uebergang_unzulaessige_zellinhalte_werden_abgewiesen(self):
        """Weist zusätzliche Inhalte und unvollständige Übergangszeilen vollständig ab."""
        # Spalte F liegt außerhalb der Vorlage. Leerzeichen in B101 erzeugen dagegen
        # eine unvollständige Eingabezeile, die auch hinter der Lücke geprüft werden muss.
        for adresse, wert in [
            ("F1", "Zusatz"),
            ("F2", "Zusatz"),
            ("F2", "   "),
            ("F2", 0),
            ("F2", False),
            ("B101", "   "),
        ]:
            with self.subTest(zelle=adresse, inhalt=wert):
                self.uebergangsblatt[adresse] = wert
                with BytesIO() as puffer:
                    self.uebergangsmappe.save(puffer)
                    datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

                fehlerliste = importiere_uebergangsdatei(datei)

                # Den leeren Ausgangszustand für den nächsten Teilfall wiederherstellen.
                self.uebergangsblatt[adresse] = None
                self.assertTrue(
                    any(m.startswith(f"{adresse}:") for m in fehlerliste),
                    fehlerliste,
                )
                self.assertEqual(
                    list(Untersuchung.objects.order_by("pk").values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.order_by("pk").values()),
                    self.teilnahmebestand,
                )

    def test_fa06_regulaer_genau_100_matrikelnummern_sind_zulaessig(self):
        """Übernimmt genau 100 Nummern trotz freiem Blattnamen und formatierter Leerzellen."""
        # Freier Blattname und reine Formatierung außerhalb des Eingabebereichs sind zulässig.
        self.regulaeres_blatt.title = "Anderer Blattname"
        for adresse in ["A107", "B7", "Z200"]:
            self.regulaeres_blatt[adresse].number_format = "0.00"

        # Genau die zulässigen Zeilen A7 bis A106 mit unterschiedlichen Nummern füllen.
        for nummer in range(100):
            self.regulaeres_blatt.cell(
                row=nummer + 7, column=1, value=1_000_000 + nummer
            )

        with BytesIO() as puffer:
            self.regulaere_mappe.save(puffer)
            datei = SimpleUploadedFile("regulaer.xlsx", puffer.getvalue())

        fehlerliste = importiere_untersuchung(datei)

        self.assertEqual(fehlerliste, [])
        neue_untersuchung = Untersuchung.objects.get(
            studienname_vergleich="neue reguläre studie"
        )
        self.assertEqual(
            list(
                Teilnahme.objects.filter(
                    untersuchung=neue_untersuchung
                ).order_by("pk").values_list("matrikelnummer", flat=True)
            ),
            [str(1_000_000 + nummer) for nummer in range(100)],
        )
        self.assertEqual(
            list(
                Untersuchung.objects.exclude(
                    pk=neue_untersuchung.pk
                ).order_by("pk").values()
            ),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(
                Teilnahme.objects.exclude(
                    untersuchung=neue_untersuchung
                ).order_by("pk").values()
            ),
            self.teilnahmebestand,
        )

    def test_fa11_uebergang_genau_100_teilnahmezeilen_sind_zulaessig(self):
        """Übernimmt genau 100 Übergangszeilen trotz freiem Blattnamen und formatierter Leerzellen."""
        # Auch hier dürfen freier Blattname und reine Formatierung den Import nicht verhindern.
        self.uebergangsblatt.title = "Anderer Blattname"
        for adresse in ["A102", "F2", "Z200"]:
            self.uebergangsblatt[adresse].number_format = "0.00"

        # Zeilen 2 bis 101: unterschiedliche Untersuchungen für dieselbe Matrikelnummer.
        for nummer in range(1, 101):
            zeile = nummer + 1
            self.uebergangsblatt.cell(zeile, 1, 1234567)
            self.uebergangsblatt.cell(zeile, 2, f"Übergangsstudie {nummer}")
            self.uebergangsblatt.cell(zeile, 3, 0.5)
            self.uebergangsblatt.cell(zeile, 4, Untersuchung.PRAESENZ)
            self.uebergangsblatt.cell(zeile, 5, "Testlehrperson Neu")

        with BytesIO() as puffer:
            self.uebergangsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        self.assertEqual(fehlerliste, [])
        self.assertEqual(Untersuchung.objects.count(), 101)

        # untersuchung__studienname folgt der Beziehung und prüft jede einzelne Zuordnung.
        self.assertEqual(
            list(
                Teilnahme.objects.filter(
                    matrikelnummer="1234567"
                ).order_by("pk").values_list("untersuchung__studienname", flat=True)
            ),
            [f"Übergangsstudie {nummer}" for nummer in range(1, 101)],
        )
        self.assertEqual(
            list(
                Untersuchung.objects.filter(
                    pk=self.untersuchung.pk
                ).order_by("pk").values()
            ),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(
                Teilnahme.objects.exclude(
                    matrikelnummer="1234567"
                ).order_by("pk").values()
            ),
            self.teilnahmebestand,
        )

    def test_fa06_regulaer_101_matrikelnummern_werden_abgewiesen(self):
        """Weist die 101. Matrikelnummer ab und speichert keine der gültigen Nummern."""
        # Alle Nummern sind gültig und verschieden, nur A107 liegt außerhalb des Bereichs.
        for nummer in range(101):
            self.regulaeres_blatt.cell(
                row=nummer + 7, column=1, value=1_000_000 + nummer
            )

        with BytesIO() as puffer:
            self.regulaere_mappe.save(puffer)
            datei = SimpleUploadedFile("regulaer.xlsx", puffer.getvalue())

        fehlerliste = importiere_untersuchung(datei)

        self.assertTrue(
            any(m.startswith("A107:") for m in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa11_uebergang_101_teilnahmezeilen_werden_abgewiesen(self):
        """Weist die 101. Teilnahmezeile ab und speichert keine der gültigen Zeilen."""
        # Alle 101 Teilnahmen sind sonst gültig, nur Zeile 102 überschreitet den Bereich.
        for nummer in range(1, 102):
            zeile = nummer + 1
            self.uebergangsblatt.cell(zeile, 1, 1234567)
            self.uebergangsblatt.cell(zeile, 2, f"Übergangsstudie {nummer}")
            self.uebergangsblatt.cell(zeile, 3, 0.5)
            self.uebergangsblatt.cell(zeile, 4, Untersuchung.PRAESENZ)
            self.uebergangsblatt.cell(zeile, 5, "Testlehrperson Neu")

        with BytesIO() as puffer:
            self.uebergangsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        self.assertTrue(
            any(m.startswith("A102:") for m in fehlerliste),
            fehlerliste,
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa11_uebergang_ohne_teilnahmeangaben_wird_abgewiesen(self):
        """Weist eine vollständig leere Übergangsliste ohne Änderung am Bestand ab."""
        # Nur die Kopfzeile bleibt: Einzelne Leerzeilen sind zulässig, eine leere Liste nicht.
        for spalte in range(1, 6):
            self.uebergangsblatt.cell(row=2, column=spalte).value = None

        with BytesIO() as puffer:
            self.uebergangsmappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())

        fehlerliste = importiere_uebergangsdatei(datei)

        self.assertEqual(
            fehlerliste,
            ["Datei: Es sind keine Teilnahmeangaben vorhanden."],
        )
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )


class ImportDateiUndBlattTests(TestCase):
    """Prüft Dateiformat, Dateigröße und zusätzliche Blätter beider Importwege."""

    def setUp(self):
        """Bereitet Altbestand und zwei gültige Vorlagen mit verschiedenen Studien vor."""
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )
        self.xlsx_typ = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        regulaere_mappe = Workbook()
        self.addCleanup(regulaere_mappe.close)
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "Reguläre Dateiprüfung",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
        }.items():
            regulaere_mappe.active[adresse] = wert

        uebergangsmappe = Workbook()
        self.addCleanup(uebergangsmappe.close)
        uebergangsmappe.active.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        uebergangsmappe.active.append([
            1234567,
            "Übergangsdateiprüfung",
            0.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Neu",
        ])

        # Jeder Teilfall wird über die zugehörige öffentliche Importfunktion geprüft.
        # Verschiedene Studiennamen verhindern Kollisionen bei zwei Erfolgen im selben Test.
        self.importe = [
            ("regulaer", regulaere_mappe, importiere_untersuchung),
            ("uebergang", uebergangsmappe, importiere_uebergangsdatei),
        ]
        self.grundgroessen = {}

        # Für echte Byte-Grenzfälle wird die XLSX ohne ZIP-Kompression geschrieben.
        # Ein vorhandenes Beschreibungszeichen lässt sich später bytegenau verlängern.
        for weg, mappe, _ in self.importe:
            mappe.properties.description = "x"
            with BytesIO() as puffer:
                with ZipFile(puffer, "w", compression=ZIP_STORED) as archiv:
                    ExcelWriter(mappe, archiv).save()
                self.grundgroessen[weg] = len(puffer.getvalue())

    def test_qa05_falsche_dateiendung_wird_abgewiesen(self):
        """Eine gültige XLSX-Datei wird unter einer unzulässigen Endung abgewiesen."""
        for weg, mappe, importiere in self.importe:
            for name in ["daten.xls", "daten.xlsx.exe"]:
                with self.subTest(importweg=weg, dateiname=name):
                    with BytesIO() as puffer:
                        mappe.save(puffer)
                        datei = SimpleUploadedFile(
                            name, puffer.getvalue(), content_type=self.xlsx_typ
                        )

                    fehlerliste = importiere(datei)

                    self.assertEqual(
                        fehlerliste,
                        ["Datei: Eine makrofreie .xlsx-Datei ist erforderlich."],
                    )
                    self.assertTrue(datei.closed)
                    self.assertEqual(
                        list(Untersuchung.objects.order_by("pk").values()),
                        self.untersuchungsbestand,
                    )
                    self.assertEqual(
                        list(Teilnahme.objects.order_by("pk").values()),
                        self.teilnahmebestand,
                    )

    def test_qa05_umbenannte_excel_vorlage_wird_abgewiesen(self):
        """Eine als XLSX benannte XLTX-Vorlage wird anhand ihres tatsächlichen Formats erkannt."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                # template=True erzeugt den Inhaltstyp XLTX, obwohl der Upload XLSX heißt.
                mappe.template = True
                with BytesIO() as puffer:
                    mappe.save(puffer)
                    datei = SimpleUploadedFile(
                        "daten.xlsx", puffer.getvalue(), content_type=self.xlsx_typ
                    )
                mappe.template = False

                fehlerliste = importiere(datei)

                self.assertEqual(
                    fehlerliste,
                    ["Datei: Eine makrofreie XLSX-Arbeitsmappe ist erforderlich."],
                )
                self.assertTrue(datei.closed)
                self.assertEqual(
                    list(Untersuchung.objects.order_by("pk").values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.order_by("pk").values()),
                    self.teilnahmebestand,
                )

    def test_qa05_unlesbare_dateien_werden_kontrolliert_abgewiesen(self):
        """Leere, fremde und beschädigte Inhalte liefern nur die allgemeine Dateimeldung."""
        for weg, mappe, importiere in self.importe:
            with BytesIO() as puffer:
                mappe.save(puffer)
                gueltige_bytes = puffer.getvalue()

            for fall, inhalt in [
                ("leer", b""),
                ("Text statt XLSX", b"Keine Excel-Arbeitsmappe"),
                ("abgeschnitten", gueltige_bytes[:len(gueltige_bytes) // 2]),
            ]:
                with self.subTest(importweg=weg, fall=fall):
                    # Passender Dateiname und Medientyp ersetzen keine Inhaltsprüfung.
                    datei = SimpleUploadedFile(
                        "daten.xlsx", inhalt, content_type=self.xlsx_typ
                    )

                    fehlerliste = importiere(datei)

                    self.assertEqual(
                        fehlerliste,
                        ["Datei: Die XLSX-Datei konnte nicht gelesen werden."],
                    )
                    self.assertTrue(datei.closed)
                    self.assertEqual(
                        list(Untersuchung.objects.order_by("pk").values()),
                        self.untersuchungsbestand,
                    )
                    self.assertEqual(
                        list(Teilnahme.objects.order_by("pk").values()),
                        self.teilnahmebestand,
                    )

    def test_fa05_fa11_genau_2000000_byte_sind_zulaessig(self):
        """Die echte Dateigrößengrenze bleibt mit XLSX-Endung und unspezifischem Medientyp zulässig."""
        zielgroesse = 2_000_000
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                # Nur Beschreibungstext außerhalb der Eingabezellen verlängern.
                # +1 berücksichtigt das Zeichen, das bereits in der Grundgröße enthalten ist.
                mappe.properties.description = "x" * (
                    zielgroesse - self.grundgroessen[weg] + 1
                )
                with BytesIO() as puffer:
                    with ZipFile(puffer, "w", compression=ZIP_STORED) as archiv:
                        ExcelWriter(mappe, archiv).save()
                    inhalt = puffer.getvalue()

                # Die tatsächlichen Bytes prüfen, nicht einen überschriebenen size-Wert.
                self.assertEqual(len(inhalt), zielgroesse)
                datei = SimpleUploadedFile(
                    "daten.XLSX", inhalt, content_type="application/octet-stream"
                )

                fehlerliste = importiere(datei)

                self.assertEqual(fehlerliste, [])
                self.assertTrue(datei.closed)

        # Beide Importe müssen ihre Teilnahme tatsächlich gespeichert haben.
        self.assertEqual(Untersuchung.objects.count(), 3)
        self.assertEqual(
            list(
                Teilnahme.objects.filter(
                    matrikelnummer="1234567"
                ).order_by("pk").values_list("untersuchung__studienname", flat=True)
            ),
            ["Reguläre Dateiprüfung", "Übergangsdateiprüfung"],
        )
        self.assertEqual(
            list(
                Untersuchung.objects.filter(
                    pk=self.untersuchung.pk
                ).order_by("pk").values()
            ),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(
                Teilnahme.objects.exclude(
                    matrikelnummer="1234567"
                ).order_by("pk").values()
            ),
            self.teilnahmebestand,
        )

    def test_fa05_fa11_2000001_byte_werden_abgewiesen(self):
        """Eine ansonsten gültige XLSX-Datei ein Byte über der Grenze verändert keine Daten."""
        zielgroesse = 2_000_001
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                mappe.properties.description = "x" * (
                    zielgroesse - self.grundgroessen[weg] + 1
                )
                with BytesIO() as puffer:
                    with ZipFile(puffer, "w", compression=ZIP_STORED) as archiv:
                        ExcelWriter(mappe, archiv).save()
                    inhalt = puffer.getvalue()

                self.assertEqual(len(inhalt), zielgroesse)
                datei = SimpleUploadedFile(
                    "daten.xlsx", inhalt, content_type=self.xlsx_typ
                )

                fehlerliste = importiere(datei)

                self.assertEqual(
                    fehlerliste,
                    ["Datei: Höchstens 2.000.000 Byte sind zulässig."],
                )
                self.assertTrue(datei.closed)
                self.assertEqual(
                    list(Untersuchung.objects.order_by("pk").values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.order_by("pk").values()),
                    self.teilnahmebestand,
                )

    def test_fa05_fa11_zusaetzliche_blaetter_werden_abgewiesen(self):
        """Ein zusätzliches Blatt führt unabhängig von seiner Sichtbarkeit zur Abweisung."""
        for weg, mappe, importiere in self.importe:
            for zustand in ["visible", "hidden", "veryHidden"]:
                with self.subTest(importweg=weg, blattzustand=zustand):
                    zusatzblatt = mappe.create_sheet("Zusatz")
                    zusatzblatt.sheet_state = zustand
                    with BytesIO() as puffer:
                        mappe.save(puffer)
                        datei = SimpleUploadedFile(
                            "daten.xlsx", puffer.getvalue(), content_type=self.xlsx_typ
                        )

                    # Die Arbeitsmappe vor dem nächsten Teilfall wieder auf ein Blatt setzen.
                    mappe.remove(zusatzblatt)
                    fehlerliste = importiere(datei)

                    self.assertEqual(
                        fehlerliste,
                        ["Datei: Genau ein Tabellenblatt ohne weitere Blätter ist erforderlich."],
                    )
                    self.assertTrue(datei.closed)
                    self.assertEqual(
                        list(Untersuchung.objects.order_by("pk").values()),
                        self.untersuchungsbestand,
                    )
                    self.assertEqual(
                        list(Teilnahme.objects.order_by("pk").values()),
                        self.teilnahmebestand,
                    )


class ImportZipXmlTests(TestCase):
    """Prüft die ZIP-Grenzen und das geschützte XML-Einlesen beider Importwege."""

    def setUp(self):
        """Bereitet einen Altbestand und zwei gültige XLSX-Dateien im Speicher vor."""
        untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        regulaere_mappe = Workbook()
        self.addCleanup(regulaere_mappe.close)
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "ZIP-Prüfung regulär",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
        }.items():
            regulaere_mappe.active[adresse] = wert

        uebergangsmappe = Workbook()
        self.addCleanup(uebergangsmappe.close)
        uebergangsmappe.active.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        uebergangsmappe.active.append([
            1234567,
            "ZIP-Prüfung Übergang",
            0.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Neu",
        ])

        self.importe = []
        for weg, mappe, importiere, studienname in [
            (
                "regulaer", regulaere_mappe,
                importiere_untersuchung, "ZIP-Prüfung regulär",
            ),
            (
                "uebergang", uebergangsmappe,
                importiere_uebergangsdatei, "ZIP-Prüfung Übergang",
            ),
        ]:
            # Die Beschreibung liegt außerhalb der Eingabezellen. Ihr kurzer
            # Marker lässt sich später ohne zusätzliche XML-Elemente verlängern.
            mappe.properties.description = "ZIP_GROESSEN_TEST"
            with BytesIO() as puffer:
                mappe.save(puffer)
                self.importe.append(
                    (weg, puffer.getvalue(), importiere, studienname)
                )

    def _packe_xlsx(self, inhalt, ersetzungen=None, zusatzanzahl=0):
        """Erzeugt ein neues ZIP-Paket mit echten Inhalten und neu berechneten Prüfsummen."""
        ersetzungen = ersetzungen or {}
        with BytesIO(inhalt) as quelle, BytesIO() as ziel:
            with ZipFile(quelle) as original:
                with ZipFile(ziel, "w", compression=ZIP_DEFLATED) as neu:
                    for eintrag in original.infolist():
                        daten = ersetzungen.get(
                            eintrag.filename, original.read(eintrag.filename)
                        )
                        neu.writestr(eintrag.filename, daten)

                    # Auch leere Verzeichniseinträge zählen zur ZIP-Grenze.
                    # Sie ändern weder Vorlage noch Untersuchungsangaben.
                    for nummer in range(zusatzanzahl):
                        neu.writestr(f"testzusatz/{nummer}/", b"")
            return ziel.getvalue()

    def _pruefe_abweisung(self, inhalt, importiere, meldung):
        """Prüft die Dateimeldung, den geschlossenen Upload und den vollständigen Altbestand."""
        # Die Uploadgrenze darf die hier untersuchte ZIP-/XML-Prüfung nicht verdecken.
        self.assertLessEqual(len(inhalt), 2_000_000)
        datei = SimpleUploadedFile("daten.xlsx", inhalt)
        self.addCleanup(datei.close)

        fehlerliste = importiere(datei)

        self.assertEqual(fehlerliste, [meldung])
        self.assertTrue(datei.closed)
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def _pruefe_erfolg(self, inhalt, importiere, studienname):
        """Prüft die tatsächliche Übernahme genau einer Untersuchung und ihrer Teilnahme."""
        vorherige_untersuchungen = list(
            Untersuchung.objects.order_by("pk").values()
        )
        vorherige_teilnahmen = list(
            Teilnahme.objects.order_by("pk").values()
        )
        self.assertLessEqual(len(inhalt), 2_000_000)
        datei = SimpleUploadedFile("daten.xlsx", inhalt)
        self.addCleanup(datei.close)

        fehlerliste = importiere(datei)

        self.assertEqual(fehlerliste, [])
        self.assertTrue(datei.closed)
        neue_untersuchung = Untersuchung.objects.get(studienname=studienname)
        self.assertEqual(
            list(
                Teilnahme.objects.filter(
                    untersuchung=neue_untersuchung
                ).values_list("matrikelnummer", flat=True)
            ),
            ["1234567"],
        )
        self.assertEqual(
            list(
                Untersuchung.objects.exclude(
                    pk=neue_untersuchung.pk
                ).order_by("pk").values()
            ),
            vorherige_untersuchungen,
        )
        self.assertEqual(
            list(
                Teilnahme.objects.exclude(
                    untersuchung=neue_untersuchung
                ).order_by("pk").values()
            ),
            vorherige_teilnahmen,
        )

    def test_qa05_genau_100_zip_eintraege_sind_zulaessig(self):
        """Eine gültige XLSX-Datei mit genau 100 ZIP-Einträgen wird übernommen."""
        for weg, inhalt, importiere, studienname in self.importe:
            with self.subTest(importweg=weg):
                with ZipFile(BytesIO(inhalt)) as archiv:
                    grundanzahl = len(archiv.infolist())
                self.assertLess(grundanzahl, 100)
                inhalt = self._packe_xlsx(
                    inhalt, zusatzanzahl=100 - grundanzahl
                )

                with ZipFile(BytesIO(inhalt)) as archiv:
                    self.assertEqual(len(archiv.infolist()), 100)
                    self.assertIsNone(archiv.testzip())

                self._pruefe_erfolg(inhalt, importiere, studienname)

    def test_qa05_101_zip_eintraege_werden_abgewiesen(self):
        """Ein zusätzlicher ZIP-Eintrag führt vor der Datenübernahme zur Abweisung."""
        for weg, inhalt, importiere, _ in self.importe:
            with self.subTest(importweg=weg):
                with ZipFile(BytesIO(inhalt)) as archiv:
                    grundanzahl = len(archiv.infolist())
                self.assertLess(grundanzahl, 101)
                inhalt = self._packe_xlsx(
                    inhalt, zusatzanzahl=101 - grundanzahl
                )

                with ZipFile(BytesIO(inhalt)) as archiv:
                    self.assertEqual(len(archiv.infolist()), 101)
                    self.assertIsNone(archiv.testzip())

                self._pruefe_abweisung(
                    inhalt, importiere,
                    "Datei: Höchstens 100 ZIP-Einträge sind zulässig.",
                )

    def test_qa05_genau_20_mib_entpackt_sind_zulaessig(self):
        """Die tatsächliche entpackte Gesamtgröße darf genau 20 MiB betragen."""
        zielgroesse = 20 * 1024 * 1024
        for weg, inhalt, importiere, studienname in self.importe:
            with self.subTest(importweg=weg):
                with ZipFile(BytesIO(inhalt)) as archiv:
                    grundgroesse = sum(
                        len(archiv.read(name)) for name in archiv.namelist()
                    )
                    beschreibung = archiv.read("docProps/core.xml")

                # Komprimierbarer Beschreibungstext erhöht die tatsächlichen
                # ZIP-Inhalte, während die Uploaddatei deutlich unter 2 MB bleibt.
                # Die ZIP-Größenangaben selbst werden nicht manipuliert.
                self.assertEqual(beschreibung.count(b"ZIP_GROESSEN_TEST"), 1)
                beschreibung = beschreibung.replace(
                    b"ZIP_GROESSEN_TEST",
                    b"ZIP_GROESSEN_TEST" + b"x" * (zielgroesse - grundgroesse),
                    1,
                )
                inhalt = self._packe_xlsx(
                    inhalt, {"docProps/core.xml": beschreibung}
                )
                with ZipFile(BytesIO(inhalt)) as archiv:
                    # read() liest echte Bytes und kontrolliert zugleich die CRC.
                    self.assertEqual(
                        sum(len(archiv.read(name)) for name in archiv.namelist()),
                        zielgroesse,
                    )

                self._pruefe_erfolg(inhalt, importiere, studienname)

    def test_qa05_20_mib_plus_ein_byte_werden_abgewiesen(self):
        """Ein Byte über der entpackten ZIP-Grenze verhindert jede Datenänderung."""
        zielgroesse = 20 * 1024 * 1024 + 1
        for weg, inhalt, importiere, _ in self.importe:
            with self.subTest(importweg=weg):
                with ZipFile(BytesIO(inhalt)) as archiv:
                    grundgroesse = sum(
                        len(archiv.read(name)) for name in archiv.namelist()
                    )
                    beschreibung = archiv.read("docProps/core.xml")

                self.assertEqual(beschreibung.count(b"ZIP_GROESSEN_TEST"), 1)
                beschreibung = beschreibung.replace(
                    b"ZIP_GROESSEN_TEST",
                    b"ZIP_GROESSEN_TEST" + b"x" * (zielgroesse - grundgroesse),
                    1,
                )
                inhalt = self._packe_xlsx(
                    inhalt, {"docProps/core.xml": beschreibung}
                )
                with ZipFile(BytesIO(inhalt)) as archiv:
                    self.assertEqual(
                        sum(len(archiv.read(name)) for name in archiv.namelist()),
                        zielgroesse,
                    )

                self._pruefe_abweisung(
                    inhalt, importiere,
                    "Datei: Die entpackte Gesamtgröße überschreitet 20 MiB.",
                )

    def test_qa05_xml_entitaet_im_manifest_wird_abgewiesen(self):
        """Eine harmlose Entity-Deklaration im Paketmanifest genügt zur sicheren Abweisung."""
        for weg, inhalt, importiere, _ in self.importe:
            with self.subTest(importweg=weg):
                with ZipFile(BytesIO(inhalt)) as archiv:
                    manifest = archiv.read("[Content_Types].xml")

                # Die Entität wird nur deklariert, nicht verwendet oder verschachtelt.
                # Es gibt weder externe Adressen noch eine Textvervielfachung.
                self.assertIn(b"<Types", manifest)
                manifest = manifest.replace(
                    b"<Types",
                    b'<!DOCTYPE Types [<!ENTITY testwert "sicher">]><Types',
                    1,
                )
                with self.assertRaises(EntitiesForbidden):
                    geschuetztes_xml(manifest)
                inhalt = self._packe_xlsx(
                    inhalt, {"[Content_Types].xml": manifest}
                )

                self._pruefe_abweisung(
                    inhalt, importiere,
                    "Datei: Die XLSX-Datei konnte nicht gelesen werden.",
                )

    def test_qa05_xml_entitaet_im_tabellenblatt_wird_abgewiesen(self):
        """Auch das Einlesen der Zellen weist Entity-Deklarationen sicher zurück."""
        for weg, inhalt, importiere, _ in self.importe:
            with self.subTest(importweg=weg):
                with ZipFile(BytesIO(inhalt)) as archiv:
                    blatt_xml = archiv.read("xl/worksheets/sheet1.xml")

                # Das unveränderte Manifest passiert die erste XML-Prüfung.
                # Die Entity-Deklaration erreicht deshalb den Worksheet-Parser.
                self.assertIn(b"<worksheet", blatt_xml)
                blatt_xml = blatt_xml.replace(
                    b"<worksheet",
                    b'<!DOCTYPE worksheet [<!ENTITY testwert "sicher">]><worksheet',
                    1,
                )
                with self.assertRaises(EntitiesForbidden):
                    geschuetztes_xml(blatt_xml)
                inhalt = self._packe_xlsx(
                    inhalt, {"xl/worksheets/sheet1.xml": blatt_xml}
                )

                self._pruefe_abweisung(
                    inhalt, importiere,
                    "Datei: Die XLSX-Datei konnte nicht gelesen werden.",
                )

    def test_qa05_ohne_xml_schutz_wird_nicht_eingelesen(self):
        """Ein fehlender XML-Schutz wird gemeldet, statt ungeschützt weiterzulesen."""
        for weg, inhalt, importiere, _ in self.importe:
            with self.subTest(importweg=weg):
                # Nur die Schutzverfügbarkeit simulieren. Der ungeeignete Zustand
                # darf keinen Aufruf der eigentlichen Ladefunktion auslösen.
                with patch("vp_stunden.importe.DEFUSEDXML", False):
                    with patch("vp_stunden.importe.load_workbook") as lade_mappe:
                        self._pruefe_abweisung(
                            inhalt, importiere,
                            "Datei: Sicheres XML-Einlesen ist nicht verfügbar.",
                        )
                        lade_mappe.assert_not_called()


class ImportMakroVerknuepfungsTests(TestCase):
    """Prüft makrofreie Pakete und die Abweisung externer Verknüpfungen beider Importe."""

    def setUp(self):
        """Bereitet Altbestand und je eine gültige Vorlage für beide Importwege vor."""
        untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        regulaere_mappe = Workbook()
        self.addCleanup(regulaere_mappe.close)
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "Reguläre Makroprüfung",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
        }.items():
            regulaere_mappe.active[adresse] = wert

        uebergangsmappe = Workbook()
        self.addCleanup(uebergangsmappe.close)
        uebergangsmappe.active.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        uebergangsmappe.active.append([
            1234567,
            "Übergangsmakroprüfung",
            0.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Neu",
        ])
        self.importe = [
            ("regulaer", regulaere_mappe, importiere_untersuchung),
            ("uebergang", uebergangsmappe, importiere_uebergangsdatei),
        ]

    def _xlsx_bytes(self, mappe):
        """Schreibt die vorbereitete Arbeitsmappe als XLSX-Bytes im Arbeitsspeicher."""
        with BytesIO() as puffer:
            mappe.save(puffer)
            return puffer.getvalue()

    def _paket_aendern(self, inhalt, ersetzungen):
        """Ersetzt oder ergänzt Paketbestandteile und berechnet ihre ZIP-Prüfsummen neu."""
        with BytesIO(inhalt) as quelle, BytesIO() as ziel:
            with ZipFile(quelle) as alt, ZipFile(
                ziel, "w", compression=ZIP_DEFLATED
            ) as neu:
                for name in alt.namelist():
                    daten = ersetzungen.get(name)
                    if daten is None:
                        daten = alt.read(name)
                    neu.writestr(name, daten)
                for name, daten in ersetzungen.items():
                    if name not in alt.namelist():
                        neu.writestr(name, daten)
            return ziel.getvalue()

    def _pruefe_abweisung(self, importiere, inhalt, meldung):
        """Prüft die konkrete Dateimeldung, geschlossenen Upload und vollständigen Bestand."""
        datei = SimpleUploadedFile("daten.xlsx", inhalt)
        self.addCleanup(datei.close)

        fehlerliste = importiere(datei)

        self.assertEqual(fehlerliste, [meldung])
        self.assertTrue(datei.closed)
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa05_fa11_umbenannte_makroformate_werden_abgewiesen(self):
        """XLSM- und XLTM-Haupttypen werden auch unter dem Dateinamen daten.xlsx erkannt."""
        for weg, mappe, importiere in self.importe:
            grundinhalt = self._xlsx_bytes(mappe)
            for haupttyp in [XLSM, XLTM]:
                with self.subTest(importweg=weg, haupttyp=haupttyp):
                    with ZipFile(BytesIO(grundinhalt)) as archiv:
                        manifest = ElementTree.fromstring(
                            archiv.read("[Content_Types].xml")
                        )

                    # Nur den tatsächlichen Haupttyp ändern, keine Makros erzeugen.
                    for element in manifest:
                        if element.get("ContentType") == XLSX:
                            element.set("ContentType", haupttyp)
                    inhalt = self._paket_aendern(grundinhalt, {
                        "[Content_Types].xml": ElementTree.tostring(manifest),
                    })

                    self._pruefe_abweisung(
                        importiere,
                        inhalt,
                        "Datei: Eine makrofreie XLSX-Arbeitsmappe ist erforderlich.",
                    )

    def test_fa05_fa11_vba_deklarationen_werden_abgewiesen(self):
        """VBA-Deklarationen werden sowohl als Standardtyp als auch als Einzeltyp abgewiesen."""
        for weg, mappe, importiere in self.importe:
            grundinhalt = self._xlsx_bytes(mappe)
            for deklaration in ["Default", "Override"]:
                with self.subTest(importweg=weg, deklaration=deklaration):
                    with ZipFile(BytesIO(grundinhalt)) as archiv:
                        manifest = ElementTree.fromstring(
                            archiv.read("[Content_Types].xml")
                        )
                        beziehungen = ElementTree.fromstring(
                            archiv.read("xl/_rels/workbook.xml.rels")
                        )

                    # Default erklärt den Typ einer Dateiendung, Override
                    # den Typ eines einzelnen Bestandteils im Paket.
                    attribute = {"ContentType": VBA}
                    if deklaration == "Default":
                        attribute["Extension"] = "bin"
                    else:
                        attribute["PartName"] = "/xl/vbaProject.bin"
                    ElementTree.SubElement(
                        manifest, f"{{{CONTYPES_NS}}}{deklaration}", attribute
                    )
                    # Die Beziehung ordnet den synthetischen Marker der Mappe zu.
                    ElementTree.SubElement(
                        beziehungen,
                        f"{{{PKG_REL_NS}}}Relationship",
                        {
                            "Id": "rIdSynthetischesVba",
                            "Type": (
                                "http://schemas.microsoft.com/office/2006/"
                                "relationships/vbaProject"
                            ),
                            "Target": "vbaProject.bin",
                        },
                    )
                    # Der XLSX-Haupttyp bleibt erhalten. Der binäre Bestandteil
                    # enthält nur eine Markierung, keinen ausführbaren VBA-Code.
                    inhalt = self._paket_aendern(grundinhalt, {
                        "[Content_Types].xml": ElementTree.tostring(manifest),
                        "xl/_rels/workbook.xml.rels": ElementTree.tostring(beziehungen),
                        "xl/vbaProject.bin": b"SYNTHETISCHER VBA-TESTMARKER",
                    })

                    self._pruefe_abweisung(
                        importiere,
                        inhalt,
                        "Datei: Makrobestandteile sind nicht zulässig.",
                    )

    def test_fa05_fa11_makroblatt_deklaration_wird_abgewiesen(self):
        """Ein deklarierter Makroblattbestandteil wird trotz unverändertem XLSX-Haupttyp abgewiesen."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                grundinhalt = self._xlsx_bytes(mappe)
                with ZipFile(BytesIO(grundinhalt)) as archiv:
                    manifest = ElementTree.fromstring(
                        archiv.read("[Content_Types].xml")
                    )

                ElementTree.SubElement(
                    manifest,
                    f"{{{CONTYPES_NS}}}Override",
                    {
                        "PartName": "/xl/macrosheets/sheet1.xml",
                        "ContentType": "application/vnd.ms-excel.macrosheet+xml",
                    },
                )
                # Ein leerer XML-Marker genügt: Entscheidend ist die Deklaration.
                # Es werden keine Makrobefehle oder ausführbaren Inhalte angelegt.
                inhalt = self._paket_aendern(grundinhalt, {
                    "[Content_Types].xml": ElementTree.tostring(manifest),
                    "xl/macrosheets/sheet1.xml": (
                        b'<worksheet xmlns="http://schemas.openxmlformats.org/'
                        b'spreadsheetml/2006/main"><sheetData/></worksheet>'
                    ),
                })

                self._pruefe_abweisung(
                    importiere,
                    inhalt,
                    "Datei: Makrobestandteile sind nicht zulässig.",
                )

    def test_qa05_externe_arbeitsmappenverknuepfungen_werden_abgewiesen(self):
        """Echte externe Verweise auf synthetische Web- und Dateiziele führen zur Abweisung."""
        for weg, mappe, importiere in self.importe:
            for ziel in [
                "https://arbeitsmappe.invalid/synthetische-quelle.xlsx",
                "file:///synthetische-testquelle.xlsx",
            ]:
                with self.subTest(importweg=weg, ziel=ziel):
                    verknuepfung = ExternalLink(externalBook=ExternalBook(
                        sheetNames=ExternalSheetNames(sheetName=["Synthetische Quelle"]),
                        id="rId1",
                    ))
                    verknuepfung.file_link = Relationship(
                        Id="rId1",
                        type="externalLinkPath",
                        Target=ziel,
                        TargetMode="External",
                    )
                    # Die private openpyxl-3.1.5-Liste wird mit echten Metadaten
                    # gefüllt, nicht durch eine vorgetäuschte Ladeantwort ersetzt.
                    mappe._external_links = [verknuepfung]
                    inhalt = self._xlsx_bytes(mappe)
                    mappe._external_links = []

                    with BytesIO(inhalt) as puffer:
                        geladen = load_workbook(puffer, keep_links=True)
                        try:
                            self.assertEqual(len(geladen._external_links), 1)
                            self.assertEqual(
                                geladen._external_links[0].file_link.Target, ziel
                            )
                        finally:
                            geladen.close()

                    # openpyxl liest die Verweismetadaten aus dem ZIP-Paket.
                    # Die Zielarbeitsmappe wird dabei weder geöffnet noch geladen.
                    self._pruefe_abweisung(
                        importiere,
                        inhalt,
                        "Datei: Externe Arbeitsmappenverknüpfungen sind nicht zulässig.",
                    )


class ImportFormelTests(TestCase):
    """Prüft die vollständige Abweisung echter Excel-Formeln in beiden Importwegen."""

    def setUp(self):
        """Bereitet Altbestand und gültige Vorlagen mit je einer Teilnahme vor."""
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        regulaere_mappe = Workbook()
        self.addCleanup(regulaere_mappe.close)
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "Formelprüfung regulär",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
        }.items():
            regulaere_mappe.active[adresse] = wert

        uebergangsmappe = Workbook()
        self.addCleanup(uebergangsmappe.close)
        uebergangsmappe.active.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        uebergangsmappe.active.append([
            1234567,
            "Formelprüfung Übergang",
            0.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Neu",
        ])
        self.importe = [
            ("regulaer", regulaere_mappe, importiere_untersuchung),
            ("uebergang", uebergangsmappe, importiere_uebergangsdatei),
        ]

    def _upload(self, mappe):
        """Erzeugt für jeden Import einen eigenen XLSX-Upload im Arbeitsspeicher."""
        with BytesIO() as puffer:
            mappe.save(puffer)
            datei = SimpleUploadedFile("daten.xlsx", puffer.getvalue())
        self.addCleanup(datei.close)
        return datei

    def _pruefe_abweisung(self, mappe, importiere, erwartete_meldungen):
        """Prüft gesammelte Zellmeldungen und den vollständigen unveränderten Bestand."""
        datei = self._upload(mappe)

        fehlerliste = importiere(datei)

        # Die Reihenfolge der vorhandenen Zellen ist kein Erfolgskriterium.
        # Jede erwartete Meldung muss aber genau einmal enthalten sein.
        self.assertCountEqual(fehlerliste, erwartete_meldungen)
        self.assertTrue(datei.closed)
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_qa05_formeln_in_eingabezellen_werden_gesammelt_abgewiesen(self):
        """Erkennt Formeln in allen fünf Feldarten, auch nach Lücken am Bereichsende."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                if weg == "regulaer":
                    formeln = {
                        "B1": '="Formelstudie"',
                        "B2": "=1/2",
                        "B3": '="P"',
                        "B4": '="Testlehrperson Neu"',
                        "A106": "=1234567",
                    }
                else:
                    formeln = {
                        "A101": "=1234567",
                        "B101": '="Formelstudie"',
                        "C101": "=1/2",
                        "D101": '="P"',
                        "E101": '="Testlehrperson Neu"',
                    }

                # Die Ausdrücke würden passende Feldwerte ergeben. Abgewiesen
                # wird die Formel selbst, unabhängig von ihrem möglichen Ergebnis.
                for adresse, formel in formeln.items():
                    mappe.active[adresse] = formel
                    self.assertEqual(mappe.active[adresse].data_type, "f")

                self._pruefe_abweisung(mappe, importiere, [
                    f"{adresse}: Formeln sind nicht zulässig."
                    for adresse in formeln
                ])

    def test_qa05_formeln_in_beschriftungen_werden_abgewiesen(self):
        """Auch feste Vorlagenbeschriftungen müssen Zellwerte ohne Formeln enthalten."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                adressen = (
                    ["A1", "A2", "A3", "A4", "A6"]
                    if weg == "regulaer"
                    else ["A1", "B1", "C1", "D1", "E1"]
                )
                for adresse in adressen:
                    beschriftung = mappe.active[adresse].value
                    mappe.active[adresse] = f'="{beschriftung}"'

                self._pruefe_abweisung(mappe, importiere, [
                    f"{adresse}: Formeln sind nicht zulässig."
                    for adresse in adressen
                ])

    def test_qa05_formeln_ausserhalb_der_vorlage_werden_abgewiesen(self):
        """Die Formelprüfung erfasst auch Zellen neben und unter dem erlaubten Bereich."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                adressen = (
                    ["C5", "A107"] if weg == "regulaer" else ["F2", "A102"]
                )
                erwartete_meldungen = []
                for adresse in adressen:
                    mappe.active[adresse] = "=1+1"
                    # Beide Fehlerursachen gehören zu derselben Zellfundstelle.
                    erwartete_meldungen.extend([
                        f"{adresse}: Formeln sind nicht zulässig.",
                        f"{adresse}: Außerhalb der Vorlage sind keine Inhalte zulässig.",
                    ])

                self._pruefe_abweisung(mappe, importiere, erwartete_meldungen)

    def test_qa05_gleichheitszeichen_in_text_bleibt_zulaessig(self):
        """Ein Gleichheitszeichen innerhalb normalen Textes macht daraus keine Formel."""
        for anzahl, (weg, mappe, importiere) in enumerate(self.importe, start=1):
            with self.subTest(importweg=weg):
                studienname = f"Formelprüfung = {weg}"
                namenszelle = "B1" if weg == "regulaer" else "B2"
                lehrpersonzelle = "B4" if weg == "regulaer" else "E2"
                mappe.active[namenszelle] = studienname
                mappe.active[lehrpersonzelle] = "Testlehrperson = Neu"
                datei = self._upload(mappe)

                fehlerliste = importiere(datei)

                self.assertEqual(fehlerliste, [])
                self.assertTrue(datei.closed)
                untersuchung = Untersuchung.objects.get(studienname=studienname)
                self.assertEqual(untersuchung.lehrperson, "Testlehrperson = Neu")
                self.assertTrue(Teilnahme.objects.filter(
                    untersuchung=untersuchung, matrikelnummer="1234567"
                ).exists())
                self.assertEqual(Untersuchung.objects.count(), 1 + anzahl)
                self.assertEqual(Teilnahme.objects.count(), 1 + anzahl)
                self.assertEqual(
                    list(Untersuchung.objects.filter(
                        pk=self.untersuchung.pk
                    ).order_by("pk").values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.exclude(
                        matrikelnummer="1234567"
                    ).order_by("pk").values()),
                    self.teilnahmebestand,
                )


class ImportDateibereinigungTests(TestCase):
    """Prüft die Entfernung tatsächlicher Django-Uploaddateien nach QA-05/QA-06."""

    def setUp(self):
        """Bereitet Altbestand, Vorlagen und ein begrenztes Uploadverzeichnis vor."""
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        # Später registrierte Aufräumfunktionen laufen zuerst. Die Uploads
        # werden daher geschlossen, bevor ihr Testverzeichnis entfernt wird.
        temporaerverzeichnis = TemporaryDirectory()
        self.addCleanup(temporaerverzeichnis.cleanup)
        self.enterContext(self.settings(FILE_UPLOAD_TEMP_DIR=temporaerverzeichnis.name))

        regulaere_mappe = Workbook()
        self.addCleanup(regulaere_mappe.close)
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "Reguläre Bereinigungsprüfung",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
        }.items():
            regulaere_mappe.active[adresse] = wert

        uebergangsmappe = Workbook()
        self.addCleanup(uebergangsmappe.close)
        uebergangsmappe.active.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        uebergangsmappe.active.append([
            1234567,
            "Übergangsbereinigungsprüfung",
            0.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Neu",
        ])
        self.importe = [
            ("regulaer", regulaere_mappe, importiere_untersuchung),
            ("uebergang", uebergangsmappe, importiere_uebergangsdatei),
        ]

    def _temporaerer_upload(self, mappe, dateiname="daten.xlsx"):
        """Erzeugt einen echten Django-Datenträgerupload aus XLSX-Bytes."""
        with BytesIO() as puffer:
            mappe.save(puffer)
            inhalt = puffer.getvalue()

        # Die Bytes entstehen im Speicher. Nur der zu prüfende Django-Upload
        # liegt auf Datenträger, ohne eine weitere Excel-Zwischenkopie.
        datei = TemporaryUploadedFile(
            dateiname,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            len(inhalt),
            None,
        )
        self.addCleanup(datei.close)
        datei.write(inhalt)
        datei.seek(0)
        pfad = Path(datei.temporary_file_path())
        self.assertTrue(pfad.is_file())
        self.assertFalse(datei.closed)
        return datei, pfad

    def _pruefe_bereinigung(self, datei, pfad):
        """Prüft Schließen und Entfernung vor der allgemeinen Testbereinigung."""
        self.assertTrue(datei.closed)
        self.assertFalse(pfad.exists())

    def _pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche alten Feldwerte und erkennt zusätzliche Datensätze."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_qa06_temporaerer_upload_wird_nach_erfolg_entfernt(self):
        """Beide Importe speichern ihre Daten und entfernen den Datenträgerupload."""
        for anzahl, (weg, mappe, importiere) in enumerate(self.importe, start=1):
            with self.subTest(importweg=weg):
                datei, pfad = self._temporaerer_upload(mappe)

                fehlerliste = importiere(datei)

                self.assertEqual(fehlerliste, [])
                self._pruefe_bereinigung(datei, pfad)
                self.assertEqual(Untersuchung.objects.count(), 1 + anzahl)
                self.assertEqual(Teilnahme.objects.count(), 1 + anzahl)

                # Die neuen Namen unterscheiden sich. Alte Untersuchungswerte
                # und die alte Teilnahme müssen in beiden Teilfällen gleich bleiben.
                self.assertEqual(
                    list(Untersuchung.objects.filter(
                        pk=self.untersuchung.pk
                    ).values()),
                    self.untersuchungsbestand,
                )
                self.assertEqual(
                    list(Teilnahme.objects.exclude(
                        matrikelnummer="1234567"
                    ).order_by("pk").values()),
                    self.teilnahmebestand,
                )

    def test_qa06_temporaerer_upload_wird_nach_frueher_abweisung_entfernt(self):
        """Bereinigt den Upload auch bei einer Abweisung vor dem XLSX-Einlesen."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                datei, pfad = self._temporaerer_upload(mappe, "daten.xls")

                fehlerliste = importiere(datei)

                self.assertEqual(
                    fehlerliste,
                    ["Datei: Eine makrofreie .xlsx-Datei ist erforderlich."],
                )
                self._pruefe_bereinigung(datei, pfad)
                self._pruefe_unveraenderten_bestand()

    def test_qa06_temporaerer_upload_wird_nach_zellfehler_entfernt(self):
        """Bereinigt nach einer Zellabweisung die geladene Uploaddatei ohne Datenänderung."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                adresse = "B2" if weg == "regulaer" else "C2"
                mappe.active[adresse] = 0.25
                datei, pfad = self._temporaerer_upload(mappe)

                fehlerliste = importiere(datei)

                self.assertTrue(
                    any(m.startswith(f"{adresse}:") for m in fehlerliste),
                    fehlerliste,
                )
                self._pruefe_bereinigung(datei, pfad)
                self._pruefe_unveraenderten_bestand()

    def test_qa06_temporaerer_upload_wird_nach_speicherfehler_entfernt(self):
        """Bereinigt den Upload und nimmt eine bereits gespeicherte Untersuchung zurück."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                datei, pfad = self._temporaerer_upload(mappe)

                # wraps lässt die neue Untersuchung tatsächlich speichern.
                # Der Fehler beim Anlegen der Teilnahme muss sie wieder zurücknehmen.
                with patch(
                    "vp_stunden.importe.Untersuchung.objects.create",
                    wraps=Untersuchung.objects.create,
                ) as untersuchungsspeicherung, patch(
                    "vp_stunden.importe.Teilnahme.objects.create",
                    side_effect=DatabaseError("INTERNER_TESTFEHLER"),
                ):
                    fehlerliste = importiere(datei)

                self.assertEqual(untersuchungsspeicherung.call_count, 1)
                bezeichnung = "Der Import" if weg == "regulaer" else "Der Übergangsimport"
                self.assertEqual(
                    fehlerliste,
                    [f"Datei: {bezeichnung} konnte nicht übernommen werden."],
                )
                self._pruefe_bereinigung(datei, pfad)
                self._pruefe_unveraenderten_bestand()

    def test_qa06_temporaerer_upload_wird_nach_unerwartetem_fehler_entfernt(self):
        """Bereinigt nach unerwartetem Verarbeitungsfehler und verbirgt dessen Interna."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                datei, pfad = self._temporaerer_upload(mappe)

                with patch(
                    "vp_stunden.importe.pruefe_studienname",
                    side_effect=RuntimeError("INTERNER_TESTFEHLER"),
                ):
                    fehlerliste = importiere(datei)

                bezeichnung = "Der Import" if weg == "regulaer" else "Der Übergangsimport"
                self.assertEqual(
                    fehlerliste,
                    [f"Datei: {bezeichnung} konnte nicht verarbeitet werden."],
                )
                self._pruefe_bereinigung(datei, pfad)
                self._pruefe_unveraenderten_bestand()

    def test_qa06_upload_schliesst_auch_wenn_arbeitsmappenschliessen_scheitert(self):
        """Prüft die Uploadbereinigung trotz einer weitergegebenen close-Ausnahme."""
        for weg, mappe, importiere in self.importe:
            with self.subTest(importweg=weg):
                adresse = "B2" if weg == "regulaer" else "C2"
                mappe.active[adresse] = 0.25
                datei, pfad = self._temporaerer_upload(mappe)

                # Die ungültige Vorlage verhindert Speicherung. Der zusätzliche
                # close-Fehler verlässt finally, dessen innerer Aufräumschritt
                # muss den Upload dennoch schließen und entfernen.
                with patch(
                    "openpyxl.workbook.workbook.Workbook.close",
                    side_effect=RuntimeError("INTERNER_SCHLIESSFEHLER"),
                ):
                    with self.assertRaisesMessage(RuntimeError, "INTERNER_SCHLIESSFEHLER"):
                        importiere(datei)

                self._pruefe_bereinigung(datei, pfad)
                self._pruefe_unveraenderten_bestand()


class RegulaererImportViewTests(TestCase):
    """Prüft reguläre Uploads über URL, Sitzung, CSRF-Schutz und Importfunktion."""

    def setUp(self):
        """Bereitet Altbestand, gültige Vorlage und einen Client mit aktivem CSRF-Schutz vor."""
        self.url = "/untersuchung_importieren/"
        untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        self.mappe = Workbook()
        self.addCleanup(self.mappe.close)
        for adresse, wert in {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
            "B1": "Neue Uploadstudie",
            "B2": 0.5,
            "B3": Untersuchung.PRAESENZ,
            "B4": "Testlehrperson Neu",
            "A7": 1234567,
            "A8": 2345678,
        }.items():
            self.mappe.active[adresse] = wert

        # Djangos gewöhnlicher Testclient überspringt die CSRF-Prüfung.
        # Hier wird sie aktiviert und bereitet ein gültiges Testcookie vor.
        self.client = Client(enforce_csrf_checks=True)
        anfrage = HttpRequest()
        get_token(anfrage)
        self.client.cookies[settings.CSRF_COOKIE_NAME] = anfrage.META["CSRF_COOKIE"]

    def _melde_testkonto_an(self, rolle):
        """Meldet ein synthetisches Rollenkonto mit den erforderlichen Sitzungszeitwerten an."""
        benutzername = f"upload_{rolle}"
        passwort = "a" * 64
        konto, fehlerliste = lege_konto_an(benutzername, passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertEqual(fehlerliste, [])

        antwort = self.client.post(
            "/anmelden/",
            data={"benutzername": benutzername, "passwort": passwort},
            HTTP_X_CSRFTOKEN=self.client.cookies[settings.CSRF_COOKIE_NAME].value,
        )
        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()["erfolg"])

    def _upload(self):
        """Erzeugt für jede Anfrage einen eigenen XLSX-Upload im Arbeitsspeicher."""
        with BytesIO() as puffer:
            self.mappe.save(puffer)
            datei = SimpleUploadedFile("untersuchung.xlsx", puffer.getvalue())
        self.addCleanup(datei.close)
        return datei

    def _post(self, daten):
        """Sendet eine Uploadanfrage mit dem aktuell gültigen CSRF-Token."""
        # Die Anmeldung wechselt das CSRF-Geheimnis. Deshalb immer den
        # aktuellen Cookiewert verwenden, nicht den Wert vor der Anmeldung.
        return self.client.post(
            self.url,
            data=daten,
            HTTP_X_CSRFTOKEN=self.client.cookies[settings.CSRF_COOKIE_NAME].value,
        )

    def _pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche Untersuchungs- und Teilnahmefeldwerte mit dem Altbestand."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa05_fa10_beide_rollen_koennen_regulaer_importieren(self):
        """Beide Rollen speichern eine echte XLSX-Datei und erhalten nur die Erfolgsbestätigung."""
        for rolle, studienname in [
            (Benutzerkonto.LEHRPERSON, "Upload Lehrperson"),
            (Benutzerkonto.ADMINISTRATOR, "Upload Administration"),
        ]:
            with self.subTest(rolle=rolle):
                self._melde_testkonto_an(rolle)
                self.mappe.active["B1"] = studienname
                vorherige_untersuchungen = list(
                    Untersuchung.objects.order_by("pk").values()
                )
                vorherige_teilnahmen = list(
                    Teilnahme.objects.order_by("pk").values()
                )

                antwort = self._post({"excel_datei": self._upload()})

                self.assertEqual(antwort.status_code, 200)
                self.assertEqual(antwort.json(), {
                    "erfolg": True,
                    "nachricht": "Import erfolgreich.",
                })
                # Die Antwort allein belegt keine Speicherung. Die echte
                # Importfunktion muss Untersuchung und beide Teilnahmen anlegen.
                neu = Untersuchung.objects.get(studienname=studienname)
                self.assertEqual(neu.stundenwert, 0.5)
                self.assertEqual(neu.einordnung, Untersuchung.PRAESENZ)
                self.assertEqual(neu.lehrperson, "Testlehrperson Neu")
                self.assertEqual(
                    list(Teilnahme.objects.filter(
                        untersuchung=neu
                    ).order_by("matrikelnummer").values_list("matrikelnummer", flat=True)),
                    ["1234567", "2345678"],
                )
                self.assertEqual(
                    list(Untersuchung.objects.exclude(
                        pk=neu.pk
                    ).order_by("pk").values()),
                    vorherige_untersuchungen,
                )
                self.assertEqual(
                    list(Teilnahme.objects.exclude(
                        untersuchung=neu
                    ).order_by("pk").values()),
                    vorherige_teilnahmen,
                )

    def test_fa05_get_anfrage_wird_abgewiesen(self):
        """GET wird mit POST als erlaubter Methode vor dem Importaufruf abgewiesen."""
        with patch("vp_stunden.views.importiere_untersuchung") as importfunktion:
            antwort = self.client.get(self.url)

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort.headers["Allow"], "POST")
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "HTTP-Methode nicht erlaubt.",
        })
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_qa04_upload_ohne_anmeldung_wird_abgewiesen(self):
        """Ein gültiger Upload ohne Anmeldung erreicht die Importfunktion nicht."""
        # Ein gültiger CSRF-Token isoliert die fehlende Anmeldung als Ursache.
        with patch("vp_stunden.views.importiere_untersuchung") as importfunktion:
            antwort = self._post({"excel_datei": self._upload()})

        self.assertEqual(antwort.status_code, 403)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Zugriff verweigert.",
        })
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_qa03_abgelaufene_sitzung_verhindert_upload(self):
        """Ein Upload mit abgelaufener Sitzung wird vor dem Import abgewiesen."""
        self._melde_testkonto_an(Benutzerkonto.LEHRPERSON)
        sitzung = self.client.session
        # Nur den Sitzungszeitwert zurücksetzen, statt die Frist abzuwarten.
        sitzung["letzte_aktivitaet"] = (
            timezone.now().timestamp() - settings.SITZUNG_INAKTIVITAET_SEKUNDEN - 1
        )
        sitzung.save()

        with patch("vp_stunden.views.importiere_untersuchung") as importfunktion:
            antwort = self._post({"excel_datei": self._upload()})

        self.assertEqual(antwort.status_code, 403)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Zugriff verweigert.",
        })
        self.assertNotIn(SESSION_KEY, self.client.session)
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_fa05_fehlender_dateiupload_wird_abgewiesen(self):
        """Eine fehlende Datei und ein bloßer Dateiname im Textfeld genügen nicht."""
        self._melde_testkonto_an(Benutzerkonto.LEHRPERSON)
        for daten in [{}, {"excel_datei": "untersuchung.xlsx"}]:
            with self.subTest(daten=daten):
                with patch("vp_stunden.views.importiere_untersuchung") as importfunktion:
                    antwort = self._post(daten)

                self.assertEqual(antwort.status_code, 400)
                self.assertEqual(antwort.json(), {
                    "erfolg": False,
                    "nachricht": "Keine Datei hochgeladen.",
                })
                importfunktion.assert_not_called()
                self._pruefe_unveraenderten_bestand()

    def test_fa09_importfehler_werden_vollstaendig_zurueckgegeben(self):
        """Die View gibt mehrere echte Zellfehler zurück, ohne gültige Teile zu speichern."""
        self._melde_testkonto_an(Benutzerkonto.LEHRPERSON)
        self.mappe.active["B2"] = 0.25
        self.mappe.active["B3"] = "X"

        antwort = self._post({"excel_datei": self._upload()})

        self.assertEqual(antwort.status_code, 400)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Import fehlgeschlagen.",
            "fehlerliste": [
                "B2: Der Stundenwert muss ein Vielfaches von 0,5 sein.",
                "B3: Die Einordnung muss exakt P oder O sein.",
            ],
        })
        self._pruefe_unveraenderten_bestand()

    def test_csrf_fehlender_oder_falscher_token_verhindert_upload(self):
        """Die CSRF-Middleware weist angemeldete Uploads ohne passenden Token ab."""
        self._melde_testkonto_an(Benutzerkonto.LEHRPERSON)
        token = self.client.cookies[settings.CSRF_COOKIE_NAME].value
        # Ein Zeichen ändern: Der falsche Token bleibt formal gültig,
        # stimmt aber sicher nicht mehr mit dem vorhandenen Cookie überein.
        falsches_token = ("a" if token[0] != "a" else "b") + token[1:]
        for prueftoken in [None, falsches_token]:
            with self.subTest(csrf_token=prueftoken):
                kopfdaten = (
                    {} if prueftoken is None else {"HTTP_X_CSRFTOKEN": prueftoken}
                )
                with patch("vp_stunden.views.importiere_untersuchung") as importfunktion:
                    antwort = self.client.post(
                        self.url,
                        data={"excel_datei": self._upload()},
                        **kopfdaten,
                    )

                # Diese Abweisung kommt vor der View von Django und ist
                # deshalb keine JSON-Fehlerantwort der Import-View.
                self.assertEqual(antwort.status_code, 403)
                importfunktion.assert_not_called()
                self._pruefe_unveraenderten_bestand()

    def test_qa06_falsch_zugeordneter_temporaerer_upload_wird_entfernt(self):
        """Django entfernt auch eine vor dem Import abgewiesene temporäre Uploaddatei."""
        self._melde_testkonto_an(Benutzerkonto.LEHRPERSON)
        with TemporaryDirectory() as verzeichnis:
            with self.settings(
                FILE_UPLOAD_MAX_MEMORY_SIZE=0,
                FILE_UPLOAD_TEMP_DIR=verzeichnis,
            ):
                with patch("vp_stunden.views.importiere_untersuchung") as importfunktion:
                    antwort = self._post({"anderes_feld": self._upload()})

                self.assertEqual(antwort.status_code, 400)
                self.assertEqual(antwort.json(), {
                    "erfolg": False,
                    "nachricht": "Keine Datei hochgeladen.",
                })
                importfunktion.assert_not_called()

                # Der andere Feldname verhindert den Importaufruf. Das Lesen
                # der Anfrage hat dennoch einen echten Datenträgerupload erzeugt.
                # Der Testclient schließt die Anfrage vor Rückgabe der Antwort.
                datei = antwort.wsgi_request.FILES["anderes_feld"]
                self.assertIsInstance(datei, TemporaryUploadedFile)
                self.assertTrue(datei.closed)
                self.assertFalse(Path(datei.temporary_file_path()).exists())
                self._pruefe_unveraenderten_bestand()


class UebergangsimportViewTests(TestCase):
    """Prüft den administrativen Übergangsimport über URL, Sitzung und CSRF-Schutz."""

    def setUp(self):
        """Bereitet Altbestand, eine gemischte Übergangsdatei und einen geschützten Client vor."""
        self.url = "/uebergangsdatei_importieren/"
        self.untersuchung = Untersuchung.objects.create(
            studienname="Bestandsstudie",
            studienname_vergleich="bestandsstudie",
            stundenwert=1.5,
            einordnung=Untersuchung.PRAESENZ,
            lehrperson="Testlehrperson Alt",
        )
        Teilnahme.objects.create(
            untersuchung=self.untersuchung,
            matrikelnummer="9876543",
        )
        self.untersuchungsbestand = list(
            Untersuchung.objects.order_by("pk").values()
        )
        self.teilnahmebestand = list(
            Teilnahme.objects.order_by("pk").values()
        )

        self.mappe = Workbook()
        self.addCleanup(self.mappe.close)
        self.mappe.active.append([
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        ])
        # Beide Zeilen betreffen dieselbe Matrikelnummer. Die erste ergänzt
        # eine vorhandene Untersuchung, die zweite legt eine neue an.
        self.mappe.active.append([
            1234567,
            "Bestandsstudie",
            1.5,
            Untersuchung.PRAESENZ,
            "Testlehrperson Alt",
        ])
        self.mappe.active.append([
            1234567,
            "Neue Übergangsstudie",
            0.5,
            Untersuchung.ONLINE,
            "Testlehrperson Neu",
        ])

        # Die CSRF-Prüfung bleibt auch bei den übrigen Uploadtests aktiv.
        # Das Testcookie ermöglicht zunächst die geschützte Anmeldung.
        self.client = Client(enforce_csrf_checks=True)
        anfrage = HttpRequest()
        get_token(anfrage)
        self.client.cookies[settings.CSRF_COOKIE_NAME] = anfrage.META["CSRF_COOKIE"]

    def _melde_testkonto_an(self, rolle):
        """Meldet ein synthetisches Rollenkonto mit den erforderlichen Sitzungszeitwerten an."""
        benutzername = f"uebergang_{rolle}"
        passwort = "a" * 64
        konto, fehlerliste = lege_konto_an(benutzername, passwort, rolle)
        self.assertIsNotNone(konto)
        self.assertEqual(fehlerliste, [])

        antwort = self.client.post(
            "/anmelden/",
            data={"benutzername": benutzername, "passwort": passwort},
            HTTP_X_CSRFTOKEN=self.client.cookies[settings.CSRF_COOKIE_NAME].value,
        )
        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()["erfolg"])

    def _upload(self):
        """Erzeugt für jede Anfrage einen eigenen XLSX-Upload im Arbeitsspeicher."""
        with BytesIO() as puffer:
            self.mappe.save(puffer)
            datei = SimpleUploadedFile("uebergang.xlsx", puffer.getvalue())
        self.addCleanup(datei.close)
        return datei

    def _post(self, daten):
        """Sendet eine Übergangsimportanfrage mit dem aktuell gültigen CSRF-Token."""
        # Die Anmeldung wechselt das CSRF-Geheimnis. Deshalb immer den
        # aktuellen Cookiewert verwenden, nicht den Wert vor der Anmeldung.
        return self.client.post(
            self.url,
            data=daten,
            HTTP_X_CSRFTOKEN=self.client.cookies[settings.CSRF_COOKIE_NAME].value,
        )

    def _pruefe_unveraenderten_bestand(self):
        """Vergleicht sämtliche Untersuchungs- und Teilnahmefeldwerte mit dem Altbestand."""
        self.assertEqual(
            list(Untersuchung.objects.order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa11_administrator_importiert_neue_und_vorhandene_untersuchung(self):
        """Eine administrative Anfrage legt neue Daten an und erhält die vorhandene Untersuchung."""
        self._melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)

        antwort = self._post({"uebergangsdatei": self._upload()})

        self.assertEqual(antwort.status_code, 200)
        self.assertEqual(antwort.json(), {
            "erfolg": True,
            "nachricht": "Import erfolgreich.",
        })
        # Die echte Importfunktion muss beide Zeilen gemeinsam übernehmen.
        # Die Antwort enthält dabei keine gespeicherten Untersuchungsdaten.
        neu = Untersuchung.objects.get(studienname="Neue Übergangsstudie")
        self.assertEqual(neu.stundenwert, 0.5)
        self.assertEqual(neu.einordnung, Untersuchung.ONLINE)
        self.assertEqual(neu.lehrperson, "Testlehrperson Neu")
        self.assertCountEqual(
            list(Teilnahme.objects.filter(
                matrikelnummer="1234567"
            ).values_list("untersuchung_id", flat=True)),
            [self.untersuchung.pk, neu.pk],
        )
        # Der vollständige Vergleich umfasst auch Schreibweise und Uploadzeitpunkt.
        self.assertEqual(
            list(Untersuchung.objects.exclude(
                pk=neu.pk
            ).order_by("pk").values()),
            self.untersuchungsbestand,
        )
        self.assertEqual(
            list(Teilnahme.objects.exclude(
                matrikelnummer="1234567"
            ).order_by("pk").values()),
            self.teilnahmebestand,
        )

    def test_fa11_get_anfrage_wird_abgewiesen(self):
        """GET wird mit POST als erlaubter Methode vor dem Importaufruf abgewiesen."""
        with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
            antwort = self.client.get(self.url)

        self.assertEqual(antwort.status_code, 405)
        self.assertEqual(antwort.headers["Allow"], "POST")
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "HTTP-Methode nicht erlaubt.",
        })
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_qa04_upload_ohne_anmeldung_wird_abgewiesen(self):
        """Ein gültiger Übergangsupload ohne Anmeldung erreicht die Importfunktion nicht."""
        # Ein gültiger CSRF-Token isoliert die fehlende Anmeldung als Ursache.
        with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
            antwort = self._post({"uebergangsdatei": self._upload()})

        self.assertEqual(antwort.status_code, 403)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Zugriff verweigert.",
        })
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_qa04_lehrperson_darf_keinen_uebergangsimport_ausfuehren(self):
        """Eine angemeldete Lehrperson darf auch eine gültige Übergangsdatei nicht importieren."""
        self._melde_testkonto_an(Benutzerkonto.LEHRPERSON)

        with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
            antwort = self._post({"uebergangsdatei": self._upload()})

        self.assertEqual(antwort.status_code, 403)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Zugriff verweigert.",
        })
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_qa03_abgelaufene_sitzung_verhindert_uebergangsimport(self):
        """Ein administrativer Upload mit abgelaufener Sitzung wird vor dem Import abgewiesen."""
        self._melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        sitzung = self.client.session
        # Nur den Sitzungszeitwert zurücksetzen, statt die Frist abzuwarten.
        sitzung["letzte_aktivitaet"] = (
            timezone.now().timestamp() - settings.SITZUNG_INAKTIVITAET_SEKUNDEN - 1
        )
        sitzung.save()

        with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
            antwort = self._post({"uebergangsdatei": self._upload()})

        self.assertEqual(antwort.status_code, 403)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Zugriff verweigert.",
        })
        self.assertNotIn(SESSION_KEY, self.client.session)
        importfunktion.assert_not_called()
        self._pruefe_unveraenderten_bestand()

    def test_fa11_fehlender_dateiupload_wird_abgewiesen(self):
        """Eine fehlende Datei und ein bloßer Dateiname im Textfeld genügen nicht."""
        self._melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        for daten in [{}, {"uebergangsdatei": "uebergang.xlsx"}]:
            with self.subTest(daten=daten):
                with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
                    antwort = self._post(daten)

                self.assertEqual(antwort.status_code, 400)
                self.assertEqual(antwort.json(), {
                    "erfolg": False,
                    "nachricht": "Keine Datei hochgeladen.",
                })
                importfunktion.assert_not_called()
                self._pruefe_unveraenderten_bestand()

    def test_fa11_importfehler_verhindern_auch_die_bestandsergaenzung(self):
        """Fehler in der neuen Untersuchung verhindern auch die gültige Ergänzung der vorhandenen."""
        self._melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        self.mappe.active["C3"] = 0.25
        self.mappe.active["D3"] = "X"

        antwort = self._post({"uebergangsdatei": self._upload()})

        self.assertEqual(antwort.status_code, 400)
        self.assertEqual(antwort.json(), {
            "erfolg": False,
            "nachricht": "Import fehlgeschlagen.",
            "fehlerliste": [
                "C3: Der Stundenwert muss ein Vielfaches von 0,5 sein.",
                "D3: Die Einordnung muss exakt P oder O sein.",
            ],
        })
        self._pruefe_unveraenderten_bestand()

    def test_csrf_fehlender_oder_falscher_token_verhindert_uebergangsimport(self):
        """Die CSRF-Middleware weist administrative Uploads ohne passenden Token ab."""
        self._melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        token = self.client.cookies[settings.CSRF_COOKIE_NAME].value
        # Ein Zeichen ändern: Der falsche Token bleibt formal gültig,
        # stimmt aber sicher nicht mehr mit dem vorhandenen Cookie überein.
        falsches_token = ("a" if token[0] != "a" else "b") + token[1:]
        for prueftoken in [None, falsches_token]:
            with self.subTest(csrf_token=prueftoken):
                kopfdaten = (
                    {} if prueftoken is None else {"HTTP_X_CSRFTOKEN": prueftoken}
                )
                with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
                    antwort = self.client.post(
                        self.url,
                        data={"uebergangsdatei": self._upload()},
                        **kopfdaten,
                    )

                # Diese Abweisung kommt vor der View von Django und ist
                # deshalb keine JSON-Fehlerantwort der Import-View.
                self.assertEqual(antwort.status_code, 403)
                importfunktion.assert_not_called()
                self._pruefe_unveraenderten_bestand()

    def test_qa06_falsch_zugeordneter_temporaerer_upload_wird_entfernt(self):
        """Django entfernt auch eine vor dem Übergangsimport abgewiesene temporäre Uploaddatei."""
        self._melde_testkonto_an(Benutzerkonto.ADMINISTRATOR)
        with TemporaryDirectory() as verzeichnis:
            with self.settings(
                FILE_UPLOAD_MAX_MEMORY_SIZE=0,
                FILE_UPLOAD_TEMP_DIR=verzeichnis,
            ):
                with patch("vp_stunden.views.importiere_uebergangsdatei") as importfunktion:
                    antwort = self._post({"excel_datei": self._upload()})

                self.assertEqual(antwort.status_code, 400)
                self.assertEqual(antwort.json(), {
                    "erfolg": False,
                    "nachricht": "Keine Datei hochgeladen.",
                })
                importfunktion.assert_not_called()

                # Die Übergangsdatei liegt im Feld des regulären Uploads.
                # Django hat sie eingelesen, die View akzeptiert dieses Feld aber nicht.
                # Der Testclient schließt die Anfrage vor Rückgabe der Antwort.
                datei = antwort.wsgi_request.FILES["excel_datei"]
                self.assertIsInstance(datei, TemporaryUploadedFile)
                self.assertTrue(datei.closed)
                self.assertFalse(Path(datei.temporary_file_path()).exists())
                self._pruefe_unveraenderten_bestand()


class OberflaecheViewTests(TestCase):
    """Prüft Oberflächenzugriff, Sitzungsende und das lokale HTTP-Cookieprofil."""

    def setUp(self):
        """Legt zwei synthetische Rollenkonten für die Oberflächenprüfungen an."""
        self.url = "/?bereich=administration"
        self.passwort = "a" * 64

        self.administratorkonto, fehlerliste = lege_konto_an(
            "oberflaeche_admin",
            self.passwort,
            Benutzerkonto.ADMINISTRATOR,
        )
        self.assertIsNotNone(self.administratorkonto)
        self.assertEqual(fehlerliste, [])

        self.lehrpersonenkonto, fehlerliste = lege_konto_an(
            "oberflaeche_lehrperson",
            self.passwort,
            Benutzerkonto.LEHRPERSON,
        )
        self.assertIsNotNone(self.lehrpersonenkonto)
        self.assertEqual(fehlerliste, [])

    def _melde_testkonto_an(self, konto):
        """Meldet ein vorbereitetes Konto an und gibt die HTTP-Antwort zurück."""
        # Die reguläre Anmeldung setzt auch Beginn und letzte Aktivität der Sitzung.
        antwort = self.client.post(
            "/anmelden/",
            data={
                "benutzername": konto.username,
                "passwort": self.passwort,
            },
        )

        self.assertEqual(antwort.status_code, 200)
        self.assertTrue(antwort.json()["erfolg"])
        self.assertEqual(self.client.session[SESSION_KEY], str(konto.pk))
        return antwort

    def _pruefe_keine_administrationsdaten(self, antwort):
        """Prüft, dass weder Kontendaten noch administrative Kontenformulare erscheinen."""
        # Eine eingeschränkte HTML-Ansicht darf weiterhin mit HTTP 200 antworten.
        self.assertEqual(antwort.status_code, 200)
        self.assertEqual(antwort.context["konten"], [])

        # Die eindeutigen Namen machen eine versehentliche Kontenausgabe erkennbar.
        for konto in [self.administratorkonto, self.lehrpersonenkonto]:
            self.assertNotContains(antwort, konto.username)

        self.assertNotContains(antwort, 'id="kontoformular"')
        self.assertNotContains(antwort, 'id="ruecksetzformular"')

    def test_fa04_administrator_erhaelt_kontenauswahl(self):
        """Eine gültige Administrationssitzung erhält die vollständige Kontenauswahl."""
        self._melde_testkonto_an(self.administratorkonto)

        antwort = self.client.get(self.url)

        self.assertEqual(antwort.status_code, 200)
        self.assertTemplateUsed(antwort, "vp_stunden/oberflaeche.html")

        # Die View liefert genau die für die Auswahl vorgesehenen Kontenangaben.
        self.assertEqual(
            list(antwort.context["konten"]),
            [
                {
                    "id": self.administratorkonto.pk,
                    "username": self.administratorkonto.username,
                    "rolle": Benutzerkonto.ADMINISTRATOR,
                },
                {
                    "id": self.lehrpersonenkonto.pk,
                    "username": self.lehrpersonenkonto.username,
                    "rolle": Benutzerkonto.LEHRPERSON,
                },
            ],
        )

        # Zusätzlich zum Kontext die tatsächlich ausgegebenen HTML-Inhalte prüfen.
        self.assertContains(antwort, self.administratorkonto.username)
        self.assertContains(antwort, self.lehrpersonenkonto.username)
        self.assertContains(antwort, 'id="kontoformular"')
        self.assertContains(antwort, 'id="ruecksetzformular"')
        self.assertContains(antwort, 'id="passwortformular"')

    def test_qa04_ohne_anmeldung_erscheinen_keine_kontendaten(self):
        """Ohne Anmeldung erscheint nur die Anmeldeansicht ohne Kontendaten."""
        antwort = self.client.get(self.url)

        self._pruefe_keine_administrationsdaten(antwort)
        self.assertContains(antwort, 'id="anmeldeformular"')
        self.assertNotContains(antwort, 'id="passwortformular"')
        self.assertFalse(antwort.wsgi_request.user.is_authenticated)

    def test_qa04_lehrperson_erhaelt_keine_kontendaten(self):
        """Eine Lehrperson erhält nur die eigene Passwortfunktion ohne Kontenauswahl."""
        self._melde_testkonto_an(self.lehrpersonenkonto)

        antwort = self.client.get(self.url)

        self._pruefe_keine_administrationsdaten(antwort)
        self.assertContains(antwort, 'id="passwortformular"')
        self.assertNotContains(antwort, 'id="anmeldeformular"')
        self.assertTrue(antwort.wsgi_request.user.is_authenticated)

    def test_qa03_abgelaufene_sitzung_verhindert_kontenausgabe(self):
        """Beide Sitzungsfristen verhindern nach ihrem Erreichen die Kontenausgabe."""
        jetzt = timezone.now()
        zeitwert = jetzt.timestamp()
        testfaelle = [
            (
                "inaktivitaet",
                zeitwert - settings.SITZUNG_INAKTIVITAET_SEKUNDEN,
                zeitwert - settings.SITZUNG_INAKTIVITAET_SEKUNDEN,
            ),
            (
                "absolute_dauer",
                zeitwert - settings.SITZUNG_MAX_DAUER_SEKUNDEN,
                zeitwert,
            ),
        ]

        for frist, sitzungsbeginn, letzte_aktivitaet in testfaelle:
            with self.subTest(frist=frist):
                # Jeder Teilfall erhält eine frisch angemeldete Sitzung.
                self.client = Client()

                # Eine feste Testzeit vermeidet Wartezeiten und ungenaue Grenzvergleiche.
                with patch(
                    "vp_stunden.sicherheit.timezone.now",
                    return_value=jetzt,
                ):
                    self._melde_testkonto_an(self.administratorkonto)
                    sitzung = self.client.session
                    sitzung["sitzungsbeginn"] = sitzungsbeginn
                    sitzung["letzte_aktivitaet"] = letzte_aktivitaet
                    sitzung.save()

                    antwort = self.client.get(self.url)

                self._pruefe_keine_administrationsdaten(antwort)
                self.assertContains(antwort, 'id="anmeldeformular"')
                self.assertFalse(antwort.wsgi_request.user.is_authenticated)

    def test_qa03_alte_sitzungskennung_erlaubt_keinen_kontenzugriff(self):
        """Die vor der Abmeldung gesicherte Kennung ermöglicht keinen erneuten Zugriff."""
        self._melde_testkonto_an(self.administratorkonto)
        alte_sitzungskennung = self.client.cookies[
            settings.SESSION_COOKIE_NAME
        ].value

        antwort_abmeldung = self.client.post("/abmelden/")
        self.assertEqual(antwort_abmeldung.status_code, 200)
        self.assertTrue(antwort_abmeldung.json()["erfolg"])

        # Die alte Kennung absichtlich mit einem getrennten Client erneut senden.
        neuer_client = Client()
        neuer_client.cookies[settings.SESSION_COOKIE_NAME] = alte_sitzungskennung
        antwort = neuer_client.get(self.url)

        self._pruefe_keine_administrationsdaten(antwort)
        self.assertContains(antwort, 'id="anmeldeformular"')
        self.assertFalse(antwort.wsgi_request.user.is_authenticated)

    def test_http_cookieprofil_entspricht_lokalem_pruefbetrieb(self):
        """Die Anmeldeantwort setzt die vorgesehenen Cookieattribute für lokales HTTP."""
        antwort = self._melde_testkonto_an(self.administratorkonto)

        # Die tatsächlich gesetzten Antwortcookies statt nur Einstellungen prüfen.
        self.assertIn(settings.SESSION_COOKIE_NAME, antwort.cookies)
        self.assertIn(settings.CSRF_COOKIE_NAME, antwort.cookies)
        sitzungscookie = antwort.cookies[settings.SESSION_COOKIE_NAME]
        csrf_cookie = antwort.cookies[settings.CSRF_COOKIE_NAME]

        # Eine leere Domain entspricht einem hostgebundenen Sitzungscookie.
        self.assertEqual(sitzungscookie["domain"], "")
        self.assertTrue(sitzungscookie["httponly"])
        self.assertEqual(sitzungscookie["samesite"], "Lax")

        # Im bestätigten HTTP-Prüfbetrieb darf Secure bei beiden Cookies fehlen.
        self.assertFalse(sitzungscookie["secure"])
        self.assertFalse(csrf_cookie["secure"])