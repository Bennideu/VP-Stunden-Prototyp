from zipfile import ZipFile  # Liest das ZIP-Paket direkt aus der Uploaddatei.

from defusedxml.ElementTree import fromstring  # Liest das Paketmanifest mit Schutz gegen XML-Entitäten.
from django.core.exceptions import ValidationError  # Ausnahmetyp der vorhandenen Feldprüfungen.
from django.core.files.uploadedfile import UploadedFile  # Typ einer von Django bereitgestellten Uploaddatei.
from django.db import DatabaseError, transaction  # Behandelt Datenbankfehler und atomare Speicheroperationen.
from openpyxl import DEFUSEDXML, Workbook, load_workbook  # XML-Schutzstatus, Workbook-Typ und Ladefunktion.
from openpyxl.xml.constants import XLSM, XLSX, XLTM, XLTX  # Inhaltstypen für XLSX und die abzuweisenden Makro-/Vorlagenformate.

from .models import Untersuchung, Teilnahme  # Modellklassen für Untersuchungen und Teilnahmezuordnungen.
from .vp_daten import (
    pruefe_einordnung,  # Prüft die Kennzeichnung P oder O.
    pruefe_lehrperson,  # Prüft und bereinigt den Namen der Lehrperson.
    pruefe_matrikelnummer,  # Prüft den numerischen Matrikelnummernwert.
    pruefe_studienname,  # Liefert den bereinigten Studiennamen und seinen Vergleichswert.
    pruefe_stundenwert,  # Prüft den positiven Stundenwert in 0,5-Schritten.
)


MAX_DATEIGROESSE = 2_000_000  # Verbindliche Uploadgrenze in Byte einschließlich Grenzwert.
MAX_ZIP_EINTRAEGE = 100  # Vorläufige Obergrenze für sämtliche ZIP-Einträge.
MAX_ENTPACKTE_GROESSE = 20 * 1024 * 1024  # Vorläufige Obergrenze von 20 MiB.


def _lese_xlsx(datei: UploadedFile) -> tuple[Workbook | None, list[str]]:
    """
    Prüft die gemeinsamen Dateigrenzen und liest eine makrofreie XLSX-Datei.

    :param datei: Die direkt übergebene Uploaddatei.
    :return: Arbeitsmappe und leere Fehlerliste oder None und Fehlermeldungen.

    Die ZIP- und Manifestprüfung prüft Grenzen, Prüfsummen und relevante
    Inhaltstypen. Sie ist keine vollständige OOXML-Schemavalidierung.
    Nach dem Laden muss genau ein sichtbares Tabellenblatt vorliegen.

    Bei Erfolg erhält der aufrufende Import die geladene Arbeitsmappe zurück.
    Die Importfunktionen prüfen ihren Vorlagenaufbau, Formeln in sämtlichen
    vorhandenen Zellen und die Zellwerte. Sie schließen anschließend
    Arbeitsmappe und Uploaddatei auch bei einer Abweisung.
    """
    arbeitsmappe = None

    try:
        # size ist die von Django ermittelte Uploadgröße in Byte.
        # Mit > bleibt eine Datei genau am Grenzwert zulässig.
        if datei.size > MAX_DATEIGROESSE:
            return None, ["Datei: Höchstens 2.000.000 Byte sind zulässig."]

        if not (datei.name or "").lower().endswith(".xlsx"):
            return None, ["Datei: Eine makrofreie .xlsx-Datei ist erforderlich."]

        # Das Einlesen darf nicht ohne den vorgesehenen XML-Schutz erfolgen.
        if not DEFUSEDXML:
            return None, ["Datei: Sicheres XML-Einlesen ist nicht verfügbar."]

        # seek(0) setzt die Leseposition auf den Anfang des vorhandenen Dateiobjekts.
        datei.seek(0)

        # Das vorhandene Dateiobjekt verwenden, ohne eine eigene Kopie anzulegen.
        with ZipFile(datei) as archiv:
            # infolist() liefert je ZIP-Eintrag ein Objekt mit dessen Metadaten.
            # Auch Verzeichniseinträge zählen zur festgelegten Eintragsgrenze.
            eintraege = archiv.infolist()

            if len(eintraege) > MAX_ZIP_EINTRAEGE:
                return None, ["Datei: Höchstens 100 ZIP-Einträge sind zulässig."]

            # Der Ausdruck in sum() liefert nacheinander die file_size jedes Eintrags.
            # sum() addiert diese angegebenen unkomprimierten Größen.
            # Die Grenze betrifft die ZIP-Inhalte, nicht den gesamten Arbeitsspeicher.
            entpackte_groesse = sum(eintrag.file_size for eintrag in eintraege)
            if entpackte_groesse > MAX_ENTPACKTE_GROESSE:
                return None, ["Datei: Die entpackte Gesamtgröße überschreitet 20 MiB."]

            # testzip() liest die Paketbestandteile und kontrolliert ihre ZIP-Prüfsummen.
            # None bedeutet, dass dabei kein beschädigter Eintrag gefunden wurde.
            if archiv.testzip() is not None:
                return None, ["Datei: Das ZIP-Paket enthält beschädigte Daten."]

            # Nach der Größen- und Prüfsummenprüfung das Paketmanifest auswerten.
            # read() liefert seine XML-Bytes, fromstring() das XML-Wurzelelement.
            # Dafür wird weder eine Datei extrahiert noch eine Internetadresse aufgerufen.
            manifest = fromstring(archiv.read("[Content_Types].xml"))

            # Jedes Kind des Manifests kann ein ContentType-Attribut enthalten.
            # get(..., "") liefert bei fehlendem Attribut eine leere Zeichenfolge.
            typen = [element.get("ContentType", "") for element in manifest]
            # Aus allen Pakettypen die bekannten Arbeitsmappenformate auswählen.
            # Genau ein XLSX-Haupttyp ist zulässig. Die Dateiendung allein genügt nicht.
            arbeitsmappentypen = [
                typ for typ in typen if typ in (XLSX, XLSM, XLTX, XLTM)
            ]
            if arbeitsmappentypen != [XLSX]:
                return None, ["Datei: Eine makrofreie XLSX-Arbeitsmappe ist erforderlich."]

            # any() wird wahr, sobald ein deklarierter Inhaltstyp auf Makros hinweist.
            # lower() vereinheitlicht dabei die Schreibweise für diesen Vergleich.
            if any("macro" in typ.lower() or "vba" in typ.lower() for typ in typen):
                return None, ["Datei: Makrobestandteile sind nicht zulässig."]

        # Nach der Archivprüfung muss openpyxl wieder am Dateianfang lesen können.
        datei.seek(0)
        arbeitsmappe = load_workbook(
            datei,
            read_only=False,  # Die begrenzte Arbeitsmappe vollständig einlesen.
            data_only=False,  # Formeln für die spätere Abweisung erkennbar halten.
            keep_vba=False,  # VBA-Inhalte nicht übernehmen.
            keep_links=True,  # Erkannte Verknüpfungen für die Abweisung erhalten.
        )

        fehlerliste = []

        # sheetnames umfasst auch Diagrammblätter, worksheets nur normale Tabellenblätter.
        # Beide Prüfungen zusammen schließen zusätzliche und reine Diagrammblätter aus.
        if len(arbeitsmappe.sheetnames) != 1 or len(arbeitsmappe.worksheets) != 1:
            fehlerliste.append("Datei: Genau ein Tabellenblatt ohne weitere Blätter ist erforderlich.")
        elif arbeitsmappe.worksheets[0].sheet_state != "visible":
            fehlerliste.append("Datei: Das einzige Tabellenblatt muss sichtbar sein.")

        # Private openpyxl-3.1.5-Schnittstelle: Die Liste enthält die beim Laden
        # erkannten externen Arbeitsmappenverweise. keep_links=True erhält sie.
        # Eine gefüllte Liste ist in einer if-Bedingung wahr und führt zur Abweisung.
        if arbeitsmappe._external_links:
            fehlerliste.append("Datei: Externe Arbeitsmappenverknüpfungen sind nicht zulässig.")

        if fehlerliste:
            arbeitsmappe.close()
            return None, fehlerliste

        # Das Tupel enthält den Ergebniswert und getrennt davon die Fehlerliste.
        return arbeitsmappe, []

    except Exception:
        # ZIP-, XML- und Bibliotheksfehler an der äußeren Lesegrenze abfangen.
        # Die Originalausnahme wird nicht in die Meldung an die Oberfläche übernommen.
        if arbeitsmappe is not None:
            arbeitsmappe.close()
        return None, ["Datei: Die XLSX-Datei konnte nicht gelesen werden."]


def importiere_untersuchung(datei: UploadedFile) -> list[str]:
    """
    Prüft und speichert den regulären Import genau einer Untersuchung.

    :param datei: Die von Django bereitgestellte Uploaddatei.
    :return: Gesammelte Fehler oder eine leere Liste nach erfolgreicher Speicherung.

    Die Vorlage enthält gemeinsame Untersuchungsangaben und bis zu 100
    unterschiedliche Matrikelnummern. Fehler werden gesammelt. Erst wenn
    die gesamte Datei gültig und der Studienname noch nicht gespeichert ist,
    werden die Untersuchung und alle Teilnahmen gemeinsam angelegt.
    """
    arbeitsmappe = None
    fehlerliste = []

    try:
        arbeitsmappe, fehlerliste = _lese_xlsx(datei)
        if arbeitsmappe is None:
            return fehlerliste

        blatt = arbeitsmappe.worksheets[0]
        beschriftungen = {
            "A1": "Studienname",
            "A2": "Stundenwert (VP-Stunden)",
            "A3": "Einordnung (P/O)",
            "A4": "Verantwortliche Lehrperson",
            "A6": "Matrikelnummer",
        }
        # Eine Menge merkt eindeutige Zelladressen wie "B2".
        # Die Formelerkennung erfolgt erst über data_type in der folgenden Schleife.
        formelzellen = set()

        # _cells ist eine private openpyxl-3.1.5-Schnittstelle.
        # Sie besucht nur vorhandene Zellen und vermeidet riesige leere
        # Zwischenbereiche bei weit außen liegenden formatierten Leerzellen.
        for zelle in blatt._cells.values():
            # "f" kennzeichnet eine echte Excel-Formel. coordinate liefert ihre Adresse.
            # Die Adresse wird gemerkt, damit später kein Formeltext als Feldwert gilt.
            if zelle.data_type == "f":
                formelzellen.add(zelle.coordinate)
                fehlerliste.append(f"{zelle.coordinate}: Formeln sind nicht zulässig.")

            # Nur None gilt als leer. Leerzeichen, 0 und False bleiben prüfbarer Inhalt.
            if zelle.value is None:
                continue

            # Zulässig sind die festen Beschriftungen und die beiden Eingabebereiche.
            erlaubt = (
                zelle.coordinate in beschriftungen
                or (zelle.column == 2 and 1 <= zelle.row <= 4)
                or (zelle.column == 1 and 7 <= zelle.row <= 106)
            )
            if not erlaubt:
                fehlerliste.append(
                    f"{zelle.coordinate}: Außerhalb der Vorlage sind keine Inhalte zulässig."
                )

        # Die Beschriftungen in A1:A4 und A6 müssen exakt übereinstimmen.
        for adresse, erwartet in beschriftungen.items():
            if adresse not in formelzellen and blatt[adresse].value != erwartet:
                fehlerliste.append(
                    f"{adresse}: Die Beschriftung muss „{erwartet}“ lauten."
                )

        # Die Liste ordnet jeder Zelladresse einen Ergebnisschlüssel und eine
        # vorhandene Prüffunktion zu. Der Funktionsname wird hier noch nicht aufgerufen.
        werte = {}
        for adresse, schluessel, pruefung in [
            ("B1", "studienname", pruefe_studienname),
            ("B2", "stundenwert", pruefe_stundenwert),
            ("B3", "einordnung", pruefe_einordnung),
            ("B4", "lehrperson", pruefe_lehrperson),
        ]:
            if adresse in formelzellen:
                continue

            # Beispiel für B2: werte["stundenwert"] = pruefe_stundenwert(Zellwert).
            # Nur erfolgreich geprüfte, gegebenenfalls bereinigte Werte werden gemerkt.
            # Beim Studiennamen ist das Ergebnis ein Tupel aus Anzeige- und Vergleichswert.
            try:
                werte[schluessel] = pruefung(blatt[adresse].value)
            except ValidationError as fehler:
                # Ein Prüffehler kann mehrere Meldungen enthalten.
                # Die Zelladresse ordnet jede Meldung ihrem Eingabefeld zu.
                for meldung in fehler.messages:
                    fehlerliste.append(f"{adresse}: {meldung}")
            except OverflowError:
                # Extrem große Zahlen können bei der Umwandlung in float scheitern.
                # Der Feldfehler beendet die Prüfung anderer Zellen nicht.
                fehlerliste.append(f"{adresse}: Der Zahlenwert ist nicht darstellbar.")

        matrikelnummern = []
        # Die Menge erkennt mehrfach vorkommende Matrikelnummern in dieser Datei.
        # Die Liste darüber erhält dagegen ihre Reihenfolge für die spätere Speicherung.
        vorhandene_nummern = set()

        # range(7, 107) umfasst die Zeilen 7 bis 106, da die obere Grenze exklusiv ist.
        # Einzelne Leerzellen beenden diese Prüfung nicht.
        for zeile in range(7, 107):
            zelle = blatt.cell(row=zeile, column=1)
            if zelle.value is None or zelle.coordinate in formelzellen:
                continue

            try:
                matrikelnummer = pruefe_matrikelnummer(zelle.value)
            except ValidationError as fehler:
                for meldung in fehler.messages:
                    fehlerliste.append(f"{zelle.coordinate}: {meldung}")
                continue

            if matrikelnummer in vorhandene_nummern:
                fehlerliste.append(
                    f"{zelle.coordinate}: Die Matrikelnummer kommt innerhalb der Datei mehrfach vor."
                )
            else:
                vorhandene_nummern.add(matrikelnummer)
                matrikelnummern.append(matrikelnummer)

        if not matrikelnummern:
            fehlerliste.append("A7:A106: Mindestens eine gültige Matrikelnummer ist erforderlich.")

        # Ein Bestandsvergleich ist nur mit einem gültigen Studiennamen möglich.
        if "studienname" in werte:
            # Der Prüfer hat äußere Leerzeichen entfernt und den Vergleichswert gebildet.
            studienname, vergleichswert = werte["studienname"]
            if Untersuchung.objects.filter(studienname_vergleich=vergleichswert).exists():
                fehlerliste.append("B1: Eine Untersuchung mit diesem Studiennamen ist bereits gespeichert.")

        # Bis hierhin wurde nichts in der Datenbank gespeichert.
        # Schon eine Meldung verhindert das Schreiben der gesamten Datei.
        if fehlerliste:
            return fehlerliste

        # atomic() übernimmt alle Schreibschritte gemeinsam.
        # Verlässt eine Ausnahme den Block, nimmt Django dessen Änderungen zurück.
        with transaction.atomic():
            untersuchung = Untersuchung.objects.create(
                studienname=studienname,
                studienname_vergleich=vergleichswert,
                stundenwert=werte["stundenwert"],
                einordnung=werte["einordnung"],
                lehrperson=werte["lehrperson"],
            )
            for matrikelnummer in matrikelnummern:
                Teilnahme.objects.create(
                    untersuchung=untersuchung,
                    matrikelnummer=matrikelnummer,
                )

        return []

    except DatabaseError:
        # Fehler innerhalb von atomic() gelangen erst nach dessen Rückabwicklung hierher.
        # Auch Datenbankfehler bei den vorherigen Bestandsabfragen werden hier behandelt.
        # Bereits erkannte Zellfehler bleiben durch die Listenverkettung erhalten.
        return fehlerliste + ["Datei: Der Import konnte nicht übernommen werden."]
    except Exception:
        # Auch unerwartete Verarbeitungsfehler erhalten eine allgemeine Dateimeldung.
        return fehlerliste + ["Datei: Der Import konnte nicht verarbeitet werden."]
    finally:
        # finally läuft auch vor einer Rückgabe durch return oder bei einer Ausnahme.
        # Das innere finally ruft datei.close() auch auf, falls Workbook.close() scheitert.
        try:
            if arbeitsmappe is not None:
                arbeitsmappe.close()
        finally:
            # Ein Speicherupload schließt seinen Puffer. Bei TemporaryUploadedFile
            # entfernt close() die Django-Temporärdatei, nicht die lokale Originaldatei.
            datei.close()


def importiere_uebergangsdatei(datei: UploadedFile) -> list[str]:
    """
    Prüft und importiert bis zu 100 Teilnahmen derselben Matrikelnummer.

    :param datei: Die von Django bereitgestellte Uploaddatei.
    :return: Gesammelte Fehler oder eine leere Liste nach erfolgreicher Speicherung.

    Jede Zeile enthält eine Teilnahme mit ihren Untersuchungsangaben.
    Vorhandene Untersuchungen werden über den Studiennamen zugeordnet,
    unbekannte Untersuchungen neu angelegt. Vorhandene Untersuchungsangaben
    und Uploadzeitpunkte bleiben erhalten. Jede Fehlermeldung verhindert
    die Speicherung der gesamten Datei.
    """
    arbeitsmappe = None
    fehlerliste = []

    try:
        arbeitsmappe, fehlerliste = _lese_xlsx(datei)
        if arbeitsmappe is None:
            return fehlerliste

        blatt = arbeitsmappe.worksheets[0]
        beschriftungen = (
            "Matrikelnummer",
            "Studienname",
            "Stundenwert (VP-Stunden)",
            "Einordnung (P/O)",
            "Verantwortliche Lehrperson",
        )
        # Die Menge speichert Adressen bereits erkannter Formelzellen ohne Dopplungen.
        formelzellen = set()

        # _cells ist dieselbe private openpyxl-Schnittstelle wie beim regulären Import.
        # Nur vorhandene Zellen prüfen, damit entfernte formatierte Leerzellen erlaubt bleiben.
        for zelle in blatt._cells.values():
            if zelle.data_type == "f":
                formelzellen.add(zelle.coordinate)
                fehlerliste.append(f"{zelle.coordinate}: Formeln sind nicht zulässig.")

            if zelle.value is not None and not (
                1 <= zelle.row <= 101 and 1 <= zelle.column <= 5
            ):
                fehlerliste.append(
                    f"{zelle.coordinate}: Außerhalb der Vorlage sind keine Inhalte zulässig."
                )

        # enumerate(..., start=1) liefert Paare aus Spaltennummer und Beschriftung.
        # Die Nummerierung 1 bis 5 entspricht den Excel-Spalten A bis E.
        for spalte, erwartet in enumerate(beschriftungen, start=1):
            zelle = blatt.cell(row=1, column=spalte)
            if zelle.coordinate not in formelzellen and zelle.value != erwartet:
                fehlerliste.append(
                    f"{zelle.coordinate}: Die Beschriftung muss „{erwartet}“ lauten."
                )

        # Die Reihenfolge entspricht den Spalten A bis E.
        # So kann jede Zelle mit ihrem Ergebnisschlüssel und ihrem Prüfer verbunden werden.
        pruefungen = (
            ("matrikelnummer", pruefe_matrikelnummer),
            ("studienname", pruefe_studienname),
            ("stundenwert", pruefe_stundenwert),
            ("einordnung", pruefe_einordnung),
            ("lehrperson", pruefe_lehrperson),
        )
        # None bedeutet: Es wurde noch keine gültige Matrikelnummer in der Datei gefunden.
        datei_matrikelnummer = None
        studiennamen = set()
        daten_zum_speichern = []
        teilnahmezeile_vorhanden = False

        # range(2, 102) umfasst die Zeilen 2 bis 101. Lücken beenden den Durchlauf nicht.
        for zeile in range(2, 102):
            zellen = [
                blatt.cell(row=zeile, column=spalte)
                for spalte in range(1, 6)
            ]
            # all() wird nur wahr, wenn jede der fünf Zellen None enthält.
            # Eine Zeile aus Leerzeichen ist deshalb keine vollständig leere Zeile.
            if all(zelle.value is None for zelle in zellen):
                continue

            teilnahmezeile_vorhanden = True
            werte = {}

            # zip() verbindet jede der fünf Zellen mit dem passenden Paar aus
            # Ergebnisschlüssel und Prüffunktion. Auch teilweise gefüllte Zeilen prüfen.
            for zelle, (schluessel, pruefung) in zip(zellen, pruefungen):
                if zelle.coordinate in formelzellen:
                    continue

                try:
                    # Der Prüfer erhält den ursprünglichen Excel-Zellwert.
                    # Sein gültiges Ergebnis wird unter dem zugeordneten Schlüssel abgelegt.
                    werte[schluessel] = pruefung(zelle.value)
                except ValidationError as fehler:
                    # Weitere Zellen bleiben unabhängig von diesem Feldfehler prüfbar.
                    for meldung in fehler.messages:
                        fehlerliste.append(f"{zelle.coordinate}: {meldung}")
                except OverflowError:
                    # Einen technisch nicht darstellbaren Zahlenwert feldbezogen melden.
                    fehlerliste.append(
                        f"{zelle.coordinate}: Der Zahlenwert ist nicht darstellbar."
                    )

            # Alle Teilnahmezeilen müssen dieselbe Matrikelnummer enthalten.
            if "matrikelnummer" in werte:
                matrikelnummer = werte["matrikelnummer"]
                if datei_matrikelnummer is None:
                    datei_matrikelnummer = matrikelnummer
                elif matrikelnummer != datei_matrikelnummer:
                    fehlerliste.append(
                        f"A{zeile}: Alle Teilnahmezeilen müssen dieselbe Matrikelnummer enthalten."
                    )

            # Ohne gültigen Studiennamen kann weder eine Dublette noch der Bestand
            # zuverlässig bestimmt werden. Andere Felder wurden bereits geprüft.
            if "studienname" not in werte:
                continue

            anzeigename, vergleichswert = werte["studienname"]
            # Für Dubletten zählt der bereinigte Vergleichswert, nicht die Anzeigeschreibweise.
            if vergleichswert in studiennamen:
                fehlerliste.append(
                    f"B{zeile}: Die Untersuchung kommt innerhalb der Datei mehrfach vor."
                )
            studiennamen.add(vergleichswert)

            # first() liefert das vorhandene Objekt oder None, wenn kein Treffer existiert.
            untersuchung = Untersuchung.objects.filter(
                studienname_vergleich=vergleichswert
            ).first()

            # Bei vorhandenem Objekt die bereits gültigen Angaben mit dem Bestand vergleichen.
            # Unbekannte Untersuchungen sind zulässig und werden erst nach allen Prüfungen angelegt.
            if untersuchung is not None:
                # getattr() liest das durch schluessel benannte Modellattribut,
                # zum Beispiel untersuchung.stundenwert. Das vermeidet doppelte Vergleichsblöcke.
                for schluessel, spalte in [
                    ("stundenwert", "C"),
                    ("einordnung", "D"),
                ]:
                    if (
                        schluessel in werte
                        and werte[schluessel] != getattr(untersuchung, schluessel)
                    ):
                        fehlerliste.append(
                            f"{spalte}{zeile}: Die Angabe weicht von der gespeicherten Untersuchung ab."
                        )

                # Der Importwert ist bereits außen bereinigt. Auch den gespeicherten
                # Namen außen bereinigen und beide Schreibweisen mit casefold() vergleichen.
                # Innere Leerzeichen und abweichende Namensbestandteile bleiben relevant.
                if (
                    "lehrperson" in werte
                    and werte["lehrperson"].casefold()
                    != untersuchung.lehrperson.strip().casefold()
                ):
                    fehlerliste.append(
                        f"E{zeile}: Die Lehrperson weicht von der gespeicherten Untersuchung ab."
                    )

                # Die gleiche Kombination aus Untersuchung und Matrikelnummer darf
                # nicht erneut angelegt werden. Die Feldwerte müssen dafür gültig sein.
                if (
                    "matrikelnummer" in werte
                    and Teilnahme.objects.filter(
                        untersuchung=untersuchung,
                        matrikelnummer=werte["matrikelnummer"],
                    ).exists()
                ):
                    fehlerliste.append(
                        f"Zeile {zeile}: Die Teilnahmezuordnung ist bereits gespeichert."
                    )

            # Fünf Ergebnisse bedeuten, dass alle fünf Feldprüfungen erfolgreich waren.
            # Das Paar merkt das vorhandene Objekt oder None zusammen mit den Werten.
            # Es handelt sich weiterhin nur um eine Vorbereitung im Arbeitsspeicher.
            if len(werte) == 5:
                daten_zum_speichern.append((untersuchung, werte))

        if not teilnahmezeile_vorhanden:
            fehlerliste.append("Datei: Es sind keine Teilnahmeangaben vorhanden.")

        if fehlerliste:
            return fehlerliste

        # Erst nach fehlerfreier Gesamtprüfung alle neuen Datensätze gemeinsam anlegen.
        with transaction.atomic():
            # Die Reihenfolge der geprüften Excel-Zeilen für die Anlage beibehalten.
            for untersuchung, werte in daten_zum_speichern:

                # None kennzeichnet eine unbekannte Untersuchung: neu anlegen.
                # Bei vorhandenem Objekt wird dieser Zweig übersprungen.
                if untersuchung is None:
                    anzeigename, vergleichswert = werte["studienname"]
                    untersuchung = Untersuchung.objects.create(
                        studienname=anzeigename,
                        studienname_vergleich=vergleichswert,
                        stundenwert=werte["stundenwert"],
                        einordnung=werte["einordnung"],
                        lehrperson=werte["lehrperson"],
                    )

                # create() speichert die Teilnahme und verknüpft sie über den Fremdschlüssel.
                # Eine bereits vorhandene Untersuchung selbst wird dabei nicht verändert.
                Teilnahme.objects.create(
                    untersuchung=untersuchung,
                    matrikelnummer=werte["matrikelnummer"],
                )

        return []

    except DatabaseError:
        # Bei einem Fehler im atomic()-Block sind dessen Schreibschritte zurückgenommen.
        # Ein Fehler beim vorherigen Lesen hat noch keine Speicherung ausgelöst.
        return fehlerliste + ["Datei: Der Übergangsimport konnte nicht übernommen werden."]
    except Exception:
        # Vorher gesammelte Meldungen erhalten und technische Interna ausblenden.
        return fehlerliste + ["Datei: Der Übergangsimport konnte nicht verarbeitet werden."]
    finally:
        # Arbeitsmappe und Upload auch bei einer frühen Abweisung schließen.
        try:
            if arbeitsmappe is not None:
                arbeitsmappe.close()
        finally:
            datei.close()
